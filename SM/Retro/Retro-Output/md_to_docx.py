"""Convert a Retrospective markdown report to a Word document.

Usage: python md_to_docx.py "Retrospective-Report-03-July-2026.md"
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

SCRIPT_DIR = os.path.dirname(__file__)
md_name = sys.argv[1] if len(sys.argv) > 1 else "Retrospective-Report-03-July-2026.md"
INPUT_FILE = os.path.join(SCRIPT_DIR, md_name)
OUTPUT_FILE = os.path.join(SCRIPT_DIR, os.path.splitext(md_name)[0] + ".docx")

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Calibri'
    hs.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)


def add_table_from_rows(doc, header_row, data_rows):
    cols = len(header_row)
    table = doc.add_table(rows=1, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, cell_text in enumerate(header_row):
        cell = table.rows[0].cells[i]
        cell.text = cell_text.strip()
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for row_data in data_rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            if i < cols:
                row_cells[i].text = cell_text.strip()
                for p in row_cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
    doc.add_paragraph()


def add_rich_paragraph(doc, text, style_name=None, bold=False):
    p = doc.add_paragraph(style=style_name)
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            run = p.add_run(part)
    if bold:
        for run in p.runs:
            run.bold = True
    return p


with open(INPUT_FILE, 'r', encoding='utf-8') as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip('\n')

    if re.match(r'^-{3,}$', line.strip()):
        i += 1
        continue

    heading_match = re.match(r'^(#{1,6})\s+(.*)', line)
    if heading_match:
        level = len(heading_match.group(1))
        text = heading_match.group(2).strip()
        doc_level = min(level, 3)
        doc.add_heading(text, level=doc_level)
        i += 1
        continue

    if '|' in line and i + 1 < len(lines) and re.match(r'^\|[-\s|:]+\|$', lines[i + 1].strip()):
        header_cells = [c.strip() for c in line.strip().strip('|').split('|')]
        i += 2
        data_rows = []
        while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
            row_cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            data_rows.append(row_cells)
            i += 1
        add_table_from_rows(doc, header_cells, data_rows)
        continue

    list_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)', line)
    if list_match:
        marker = list_match.group(2)
        text = list_match.group(3)
        style_name = 'List Bullet' if marker in ['-', '*'] else 'List Number'
        text = re.sub(r'^\[[ x]\]\s*', '', text)
        add_rich_paragraph(doc, text, style_name=style_name)
        i += 1
        continue

    if line.strip() == '':
        i += 1
        continue

    add_rich_paragraph(doc, line)
    i += 1

doc.save(OUTPUT_FILE)
print(f"Word document saved: {OUTPUT_FILE}")
