"""
Generate comprehensive test cases for US8992 - Streamlining Compte Statistics Logic
Based on template: FTC02-PRD to Test Cases (Comprehensive)
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Define test cases based on US8992 requirements
test_cases = [
    # Functional - Happy Path Test Cases
    {
        "TID": "TC001",
        "Category": "Functional - Happy Path",
        "Description": "Verify Activity Compte field is populated correctly when valid Business Type for Compte Statistics and Expected Contract Type are provided",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Business Type for Compte Statistics field is available (SFSUPPORTDE_BLM 8970)\n3. Expected Contract Type field is configured\n4. Mapping rules from Excel file are implemented",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity or edit existing one\n3. Select a valid 'Business Type for Compte Statistics' value\n4. Select a valid 'Expected Contract Type' value\n5. Save the Opportunity record",
        "Expected": "Activity Compte field is automatically populated with the correct value as per the mapping defined in 20260320_TranslationLogicsCompteSFDC_Salesforce@SopraSteria_v5.0.xlsx",
        "Priority": "High"
    },
    {
        "TID": "TC002",
        "Category": "Functional - Happy Path",
        "Description": "Verify Service Type Compte field is populated correctly based on input parameters",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Business Type for Compte Statistics field is available\n3. Expected Contract Type field is configured\n4. Mapping rules are implemented",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity or edit existing one\n3. Select a valid 'Business Type for Compte Statistics' value\n4. Select a valid 'Expected Contract Type' value\n5. Save the Opportunity record",
        "Expected": "Service Type Compte field is automatically populated with the correct value as per the mapping defined in the Excel specification file",
        "Priority": "High"
    },
    {
        "TID": "TC003",
        "Category": "Functional - Happy Path",
        "Description": "Verify Commitment Type Compte field is populated correctly based on input parameters",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Business Type for Compte Statistics field is available\n3. Expected Contract Type field is configured\n4. Mapping rules are implemented",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity or edit existing one\n3. Select a valid 'Business Type for Compte Statistics' value\n4. Select a valid 'Expected Contract Type' value\n5. Save the Opportunity record",
        "Expected": "Commitment Type Compte field is automatically populated with the correct value as per the mapping defined in the Excel specification file",
        "Priority": "High"
    },
    {
        "TID": "TC004",
        "Category": "Functional - Happy Path",
        "Description": "Verify Contract Type Compte field is populated correctly based on input parameters",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Business Type for Compte Statistics field is available\n3. Expected Contract Type field is configured\n4. Mapping rules are implemented",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity or edit existing one\n3. Select a valid 'Business Type for Compte Statistics' value\n4. Select a valid 'Expected Contract Type' value\n5. Save the Opportunity record",
        "Expected": "Contract Type Compte field is automatically populated with the correct value as per the mapping defined in the Excel specification file",
        "Priority": "High"
    },
    {
        "TID": "TC005",
        "Category": "Functional - Happy Path",
        "Description": "Verify all four Compte Statistics fields are populated simultaneously when Opportunity is created",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Both input parameters are configured\n3. Translation Salesforce Compte Statistics flow is active",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Click 'New Opportunity'\n3. Fill in required fields\n4. Select 'Business Type for Compte Statistics' value\n5. Select 'Expected Contract Type' value\n6. Save the Opportunity",
        "Expected": "All four Compte Statistics fields (Activity Compte, Service Type Compte, Commitment Type Compte, Contract Type Compte) are populated correctly in a single save operation",
        "Priority": "High"
    },
    {
        "TID": "TC006",
        "Category": "Functional - Happy Path",
        "Description": "Verify Compte Statistics fields are updated when Business Type for Compte Statistics is changed on an existing Opportunity",
        "Pre-conditions": "1. An existing Opportunity with populated Compte Statistics fields exists\n2. User has edit permissions",
        "Steps": "1. Navigate to an existing Opportunity record\n2. Edit the record\n3. Change 'Business Type for Compte Statistics' to a different valid value\n4. Keep the same 'Expected Contract Type'\n5. Save the Opportunity",
        "Expected": "All four Compte Statistics fields are updated according to the new combination as per the mapping rules",
        "Priority": "High"
    },
    {
        "TID": "TC007",
        "Category": "Functional - Happy Path",
        "Description": "Verify Compte Statistics fields are updated when Expected Contract Type is changed on an existing Opportunity",
        "Pre-conditions": "1. An existing Opportunity with populated Compte Statistics fields exists\n2. User has edit permissions",
        "Steps": "1. Navigate to an existing Opportunity record\n2. Edit the record\n3. Keep the same 'Business Type for Compte Statistics'\n4. Change 'Expected Contract Type' to a different valid value\n5. Save the Opportunity",
        "Expected": "All four Compte Statistics fields are updated according to the new combination as per the mapping rules",
        "Priority": "High"
    },
    {
        "TID": "TC008",
        "Category": "Functional - Happy Path",
        "Description": "Verify Compte Statistics values are correctly transferred to Contract when Contract is generated from Opportunity",
        "Pre-conditions": "1. An Opportunity with correctly populated Compte Statistics fields exists\n2. Flow 'Update Compte Statistics on Contract Record' is active\n3. User has permission to create Contracts",
        "Steps": "1. Navigate to the Opportunity with populated Compte Statistics\n2. Generate a Contract from the Opportunity\n3. View the newly created Contract record",
        "Expected": "The four Compte Statistics fields on the Contract record contain the same values as the source Opportunity",
        "Priority": "High"
    },
    
    # Field Renaming Test Cases
    {
        "TID": "TC009",
        "Category": "Functional - Field Renaming",
        "Description": "Verify 'Activity Compte Hidden' field has been renamed to 'Activity Compte'",
        "Pre-conditions": "1. Field renaming deployment has been completed\n2. User has access to Opportunity object",
        "Steps": "1. Navigate to Salesforce Setup\n2. Go to Object Manager > Opportunity > Fields & Relationships\n3. Search for 'Activity Compte' field\n4. Verify the field label",
        "Expected": "Field is labeled 'Activity Compte' (not 'Activity Compte Hidden')",
        "Priority": "Medium"
    },
    {
        "TID": "TC010",
        "Category": "Functional - Field Renaming",
        "Description": "Verify 'Commitment Type Compte Hidden' field has been renamed to 'Commitment Type Compte'",
        "Pre-conditions": "1. Field renaming deployment has been completed\n2. User has access to Opportunity object",
        "Steps": "1. Navigate to Salesforce Setup\n2. Go to Object Manager > Opportunity > Fields & Relationships\n3. Search for 'Commitment Type Compte' field\n4. Verify the field label",
        "Expected": "Field is labeled 'Commitment Type Compte' (not 'Commitment Type Compte Hidden')",
        "Priority": "Medium"
    },
    {
        "TID": "TC011",
        "Category": "Functional - Field Renaming",
        "Description": "Verify 'Service Type Compte Hidden' field has been renamed to 'Service Type Compte'",
        "Pre-conditions": "1. Field renaming deployment has been completed\n2. User has access to Opportunity object",
        "Steps": "1. Navigate to Salesforce Setup\n2. Go to Object Manager > Opportunity > Fields & Relationships\n3. Search for 'Service Type Compte' field\n4. Verify the field label",
        "Expected": "Field is labeled 'Service Type Compte' (not 'Service Type Compte Hidden')",
        "Priority": "Medium"
    },
    {
        "TID": "TC012",
        "Category": "Functional - Field Renaming",
        "Description": "Verify 'Vertragsart Compte Hidden' field has been renamed to 'Contract Type Compte'",
        "Pre-conditions": "1. Field renaming deployment has been completed\n2. User has access to Opportunity object",
        "Steps": "1. Navigate to Salesforce Setup\n2. Go to Object Manager > Opportunity > Fields & Relationships\n3. Search for 'Contract Type Compte' field\n4. Verify the field label",
        "Expected": "Field is labeled 'Contract Type Compte' (not 'Vertragsart Compte Hidden')",
        "Priority": "Medium"
    },
    
    # Flow and Automation Test Cases
    {
        "TID": "TC013",
        "Category": "Functional - Automation",
        "Description": "Verify 'Translation Salesforce Compte Statistics' flow is updated to use only two input parameters",
        "Pre-conditions": "1. Admin access to Salesforce Flow Builder\n2. Flow deployment has been completed",
        "Steps": "1. Navigate to Salesforce Setup > Flows\n2. Open 'Translation Salesforce Compte Statistics' flow\n3. Review the flow trigger and decision elements\n4. Verify input parameters",
        "Expected": "Flow only uses 'Business Type for Compte Statistics' and 'Expected Contract Type' as input parameters for determining Compte Statistics values",
        "Priority": "High"
    },
    {
        "TID": "TC014",
        "Category": "Functional - Automation",
        "Description": "Verify 'Update Compte Statistics on Contract Record' flow updates only the four required fields",
        "Pre-conditions": "1. Admin access to Salesforce Flow Builder\n2. Flow deployment has been completed",
        "Steps": "1. Navigate to Salesforce Setup > Flows\n2. Open 'Update Compte Statistics on Contract Record' flow\n3. Review the assignment elements\n4. Verify fields being updated",
        "Expected": "Flow only updates Activity Compte, Service Type Compte, Commitment Type Compte, and Contract Type Compte fields",
        "Priority": "High"
    },
    {
        "TID": "TC015",
        "Category": "Functional - Automation",
        "Description": "Verify all legacy/deprecated logic paths have been removed from the automation",
        "Pre-conditions": "1. Admin access to Salesforce Flow Builder\n2. Access to previous flow versions for comparison",
        "Steps": "1. Navigate to Salesforce Setup > Flows\n2. Open 'Translation Salesforce Compte Statistics' flow\n3. Review all decision elements and assignment nodes\n4. Compare with mapping table in Excel file",
        "Expected": "No legacy combinations exist in the flow; only combinations from the mapping table are present",
        "Priority": "High"
    },
    
    # Negative Test Cases
    {
        "TID": "TC016",
        "Category": "Negative Scenario",
        "Description": "Verify behavior when Business Type for Compte Statistics is blank/not selected",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Automation flows are active",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity\n3. Leave 'Business Type for Compte Statistics' blank\n4. Select a valid 'Expected Contract Type'\n5. Save the Opportunity",
        "Expected": "Compte Statistics fields remain empty or display appropriate error message (Needs clarification: exact expected behavior)",
        "Priority": "Medium"
    },
    {
        "TID": "TC017",
        "Category": "Negative Scenario",
        "Description": "Verify behavior when Expected Contract Type is blank/not selected",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Automation flows are active",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity\n3. Select a valid 'Business Type for Compte Statistics'\n4. Leave 'Expected Contract Type' blank\n5. Save the Opportunity",
        "Expected": "Compte Statistics fields remain empty or display appropriate error message (Needs clarification: exact expected behavior)",
        "Priority": "Medium"
    },
    {
        "TID": "TC018",
        "Category": "Negative Scenario",
        "Description": "Verify behavior when both input parameters are blank/not selected",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Automation flows are active",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity\n3. Leave 'Business Type for Compte Statistics' blank\n4. Leave 'Expected Contract Type' blank\n5. Save the Opportunity",
        "Expected": "Compte Statistics fields remain empty; no errors occur during save",
        "Priority": "Medium"
    },
    {
        "TID": "TC019",
        "Category": "Negative Scenario",
        "Description": "Verify invalid/legacy combinations do not populate Compte Statistics fields",
        "Pre-conditions": "1. User has access to Salesforce Opportunity object\n2. Legacy combinations have been removed from automation\n3. Test data for invalid combination is available",
        "Steps": "1. Navigate to Salesforce Opportunities\n2. Create a new Opportunity\n3. Select a 'Business Type for Compte Statistics' value\n4. Select an 'Expected Contract Type' value that forms an invalid/unmapped combination\n5. Save the Opportunity",
        "Expected": "Compte Statistics fields are not populated for invalid combinations (Needs clarification: confirm if this combination exists in test environment)",
        "Priority": "High"
    },
    
    # Boundary/Edge Cases
    {
        "TID": "TC020",
        "Category": "Edge Case",
        "Description": "Verify Compte Statistics fields when Opportunity is cloned",
        "Pre-conditions": "1. An Opportunity with populated Compte Statistics exists\n2. User has clone permission",
        "Steps": "1. Navigate to an Opportunity with populated Compte Statistics\n2. Clone the Opportunity\n3. Modify the cloned Opportunity's input parameters if needed\n4. Save the cloned Opportunity",
        "Expected": "Cloned Opportunity recalculates Compte Statistics based on its own input parameters",
        "Priority": "Medium"
    },
    {
        "TID": "TC021",
        "Category": "Edge Case",
        "Description": "Verify Compte Statistics behavior during bulk Opportunity updates",
        "Pre-conditions": "1. Multiple Opportunities exist with various input parameter combinations\n2. User has bulk edit permissions\n3. Data Loader or bulk update tool is available",
        "Steps": "1. Prepare a list of Opportunities for bulk update\n2. Update 'Business Type for Compte Statistics' for multiple records\n3. Execute bulk update\n4. Verify each Opportunity's Compte Statistics",
        "Expected": "All Opportunities are correctly updated with appropriate Compte Statistics values based on individual mappings",
        "Priority": "Medium"
    },
    {
        "TID": "TC022",
        "Category": "Edge Case",
        "Description": "Verify Compte Statistics update when input parameter values change rapidly/consecutively",
        "Pre-conditions": "1. An existing Opportunity exists\n2. User has edit permissions",
        "Steps": "1. Navigate to an existing Opportunity\n2. Change 'Business Type for Compte Statistics' value\n3. Save immediately\n4. Quickly edit again and change 'Expected Contract Type'\n5. Save again\n6. Verify final Compte Statistics values",
        "Expected": "Final Compte Statistics values reflect the last saved combination accurately",
        "Priority": "Low"
    },
    {
        "TID": "TC023",
        "Category": "Edge Case",
        "Description": "Verify Compte Statistics fields behavior when Opportunity stage changes to Closed Won",
        "Pre-conditions": "1. An Opportunity with populated Compte Statistics exists\n2. Opportunity is not yet Closed",
        "Steps": "1. Navigate to an existing Opportunity\n2. Change stage to 'Closed Won'\n3. Save the Opportunity\n4. Verify Compte Statistics fields",
        "Expected": "Compte Statistics fields retain their correct values when Opportunity is closed",
        "Priority": "Medium"
    },
    {
        "TID": "TC024",
        "Category": "Edge Case",
        "Description": "Verify Compte Statistics when Opportunity is converted to Contract multiple times (re-contracted)",
        "Pre-conditions": "1. An Opportunity that has previously generated a Contract exists\n2. Business process allows multiple contracts",
        "Steps": "1. Navigate to an Opportunity that already has an associated Contract\n2. Generate another Contract from the same Opportunity\n3. Verify Compte Statistics on the new Contract",
        "Expected": "Each Contract generated from the Opportunity inherits the correct Compte Statistics from the Opportunity at the time of generation",
        "Priority": "Low"
    },
    
    # Data Integrity Test Cases
    {
        "TID": "TC025",
        "Category": "Data Integrity",
        "Description": "Verify all mapping combinations from Excel file produce correct outputs",
        "Pre-conditions": "1. Complete mapping table from 20260320_TranslationLogicsCompteSFDC_Salesforce@SopraSteria_v5.0.xlsx is available\n2. Test environment has all valid input values configured",
        "Steps": "1. For each row in the Mapping sheet:\n   a. Create or update an Opportunity\n   b. Set Business Type for Compte Statistics as per mapping\n   c. Set Expected Contract Type as per mapping\n   d. Save and verify all four output fields\n2. Document results for each combination",
        "Expected": "Each combination produces exactly the output values specified in the Mapping sheet",
        "Priority": "Critical"
    },
    {
        "TID": "TC026",
        "Category": "Data Integrity",
        "Description": "Verify existing Opportunities retain correct Compte Statistics after migration",
        "Pre-conditions": "1. Migration sub-task has been completed\n2. Pre-migration data snapshot is available for comparison",
        "Steps": "1. Compare Compte Statistics values on existing Opportunities before and after migration\n2. Verify renamed fields contain expected data\n3. Run reports on Compte Statistics fields",
        "Expected": "Existing data is preserved correctly after field renaming and automation updates",
        "Priority": "Critical"
    },
    {
        "TID": "TC027",
        "Category": "Data Integrity",
        "Description": "Verify deprecated Compte fields are no longer being populated",
        "Pre-conditions": "1. Automation updates have been deployed\n2. List of deprecated fields is available",
        "Steps": "1. Create a new Opportunity with valid input parameters\n2. Check all deprecated Compte Statistics fields\n3. Verify deprecated fields are not populated by automation",
        "Expected": "Only the four required Compte Statistics fields (Activity Compte, Service Type Compte, Commitment Type Compte, Contract Type Compte) are populated; no deprecated fields are updated",
        "Priority": "High"
    },
    
    # Regression Test Cases
    {
        "TID": "TC028",
        "Category": "Regression",
        "Description": "Verify Opportunity creation process works correctly after changes",
        "Pre-conditions": "1. All changes have been deployed\n2. User has standard Opportunity creation permissions",
        "Steps": "1. Navigate to Opportunities tab\n2. Create a new Opportunity with all required fields\n3. Include Compte Statistics input parameters\n4. Save successfully",
        "Expected": "Opportunity is created successfully with Compte Statistics populated; no errors or unexpected behavior",
        "Priority": "High"
    },
    {
        "TID": "TC029",
        "Category": "Regression",
        "Description": "Verify Opportunity update process works correctly after changes",
        "Pre-conditions": "1. All changes have been deployed\n2. Existing Opportunities are available",
        "Steps": "1. Navigate to an existing Opportunity\n2. Edit various fields including non-Compte related fields\n3. Save the Opportunity",
        "Expected": "Opportunity is updated successfully; Compte Statistics are recalculated if input parameters changed, otherwise retained",
        "Priority": "High"
    },
    {
        "TID": "TC030",
        "Category": "Regression",
        "Description": "Verify Contract creation from Opportunity works correctly after changes",
        "Pre-conditions": "1. All changes have been deployed\n2. Contract generation process is configured",
        "Steps": "1. Navigate to an Opportunity ready for contract generation\n2. Generate a Contract using standard process\n3. Verify Contract is created with correct Compte Statistics",
        "Expected": "Contract is created successfully with Compte Statistics inherited from Opportunity",
        "Priority": "High"
    },
    
    # User Profile/Permission Test Cases
    {
        "TID": "TC031",
        "Category": "Functional - Permissions",
        "Description": "Verify Area Controlling employee can view Compte Statistics fields on Opportunity",
        "Pre-conditions": "1. User is logged in as Area Controlling employee\n2. Field level security is configured appropriately",
        "Steps": "1. Log in as Area Controlling employee\n2. Navigate to an Opportunity record\n3. View the Compte Statistics section/fields",
        "Expected": "Area Controlling employee can view all four Compte Statistics fields with correct values",
        "Priority": "Medium"
    },
    {
        "TID": "TC032",
        "Category": "Functional - Permissions",
        "Description": "Verify Compte Statistics fields are read-only for users (auto-populated by system)",
        "Pre-conditions": "1. User has standard Opportunity edit permissions\n2. Field-level security is configured",
        "Steps": "1. Navigate to an Opportunity record\n2. Click Edit\n3. Attempt to manually modify Compte Statistics fields",
        "Expected": "Compte Statistics fields are read-only and cannot be manually edited (auto-populated by automation only)",
        "Priority": "Medium"
    },
    
    # Reporting Test Cases
    {
        "TID": "TC033",
        "Category": "Functional - Reporting",
        "Description": "Verify Compte Statistics fields are available in Opportunity reports",
        "Pre-conditions": "1. Field renaming has been completed\n2. User has report creation permissions",
        "Steps": "1. Navigate to Reports\n2. Create a new Opportunity report\n3. Add new Compte Statistics field names as columns\n4. Run the report",
        "Expected": "Reports display correct field names (Activity Compte, Service Type Compte, Commitment Type Compte, Contract Type Compte) and accurate data",
        "Priority": "Medium"
    },
    {
        "TID": "TC034",
        "Category": "Functional - Reporting",
        "Description": "Verify Group Controlling reporting requirements are met with new Compte Statistics",
        "Pre-conditions": "1. All changes have been deployed\n2. Sample data with various combinations exists",
        "Steps": "1. Generate Compte Statistics report using existing report template\n2. Verify data consistency and accuracy\n3. Compare output format with Group Controlling requirements",
        "Expected": "Reports provide accurate, consistent statistical classification aligned with Group Controlling expectations",
        "Priority": "High"
    }
]

def create_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "US8992_Test_Cases"
    
    # Define headers
    headers = ["TID", "Category", "Description", "Pre-conditions", "Steps", "Expected", "Priority"]
    
    # Style definitions
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Priority color fills
    priority_fills = {
        "Critical": PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
        "High": PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
        "Medium": PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
        "Low": PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    }
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Write test cases
    for row_num, tc in enumerate(test_cases, 2):
        ws.cell(row=row_num, column=1, value=tc["TID"]).border = thin_border
        ws.cell(row=row_num, column=2, value=tc["Category"]).border = thin_border
        ws.cell(row=row_num, column=3, value=tc["Description"]).border = thin_border
        ws.cell(row=row_num, column=4, value=tc["Pre-conditions"]).border = thin_border
        ws.cell(row=row_num, column=5, value=tc["Steps"]).border = thin_border
        ws.cell(row=row_num, column=6, value=tc["Expected"]).border = thin_border
        priority_cell = ws.cell(row=row_num, column=7, value=tc["Priority"])
        priority_cell.border = thin_border
        
        # Apply priority color
        if tc["Priority"] in priority_fills:
            priority_cell.fill = priority_fills[tc["Priority"]]
        
        # Apply alignment to all cells in row
        for col_num in range(1, 8):
            ws.cell(row=row_num, column=col_num).alignment = cell_alignment
    
    # Set column widths
    column_widths = {
        'A': 8,   # TID
        'B': 25,  # Category
        'C': 50,  # Description
        'D': 40,  # Pre-conditions
        'E': 50,  # Steps
        'F': 50,  # Expected
        'G': 12   # Priority
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Freeze header row
    ws.freeze_panes = 'A2'
    
    # Add auto-filter
    ws.auto_filter.ref = f"A1:G{len(test_cases) + 1}"
    
    # Save workbook
    output_path = r"C:\Apparao Maroju\AI-SSI\Project-SF(POC)\Output\US8992_TestCases.xlsx"
    wb.save(output_path)
    print(f"Test cases successfully saved to: {output_path}")
    print(f"Total test cases generated: {len(test_cases)}")

if __name__ == "__main__":
    create_xlsx()
