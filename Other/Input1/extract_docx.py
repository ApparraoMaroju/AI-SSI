from docx import Document

doc = Document(r'c:\Apparao Maroju\AI-SSI\Other\Input1\Spec_Driven_Development_Spec_Kit_Professional.docx')

print("=== DOCUMENT CONTENT ===\n")

for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)

print("\n=== TABLES ===\n")

for i, table in enumerate(doc.tables):
    print(f"\n--- Table {i+1} ---")
    for row in table.rows:
        row_text = ' | '.join([cell.text for cell in row.cells])
        print(row_text)
