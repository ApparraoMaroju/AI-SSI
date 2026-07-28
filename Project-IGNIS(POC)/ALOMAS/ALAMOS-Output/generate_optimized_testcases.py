"""
ALAMOS Integration - Optimized Functional Integration Test Cases Generator
Generates deduplicated test cases with minimum count and maximum coverage
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

def create_optimized_test_cases_workbook():
    """Generate Excel workbook with optimized functional integration test cases (deduplicated)"""
    
    # Create workbook and worksheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    
    # Define columns
    columns = [
        "TID", "Test Case Name", "Category", "Description", "Pre-conditions",
        "Step No", "Step Description", "Expected Result", "Priority", "User Story Reference"
    ]
    
    # Style definitions
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Write header row
    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Freeze header row
    ws.freeze_panes = 'A2'
    
    # ============================================================================
    # OPTIMIZED TEST CASES - Consolidated with maximum coverage, minimum duplication
    # ============================================================================
    
    test_cases = [
        # ============================================================================
        # SITUATION REPORTS - Adapter Receive & Validate (US-ADAPTER-ALAMOS-001)
        # Consolidated from original TC-001 to TC-006
        # ============================================================================
        {
            "tid": "TC-001",
            "name": "[US-ADAPTER-ALAMOS-001] Verify complete Situation Report receive, validate, and acknowledge flow",
            "category": "Functional (Happy Path)",
            "description": "Validate complete Situation Report flow: receive from ALAMOS FE2, validate against ALAMOS contract, process through IPL, and return synchronous acknowledgment. Covers mandatory fields, headers, content-type, and authentication validation.",
            "preconditions": "1. Adapter REST endpoint is available\n2. ALAMOS FE2 is configured and connected\n3. Valid authentication credentials are available\n4. IPL endpoint is available",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-001",
            "steps": [
                ("1", "Prepare a valid Situation Report payload with all mandatory fields (external_id, incident_number, call_sign, situation_report.text, reported_at)", "Payload is correctly structured per ALAMOS contract"),
                ("2", "Send the Situation Report to the Adapter REST endpoint with valid headers (Content-Type: application/json) and authentication", "Request is received by the Adapter"),
                ("3", "Verify the Adapter validates request headers, authentication, and content-type before payload processing", "Technical prerequisites validation passes"),
                ("4", "Verify the Adapter validates the payload against the ALAMOS contract including all mandatory fields", "Payload validation passes without errors"),
                ("5", "Verify the Adapter accepts the request and forwards to FacadeService for IPL processing", "Situation Report is accepted and forwarded for processing"),
                ("6", "Verify IPL processes the request successfully", "IPL returns success response"),
                ("7", "Verify the Adapter returns a synchronous success acknowledgment to ALAMOS FE2", "HTTP 200 response with success acknowledgment is returned")
            ]
        },
        {
            "tid": "TC-002",
            "name": "[US-ADAPTER-ALAMOS-001] Verify Adapter rejects invalid requests (missing fields, invalid headers/auth/content-type, unsupported methods)",
            "category": "Negative Scenario",
            "description": "Validate that the Adapter correctly rejects requests with: missing mandatory fields, invalid authentication, invalid content-type, and unsupported HTTP methods. Consolidates all request validation negative scenarios.",
            "preconditions": "1. Adapter REST endpoint is available\n2. Authentication mechanism is enabled",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-001",
            "steps": [
                ("1", "Send a Situation Report with missing mandatory field (external_id empty)", "HTTP 400 with validation error indicating missing field"),
                ("2", "Send a request without authentication headers", "HTTP 401 Unauthorized response"),
                ("3", "Send a request with invalid Content-Type (text/plain instead of application/json)", "HTTP 415 Unsupported Media Type response"),
                ("4", "Send a request using GET method instead of POST", "HTTP 405 Method Not Allowed response"),
                ("5", "Verify all rejected requests do not proceed to downstream processing", "No Domain Event created for invalid requests")
            ]
        },
        {
            "tid": "TC-003",
            "name": "[US-ADAPTER-ALAMOS-001] Verify Adapter returns failure acknowledgment when IPL processing fails",
            "category": "Negative Scenario",
            "description": "Validate that the Adapter returns synchronous failure acknowledgment accurately reflecting IPL processing failure.",
            "preconditions": "1. Adapter REST endpoint is available\n2. IPL endpoint is configured to return a fault response",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-001",
            "steps": [
                ("1", "Configure IPL endpoint to return a business fault for the test scenario", "IPL mock/stub is configured to return fault"),
                ("2", "Send a valid Situation Report to the Adapter", "Request is received and validated by Adapter"),
                ("3", "Verify the Adapter forwards the request and waits for IPL response", "Adapter processes the request synchronously"),
                ("4", "Verify the Adapter returns synchronous failure acknowledgment to ALAMOS FE2", "HTTP response with failure acknowledgment accurately reflecting IPL fault is returned")
            ]
        },
        
        # ============================================================================
        # SITUATION REPORTS - Domain Event Generation (US-ADAPTER-ALAMOS-002)
        # Consolidated from original TC-007 to TC-011
        # ============================================================================
        {
            "tid": "TC-004",
            "name": "[US-ADAPTER-ALAMOS-002] Verify Domain Event generation with incident correlation (external_id primary, incident_number fallback)",
            "category": "Functional (Happy Path)",
            "description": "Validate complete Domain Event transformation including all required attributes per contract and incident correlation using external_id (primary) or incident_number (fallback).",
            "preconditions": "1. Adapter is running\n2. Valid Situation Report payloads are available",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-002",
            "steps": [
                ("1", "Submit a valid Situation Report with external_id populated", "Situation Report is received and validated"),
                ("2", "Verify the Adapter creates incident.situation_report_received/v1 Domain Event with correct event_type", "Domain Event is created"),
                ("3", "Verify incident correlation uses external_id from the payload (primary key)", "incident.external_id in Domain Event matches submitted external_id"),
                ("4", "Verify all required Domain Event attributes are populated: event_id (UUID), event_type, occurred_at (ISO 8601), source_system (ALAMOS), tenant_id, authoring_resource, situation_report", "All mandatory attributes present and correctly formatted"),
                ("5", "Submit a Situation Report with external_id empty/null and incident_number populated", "Situation Report is received and validated"),
                ("6", "Verify incident correlation falls back to incident_number when external_id unavailable", "incident.incident_number is used for correlation"),
                ("7", "Verify the Domain Event is validated prior to downstream processing", "Domain Event passes contract validation")
            ]
        },
        {
            "tid": "TC-005",
            "name": "[US-ADAPTER-ALAMOS-002] Verify call_sign fallback and error handling for incomplete correlation data",
            "category": "Edge Case / Negative",
            "description": "Validate call_sign fallback handling when missing, and error handling when both correlation keys (external_id and incident_number) are unavailable.",
            "preconditions": "1. Adapter is running",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-002",
            "steps": [
                ("1", "Submit a valid Situation Report with call_sign field empty or missing", "Situation Report is received by Adapter"),
                ("2", "Verify the Adapter applies fallback handling for call_sign", "authoring_resource.call_sign uses fallback value"),
                ("3", "Verify Domain Event is created successfully with fallback call_sign", "Domain Event is generated without error"),
                ("4", "Submit a Situation Report with both external_id and incident_number missing/empty", "Situation Report is received by Adapter"),
                ("5", "Verify the Adapter cannot complete incident correlation", "Both correlation keys are unavailable"),
                ("6", "Verify an appropriate error is returned to ALAMOS FE2", "Error response indicating incident correlation failure is returned")
            ]
        },
        
        # ============================================================================
        # FACADESERVICE - Routing Layer (US-ADAPTER-ALAMOS-003)
        # Consolidated from original TC-012 to TC-015
        # ============================================================================
        {
            "tid": "TC-006",
            "name": "[US-ADAPTER-ALAMOS-003] Verify FacadeService routing, synchronous processing, and response propagation",
            "category": "Functional (Happy Path)",
            "description": "Validate FacadeService correctly routes Domain Events through EVENT processing path, coordinates synchronous processing, and propagates both success and failure responses back to Adapter.",
            "preconditions": "1. Adapter and FacadeService are running\n2. Routing configuration for EVENT path is enabled\n3. IPL endpoint is available",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-003",
            "steps": [
                ("1", "Generate a valid incident.situation_report_received Domain Event", "Domain Event is created by Adapter"),
                ("2", "Verify FacadeService receives the Domain Event", "FacadeService receives the event"),
                ("3", "Verify FacadeService determines EVENT processing path based on capability-specific routing configuration", "Routing configuration resolves to EVENT path"),
                ("4", "Verify FacadeService coordinates synchronous processing via IPL", "Processing is coordinated synchronously"),
                ("5", "Verify IPL returns success response", "IPL processing completes successfully"),
                ("6", "Verify FacadeService propagates success response back to Adapter", "Adapter receives success indication"),
                ("7", "Configure IPL to return business fault, repeat steps 1-4", "Request is processed until IPL call"),
                ("8", "Verify FacadeService propagates failure response back to Adapter", "Adapter receives failure indication with fault details"),
                ("9", "Verify FacadeService operates without Inbox, Outbox, Event Bus, or Process Tracker", "No persistence mechanisms are invoked")
            ]
        },
        {
            "tid": "TC-007",
            "name": "[US-ADAPTER-ALAMOS-003] Verify FacadeService handles invalid routing configuration",
            "category": "Negative Scenario",
            "description": "Validate that the FacadeService returns error when routing configuration is invalid or capability is unsupported.",
            "preconditions": "1. Adapter and FacadeService are running\n2. Invalid routing configuration is set",
            "priority": "Medium",
            "user_story": "US-ADAPTER-ALAMOS-003",
            "steps": [
                ("1", "Configure FacadeService with invalid routing for Situation Report capability", "Invalid configuration is applied"),
                ("2", "Submit a valid Situation Report", "Request is received by Adapter"),
                ("3", "Verify FacadeService cannot resolve to valid processing path", "Invalid configuration is detected"),
                ("4", "Verify error response is returned to ALAMOS FE2", "Error indicating routing failure is propagated")
            ]
        },
        
        # ============================================================================
        # FACADESERVICE - IPL Processing (US-ADAPTER-ALAMOS-004)
        # Consolidated from original TC-016 to TC-020
        # ============================================================================
        {
            "tid": "TC-008",
            "name": "[US-ADAPTER-ALAMOS-004] Verify FacadeService IPL invocation with request mapping and response handling",
            "category": "Functional (Happy Path)",
            "description": "Validate FacadeService transforms Domain Event to IPL request structure and successfully invokes IPL erzeugeEinsatzmeldung endpoint synchronously.",
            "preconditions": "1. FacadeService is running\n2. IPL endpoint POST /einsatztransfer/erzeugeEinsatzmeldung is available",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-004",
            "steps": [
                ("1", "Generate incident.situation_report_received Domain Event with all required data", "Valid Domain Event is created"),
                ("2", "Verify FacadeService validates all required data for IPL request generation", "Validation passes"),
                ("3", "Verify FacadeService maps Domain Event data to IPL request structure", "Request mapping is performed correctly"),
                ("4", "Verify FacadeService invokes POST /einsatztransfer/erzeugeEinsatzmeldung synchronously", "SOAP request is sent to IPL"),
                ("5", "Verify IPL returns success response", "IPL processing completes successfully"),
                ("6", "Verify success result is propagated back to ALAMOS FE2", "Success acknowledgment is returned")
            ]
        },
        {
            "tid": "TC-009",
            "name": "[US-ADAPTER-ALAMOS-004] Verify FacadeService handles IPL errors (business fault, technical fault, unavailable, timeout)",
            "category": "Negative Scenario",
            "description": "Validate FacadeService correctly handles all IPL error scenarios: business faults, technical faults, endpoint unavailability, and timeouts.",
            "preconditions": "1. FacadeService is running\n2. IPL endpoint can be configured for various error scenarios",
            "priority": "High",
            "user_story": "US-ADAPTER-ALAMOS-004",
            "steps": [
                ("1", "Configure IPL to return a business fault, submit valid Situation Report", "Business fault is received"),
                ("2", "Verify business fault is propagated to ALAMOS FE2 with appropriate error details", "Failure acknowledgment with business fault details is returned"),
                ("3", "Configure IPL to return a technical fault (e.g., internal server error)", "Technical fault is received"),
                ("4", "Verify technical fault is propagated to ALAMOS FE2", "Failure acknowledgment with technical error is returned"),
                ("5", "Stop IPL endpoint, submit valid Situation Report", "Connection failure occurs"),
                ("6", "Verify communication failure is handled and propagated to ALAMOS FE2", "Failure acknowledgment indicating IPL unavailable is returned"),
                ("7", "Configure IPL to delay response beyond timeout threshold", "Timeout occurs"),
                ("8", "Verify timeout failure is propagated to ALAMOS FE2", "Failure acknowledgment indicating timeout is returned")
            ]
        },
        
        # ============================================================================
        # ALAMOS GATEWAY - Legacy Alerting (US-ALAMOS-GW-001)
        # Consolidated from original TC-021 to TC-024
        # ============================================================================
        {
            "tid": "TC-010",
            "name": "[US-ALAMOS-GW-001] Verify Gateway processes valid Alerting request and handles ALAMOS FE2 responses",
            "category": "Functional (Happy Path)",
            "description": "Validate ALAMOS Gateway receives Alerting from IPL, transforms to REST Alarm API format, forwards to ALAMOS FE2, and propagates response back to IPL.",
            "preconditions": "1. ALAMOS Gateway is running\n2. IPL is configured\n3. ALAMOS FE2 is available",
            "priority": "High",
            "user_story": "US-ALAMOS-GW-001",
            "steps": [
                ("1", "IPL sends a valid SOAP AlarmRequest to the ALAMOS Gateway", "AlarmRequest is transmitted"),
                ("2", "Verify ALAMOS Gateway receives and parses the request", "Gateway receives AlarmRequest"),
                ("3", "Verify Gateway transforms request to REST Alarm API format", "Request is mapped correctly"),
                ("4", "Verify Gateway forwards request to ALAMOS FE2", "REST Alarm API call is made"),
                ("5", "Verify ALAMOS FE2 returns success acknowledgment", "Success response is received"),
                ("6", "Verify Gateway propagates success back to IPL", "SOAP response indicating success is returned to IPL")
            ]
        },
        {
            "tid": "TC-011",
            "name": "[US-ALAMOS-GW-001] Verify Gateway handles Alerting errors (business error, communication failure, malformed request)",
            "category": "Negative Scenario",
            "description": "Validate ALAMOS Gateway correctly handles business errors from ALAMOS FE2, communication failures, and malformed requests from IPL.",
            "preconditions": "1. ALAMOS Gateway is running\n2. ALAMOS FE2 can be configured for error scenarios",
            "priority": "High",
            "user_story": "US-ALAMOS-GW-001",
            "steps": [
                ("1", "Configure ALAMOS FE2 to return a business error, send valid AlarmRequest from IPL", "Business error is received"),
                ("2", "Verify business error is propagated back to IPL", "IPL receives business error response"),
                ("3", "Simulate communication failure with ALAMOS FE2 (timeout/connectivity)", "Communication failure occurs"),
                ("4", "Verify communication failure is handled and propagated to IPL", "IPL receives appropriate error response"),
                ("5", "Send an invalid/malformed SOAP AlarmRequest from IPL", "Gateway validates the request"),
                ("6", "Verify malformed request is rejected", "Error response returned to IPL")
            ]
        },
        
        # ============================================================================
        # ALAMOS GATEWAY - Continuous Mission Updates (US-ALAMOS-GW-002)
        # Consolidated from original TC-025 to TC-029 (all 4 triggers in one test)
        # ============================================================================
        {
            "tid": "TC-012",
            "name": "[US-ALAMOS-GW-002] Verify Gateway processes Continuous Mission Updates for all triggers (incident/resource/keyword/location)",
            "category": "Functional (Happy Path)",
            "description": "Validate ALAMOS Gateway processes and forwards Continuous Mission Updates for all four supported triggers: incident status, resource status, keyword, and location changes. Tests backward compatibility with existing implementation.",
            "preconditions": "1. ALAMOS Gateway is running\n2. ALAMOS is available\n3. IPL can trigger all update types",
            "priority": "High",
            "user_story": "US-ALAMOS-GW-002",
            "steps": [
                ("1", "Trigger an incident status change in IPL", "Mission Update for incident status is generated"),
                ("2", "Verify Gateway processes and forwards the update to ALAMOS", "Update is forwarded successfully"),
                ("3", "Verify ALAMOS response is propagated back to IPL", "IPL receives acknowledgment"),
                ("4", "Trigger a resource status change in IPL", "Mission Update for resource status is generated"),
                ("5", "Verify update is processed and forwarded, response propagated to IPL", "Complete flow successful"),
                ("6", "Trigger a keyword change in IPL", "Mission Update for keyword change is generated"),
                ("7", "Verify update is processed and forwarded, response propagated to IPL", "Complete flow successful"),
                ("8", "Trigger a location change in IPL", "Mission Update for location change is generated"),
                ("9", "Verify update is processed and forwarded, response propagated to IPL", "Complete flow successful"),
                ("10", "Verify all updates maintain backward compatibility with existing gateway implementation", "Existing functionality remains unaffected")
            ]
        },
        {
            "tid": "TC-013",
            "name": "[US-ALAMOS-GW-002] Verify Gateway handles Mission Update errors from ALAMOS",
            "category": "Negative Scenario",
            "description": "Validate ALAMOS Gateway handles business and technical errors from ALAMOS for Continuous Mission Updates.",
            "preconditions": "1. ALAMOS Gateway is running\n2. ALAMOS can return error responses",
            "priority": "High",
            "user_story": "US-ALAMOS-GW-002",
            "steps": [
                ("1", "Configure ALAMOS to return business error, trigger Mission Update from IPL", "Update is sent to ALAMOS"),
                ("2", "Verify business error is propagated back to IPL", "IPL receives business error"),
                ("3", "Configure ALAMOS to return technical error, trigger Mission Update from IPL", "Update is sent to ALAMOS"),
                ("4", "Verify technical error is propagated back to IPL", "IPL receives technical error")
            ]
        },
        
        # ============================================================================
        # ALAMOS GATEWAY - Squad Information Extension (US-ALAMOS-GW-003)
        # Consolidated from original TC-030 to TC-033
        # ============================================================================
        {
            "tid": "TC-014",
            "name": "[US-ALAMOS-GW-003] Verify Gateway processes squad information (single/multiple squads) with backward compatibility",
            "category": "Functional (Happy Path)",
            "description": "Validate ALAMOS Gateway extracts and maps squad information from Alerting requests (single and multiple squads), maintains backward compatibility for requests without squads, and transmits correctly to ALAMOS FE2.",
            "preconditions": "1. ALAMOS Gateway is running\n2. ALAMOS FE2 is available",
            "priority": "High",
            "user_story": "US-ALAMOS-GW-003",
            "steps": [
                ("1", "Send Alerting request with single squad in <squads> element", "Gateway receives the request"),
                ("2", "Verify squad information is extracted and mapped to ALAMOS message structure", "Single squad is correctly mapped"),
                ("3", "Verify request is forwarded to ALAMOS FE2 with squad data", "Squad information is transmitted correctly"),
                ("4", "Send Alerting request with multiple squads in <squads> element", "Gateway receives the request"),
                ("5", "Verify all squad entries are correctly mapped to ALAMOS message structure", "Multiple squads are correctly mapped"),
                ("6", "Verify all squads are transmitted to ALAMOS FE2", "All squad data received by ALAMOS FE2"),
                ("7", "Send Alerting request without <squads> element", "Gateway receives the request"),
                ("8", "Verify request is processed successfully without errors (backward compatibility)", "Request processed, response propagated to IPL"),
                ("9", "Verify acknowledgments are propagated back to IPL without loss of information", "IPL receives complete acknowledgment")
            ]
        },
        
        # ============================================================================
        # RELIABILITY - Retry & DLQ Validation (US-ADAPTER-RELIABILITY-001)
        # Consolidated from original TC-034 to TC-038
        # ============================================================================
        {
            "tid": "TC-015",
            "name": "[US-ADAPTER-RELIABILITY-001] Verify retry, DLQ, and logging behavior across all ALAMOS flows",
            "category": "Functional (Happy Path)",
            "description": "Validate existing retry mechanisms for recoverable failures, DLQ routing when limits exceeded, business validation errors do not trigger retries, and diagnostic logging across Situation Report, Alerting, and Mission Update flows.",
            "preconditions": "1. Adapter/Gateway running\n2. Retry policies and DLQ configured\n3. Logging enabled",
            "priority": "High",
            "user_story": "US-ADAPTER-RELIABILITY-001",
            "steps": [
                ("1", "Simulate a temporary communication failure during Situation Report processing", "Initial attempt fails"),
                ("2", "Verify message is retried according to existing retry policy", "Retry is triggered"),
                ("3", "Resolve the transient failure before retry limit", "Next retry succeeds"),
                ("4", "Verify message is processed successfully after retry", "Success response propagated"),
                ("5", "Simulate persistent failure for Alerting flow exceeding retry limit", "All retry attempts fail"),
                ("6", "Verify message is routed to existing DLQ", "Message is in DLQ"),
                ("7", "Simulate a business validation error from IPL", "Business error is returned"),
                ("8", "Verify business validation error does NOT trigger unnecessary retries", "No retries triggered for business error"),
                ("9", "Verify retry attempts, DLQ events, and failures are logged with sufficient diagnostic information", "Logs contain correlation IDs, timestamps, error details"),
                ("10", "Verify success and fault responses are propagated correctly to originating caller", "Responses accurately reflect processing result")
            ]
        },
        
        # ============================================================================
        # ALAMOS GATEWAY - Mock Service (US-ALAMOSGW-MOCK-001)
        # Consolidated from original TC-039 to TC-043
        # ============================================================================
        {
            "tid": "TC-016",
            "name": "[US-ALAMOSGW-MOCK-001] Verify Gateway mock mode functionality for development/testing",
            "category": "Functional (Happy Path)",
            "description": "Validate ALAMOS Gateway runs in mock mode, returns configurable success/error responses for Alerting and Mission Updates, and does not impact normal operation when disabled.",
            "preconditions": "1. Gateway mock mode configuration available",
            "priority": "High",
            "user_story": "US-ALAMOSGW-MOCK-001",
            "steps": [
                ("1", "Enable mock mode configuration and start ALAMOS Gateway", "Gateway starts successfully in mock mode"),
                ("2", "Send a valid Alerting request with squad information", "Gateway processes in mock mode"),
                ("3", "Verify predefined success response is returned without connecting to ALAMOS FE2", "Mock success response returned"),
                ("4", "Configure mock mode to return an error response", "Error configuration applied"),
                ("5", "Send another Alerting request", "Gateway returns configured error response"),
                ("6", "Send Mission Update request in mock mode", "Gateway processes in mock mode"),
                ("7", "Verify mock response is returned for Mission Update", "Mock response returned"),
                ("8", "Disable mock mode and restart Gateway in normal mode", "Gateway starts normally"),
                ("9", "Verify mock mode configuration has no impact on normal gateway operation", "Gateway operates with real ALAMOS FE2 connectivity")
            ]
        },
        
        # ============================================================================
        # END-TO-END INTEGRATION TESTING (US-INTEGRATION-E2E-001)
        # Consolidated from original TC-044 to TC-050
        # Only keeping truly E2E scenarios that aren't covered by component tests
        # ============================================================================
        {
            "tid": "TC-017",
            "name": "[US-INTEGRATION-E2E-001] Verify complete Alerting with Squad Information flow (E2E)",
            "category": "Integration (E2E)",
            "description": "Validate complete end-to-end Alerting flow with squad information: IPL → Facade Service (LEGACY) → ALAMOS Gateway → ALAMOS FE2, verifying squad transmission and acknowledgment propagation across all components.",
            "preconditions": "1. All components running (IPL, Facade Service, ALAMOS Gateway, ALAMOS FE2)\n2. Network connectivity established",
            "priority": "High",
            "user_story": "US-INTEGRATION-E2E-001",
            "steps": [
                ("1", "Send SOAP AlarmRequest with squads from IPL", "Request enters the system"),
                ("2", "Verify request flows through: IPL → Facade Service → ALAMOS Gateway → ALAMOS FE2", "Request traverses all components"),
                ("3", "Verify squad information is correctly transmitted to ALAMOS FE2", "Squad data intact at destination"),
                ("4", "Verify acknowledgment flows back: ALAMOS FE2 → ALAMOS Gateway → Facade Service → IPL", "Acknowledgment propagated correctly"),
                ("5", "Verify no data loss or transformation errors across component boundaries", "Data integrity maintained")
            ]
        },
        {
            "tid": "TC-018",
            "name": "[US-INTEGRATION-E2E-001] Verify complete Situation Report flow (E2E)",
            "category": "Integration (E2E)",
            "description": "Validate complete end-to-end Situation Report flow: ALAMOS FE2 → Adapter → Domain Event → FacadeService → IPL, verifying event creation, transformation, and acknowledgment across all components.",
            "preconditions": "1. All components running (ALAMOS FE2, Adapter, FacadeService, IPL)\n2. Network connectivity established",
            "priority": "High",
            "user_story": "US-INTEGRATION-E2E-001",
            "steps": [
                ("1", "Send Situation Report from ALAMOS FE2 to Adapter", "Report enters the Adapter"),
                ("2", "Verify incident.situation_report_received Domain Event is created with all required attributes", "Domain Event created correctly"),
                ("3", "Verify FacadeService routes and maps event to IPL request", "Request transformation successful"),
                ("4", "Verify IPL erzeugeEinsatzmeldung is invoked and returns success", "IPL processes successfully"),
                ("5", "Verify acknowledgment is propagated back to ALAMOS FE2", "ALAMOS FE2 receives success acknowledgment"),
                ("6", "Verify logging and traceability across all components", "Correlation IDs and logs available")
            ]
        },
        {
            "tid": "TC-019",
            "name": "[US-INTEGRATION-E2E-001] Verify E2E failure propagation and retry/DLQ behavior",
            "category": "Integration (E2E / Negative)",
            "description": "Validate end-to-end failure propagation and retry/DLQ mechanisms: failures propagate correctly to originating system, retries work across component boundaries, and DLQ captures failed messages without loss.",
            "preconditions": "1. All components running\n2. Retry and DLQ mechanisms configured",
            "priority": "High",
            "user_story": "US-INTEGRATION-E2E-001",
            "steps": [
                ("1", "Trigger IPL business validation failure during Situation Report processing", "Business failure occurs"),
                ("2", "Verify failure is propagated correctly back to ALAMOS FE2", "ALAMOS FE2 receives accurate failure response"),
                ("3", "Simulate temporary ALAMOS FE2 unavailability during Alerting flow", "Communication failure triggers retry"),
                ("4", "Verify retry mechanism works end-to-end and succeeds when FE2 recovers", "Alerting completes after retry"),
                ("5", "Simulate persistent failure exceeding retry limits", "All retries fail"),
                ("6", "Verify message is routed to DLQ with no unexpected message loss", "DLQ contains failed message"),
                ("7", "Verify logging and monitoring data is available across all E2E flows", "Complete diagnostic trail available")
            ]
        },
    ]
    
    # Write test cases to worksheet
    row_num = 2
    for tc in test_cases:
        # Write first row with test case header info
        first_step = tc["steps"][0]
        ws.cell(row=row_num, column=1, value=tc["tid"])
        ws.cell(row=row_num, column=2, value=tc["name"])
        ws.cell(row=row_num, column=3, value=tc["category"])
        ws.cell(row=row_num, column=4, value=tc["description"])
        ws.cell(row=row_num, column=5, value=tc["preconditions"])
        ws.cell(row=row_num, column=6, value=first_step[0])
        ws.cell(row=row_num, column=7, value=first_step[1])
        ws.cell(row=row_num, column=8, value=first_step[2])
        ws.cell(row=row_num, column=9, value=tc["priority"])
        ws.cell(row=row_num, column=10, value=tc["user_story"])
        row_num += 1
        
        # Write remaining steps
        for step in tc["steps"][1:]:
            ws.cell(row=row_num, column=6, value=step[0])
            ws.cell(row=row_num, column=7, value=step[1])
            ws.cell(row=row_num, column=8, value=step[2])
            row_num += 1
        
        # Add empty row between test cases
        row_num += 1
    
    # Adjust column widths
    column_widths = [10, 80, 25, 60, 50, 10, 80, 60, 10, 25]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # Apply text wrapping to all cells
    for row in ws.iter_rows(min_row=1, max_row=row_num):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"ALAMOS_Functional_Integration_TestCases_Optimized_{timestamp}.xlsx"
    
    # Save workbook
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    
    print(f"Optimized test cases generated successfully!")
    print(f"File saved to: {output_path}")
    print(f"\nSummary:")
    print(f"  Original test cases: 50")
    print(f"  Optimized test cases: {len(test_cases)}")
    print(f"  Reduction: {50 - len(test_cases)} test cases ({((50 - len(test_cases)) / 50 * 100):.1f}% reduction)")
    print(f"\nCoverage maintained for all user stories:")
    user_stories = set(tc["user_story"] for tc in test_cases)
    for us in sorted(user_stories):
        count = sum(1 for tc in test_cases if tc["user_story"] == us)
        print(f"  {us}: {count} test case(s)")
    
    return output_path

if __name__ == "__main__":
    create_optimized_test_cases_workbook()
