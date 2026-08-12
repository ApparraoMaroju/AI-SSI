import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "UCRI Test Cases"

# Template columns: Test ID | Description | Pre-conditions | Steps | Expected Result | Priority
headers = ["Test ID", "Description", "Pre-conditions", "Steps", "Expected Result", "Priority"]
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
wrap = Alignment(wrap_text=True, vertical='top')

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.border = thin_border
    c.alignment = wrap

ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 50
ws.column_dimensions['C'].width = 50
ws.column_dimensions['D'].width = 70
ws.column_dimensions['E'].width = 55
ws.column_dimensions['F'].width = 10

row = 2

def tc(tid, desc, pre, steps, expected, priority):
    global row
    ws.cell(row=row, column=1, value=tid).alignment = wrap
    ws.cell(row=row, column=2, value=desc).alignment = wrap
    ws.cell(row=row, column=3, value=pre).alignment = wrap
    ws.cell(row=row, column=4, value=steps).alignment = wrap
    ws.cell(row=row, column=5, value=expected).alignment = wrap
    ws.cell(row=row, column=6, value=priority).alignment = wrap
    for c in range(1, 7):
        ws.cell(row=row, column=c).border = thin_border
    row += 1

# ============================================================
# SECTION 1: IN-MEMORY TRANSPORT TEST CASES
# ============================================================

tc("IMT-001",
   "Verify Adapter starts correctly with in-memory transport mode (Vert.x EventBus)",
   "Adapter service deployed.\nadapter.transport=in-memory configured.\nadapter.receiver.enabled=true\nadapter.handler.enabled=true\nadapter.dispatcher.enabled=true",
   "1. Set configuration property adapter.transport=in-memory\n2. Start the Adapter service (Receiver, Handler, Dispatcher)\n3. Check startup logs for transport mode indication\n4. Verify no external Kafka broker connection is attempted\n5. Verify Receiver, Handler, and Dispatcher modules are all initialized",
   "Adapter starts successfully.\nLogs indicate 'in-memory transport active' (Vert.x EventBus).\nNo Kafka connection attempts.\nAll three modules report 'started/initialized'.",
   "High")

tc("IMT-002",
   "Verify end-to-end message flow from Receiver → Handler → Dispatcher using in-memory transport",
   "Adapter running with in-memory transport.\nAll modules enabled.\ndispatch-center.ucri.enabled=true",
   "1. Send a valid event message to the Receiver via EventReceiverPort\n2. Verify Receiver validates the incoming event\n3. Verify Handler receives the event from Receiver via in-memory EventBus\n4. Verify Handler processes and transforms the event\n5. Verify Handler forwards the processed event to Dispatcher\n6. Verify Dispatcher routes the event to correct dispatch center topic (events-ucri)\n7. Verify envelope metadata (correlationId, headers) is preserved end-to-end",
   "Event flows: Receiver → Handler → Dispatcher without errors.\nDispatcher publishes to events-ucri topic.\ncorrelationId and headers match original event across all three services.",
   "High")

tc("IMT-003",
   "Verify event ordering is preserved for events with the same partition key via in-memory transport",
   "Adapter running with in-memory transport.\nAll modules enabled.",
   "1. Publish Event-A with partitionKey='PK-001' at time T1\n2. Publish Event-B with partitionKey='PK-001' at time T2 (T2 > T1)\n3. Publish Event-C with partitionKey='PK-001' at time T3 (T3 > T2)\n4. Verify Dispatcher outputs events in order: A, B, C\n5. Verify ordering is preserved based on partition key",
   "Events delivered in published order (A → B → C) at Dispatcher output.\nFIFO order maintained for same partitionKey.",
   "High")

tc("IMT-004",
   "Verify envelope metadata (correlationId, headers, partitionKey) preserved across Receiver → Handler → Dispatcher",
   "Adapter running with in-memory transport.\nAll modules enabled.",
   "1. Create event with correlationId='CORR-12345', headers={'source':'IPL','type':'CreateRequest'}, partitionKey='PK-001'\n2. Send event to Receiver\n3. Verify Handler receives event with same correlationId\n4. Verify Dispatcher receives event with same correlationId and headers\n5. Verify outbound event contains all original metadata",
   "All metadata (correlationId, headers, partitionKey) matches the original event at every service boundary.\nNo metadata lost or altered.",
   "High")

tc("IMT-005",
   "Verify invalid/malformed event is rejected at Receiver and not forwarded to Handler or Dispatcher",
   "Adapter running with in-memory transport.\nAll modules enabled.",
   "1. Send a malformed event (missing required fields) to Receiver\n2. Check Receiver logs for validation failure reason\n3. Verify Handler does NOT receive the rejected event\n4. Verify Dispatcher does NOT receive or publish the rejected event",
   "Receiver rejects the event with a clear validation error.\nNo Handler or Dispatcher processing for the rejected event.",
   "High")

tc("IMT-006",
   "Verify disabling Receiver module prevents event ingestion (adapter.receiver.enabled=false)",
   "Adapter deployed.\nadapter.receiver.enabled=false\nadapter.handler.enabled=true\nadapter.dispatcher.enabled=true",
   "1. Start Adapter with adapter.receiver.enabled=false\n2. Verify Receiver module is NOT instantiated (logs confirm disabled)\n3. Attempt to send an event to the Receiver endpoint\n4. Verify no resources consumed by Receiver (no threads or subscriptions)",
   "Adapter starts without Receiver.\nNo event subscriptions created.\nEvents are not consumed.",
   "Medium")

tc("IMT-007",
   "Verify disabling Handler module stops event processing (adapter.handler.enabled=false)",
   "Adapter deployed.\nadapter.receiver.enabled=true\nadapter.handler.enabled=false\nadapter.dispatcher.enabled=true",
   "1. Start Adapter with adapter.handler.enabled=false\n2. Send a valid event to the Receiver\n3. Verify Handler does NOT process the event\n4. Verify Dispatcher does NOT receive or dispatch the event",
   "Receiver accepts event but Handler does not process it.\nEvent stops after Receiver.\nNo Dispatcher activity.",
   "Medium")

tc("IMT-008",
   "Verify disabling Dispatcher module prevents event routing (adapter.dispatcher.enabled=false)",
   "Adapter deployed.\nadapter.receiver.enabled=true\nadapter.handler.enabled=true\nadapter.dispatcher.enabled=false",
   "1. Start Adapter with adapter.dispatcher.enabled=false\n2. Send a valid event to Receiver\n3. Verify Handler processes the event\n4. Verify Dispatcher does NOT route the event to any dispatch center",
   "Receiver and Handler work normally.\nNo outbound event published.\nDispatcher module inactive.",
   "Medium")

tc("IMT-009",
   "Verify invalid transport configuration causes startup failure with descriptive error",
   "Adapter deployed.\nadapter.transport=invalid_value",
   "1. Set adapter.transport=invalid_value in configuration\n2. Attempt to start the Adapter service\n3. Check error output/logs for descriptive error message\n4. Correct to adapter.transport=in-memory and restart",
   "Startup fails.\nError message lists valid transport options (in-memory, kafka).\nAfter correction, Adapter starts successfully.",
   "Medium")

tc("IMT-010",
   "Verify missing transport configuration property causes startup failure",
   "Adapter deployed.\nadapter.transport property removed from configuration.",
   "1. Remove adapter.transport property from configuration\n2. Attempt to start the Adapter service\n3. Verify error message indicates missing transport configuration",
   "Startup fails with descriptive error clearly stating transport property is required.",
   "Medium")

tc("IMT-011",
   "Verify retry decorator retries failed publish operations with configurable backoff",
   "Adapter running with in-memory transport.\ndispatch-center.ucri.retry-attempts=3\ndispatch-center.ucri.retry-delay-ms=1000",
   "1. Simulate a transient Dispatcher publish failure for the first attempt\n2. Verify first retry occurs after ~1000ms delay\n3. Verify second retry if first retry also fails\n4. Verify third and final retry if second also fails\n5. Verify metrics reflect all retry attempts",
   "Retry occurs per configuration (3 attempts, ~1s delay).\nRetry timestamps logged.\nPublish failure count and retry count metrics updated.",
   "Medium")

tc("IMT-012",
   "Verify events failing permanently after all retries are sent to Dead Letter Queue (DLQ)",
   "Adapter running with in-memory transport.\nDLQ configured: dispatch-center.ucri.dlq-topic=events-ucri-dlq\nretry-attempts=3",
   "1. Send event that causes permanent publish failure after all retries\n2. Verify event is sent to DLQ topic (events-ucri-dlq)\n3. Verify DLQ entry contains the original event payload\n4. Verify DLQ entry contains error context (reason for failure)",
   "Event present in DLQ topic after 3 failed retries.\nOriginal payload and error context attached to DLQ entry.",
   "Medium")

tc("IMT-013",
   "Verify comprehensive logging with correlationId traceability across Receiver, Handler, and Dispatcher",
   "Adapter running with in-memory transport.\nLogging enabled.",
   "1. Send a valid event with correlationId='TRACE-001' to Receiver\n2. Verify Receiver logs event receipt and validation with correlationId='TRACE-001'\n3. Verify Handler logs transformation with same correlationId\n4. Verify Dispatcher logs dispatch with same correlationId\n5. Verify end-to-end traceability: all log entries linkable via correlationId",
   "All three services' logs contain correlationId='TRACE-001'.\nEnd-to-end traceability confirmed.",
   "High")

# ============================================================
# SECTION 2: IPL TO UCRI (FORWARD FLOW) TEST CASES
# ============================================================

tc("FWD-001",
   "IPL to UCRI — Happy Path: Event sent and accepted by UCRM (EMPFANG ACK)",
   "Adapter running (in-memory transport).\nIPL operational (or mock).\nUCRM accessible.\nP2P token valid.\ndispatch-center.ucri.enabled=true",
   "1. IPL sends valid TransferEinsatzRequest to Adapter via POST /einsatztransfer/erzeugeEinsatz\n2. Verify Receiver validates the incoming request (format, required fields)\n3. Verify Handler authenticates to UCRM via P2P token\n4. Verify Handler transforms IPL payload to UCRM format (sourceMandant → header.sender_client_id, targetMandant → header.receiver_client_id, einsatzId → operation.id, etc.)\n5. Verify Dispatcher forwards message to UCRM via POST messaging/send\n6. UCRM processes and returns Accept response\n7. Verify Handler translates Accept response to IPL ACK format\n8. Verify Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Accept status\n9. Verify log: Einsatztransfer/quittiereTranferEinsatz contains 'EMPFANG'",
   "End-to-end forward flow completes.\nIPL receives EMPFANG ACK.\nLog entry shows EMPFANG for this transaction.",
   "High")

tc("FWD-002",
   "IPL to UCRI — Event rejected by UCRM (EINSATZ_ABGELEHNT ACK)",
   "Adapter running.\nIPL operational.\nUCRM accessible.\nP2P token valid.",
   "1. IPL sends valid TransferEinsatzRequest via POST /einsatztransfer/erzeugeEinsatz\n2. Receiver validates (passes)\n3. Handler authenticates and transforms payload\n4. Dispatcher forwards to UCRM via messaging/send\n5. UCRM rejects and returns Reject response\n6. Handler translates Reject to IPL ACK format\n7. Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Reject status\n8. Verify log: Einsatztransfer/quittiereTranferEinsatz contains 'EINSATZ_ABGELEHNT'",
   "IPL receives EINSATZ_ABGELEHNT ACK.\nLog entry shows EINSATZ_ABGELEHNT for this transaction.",
   "High")

tc("FWD-003",
   "IPL to UCRI — UCRM does not respond within allowed time (TECHNISCHER_FEHLER ACK)",
   "Adapter running.\nIPL operational.\nUCRM accessible but configured to not respond (or delayed).\ntimeout-ms configured.",
   "1. IPL sends valid TransferEinsatzRequest via POST /einsatztransfer/erzeugeEinsatz\n2. Receiver validates and Handler authenticates/transforms\n3. Dispatcher forwards to UCRM via messaging/send\n4. UCRM does NOT respond within configured timeout (e.g., 10 seconds)\n5. Handler detects timeout and prepares TECHNISCHER_FEHLER ACK\n6. Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Timeout error\n7. Verify log: Einsatztransfer/quittiereTranferEinsatz contains 'TECHNISCHER_FEHLER'",
   "IPL receives TECHNISCHER_FEHLER ACK.\nLog entry shows TECHNISCHER_FEHLER for this transaction.\nTimeout detected within configured time limit.",
   "High")

tc("FWD-004",
   "IPL to UCRI — Validation failure at Receiver (malformed TransferEinsatzRequest)",
   "Adapter running.\nIPL operational.",
   "1. IPL sends invalid/malformed TransferEinsatzRequest (e.g., missing einsatzId) via POST /einsatztransfer/erzeugeEinsatz\n2. Verify Receiver logs validation failure reason\n3. Verify IPL receives validation error response\n4. Verify NO message forwarded to Handler\n5. Verify NO message sent to UCRM via Dispatcher\n6. Verify NO Einsatztransfer/quittiereTranferEinsatz log entry generated",
   "Receiver rejects request with validation error.\nNo forwarding to Handler or UCRM.\nNo ACK log entry generated.",
   "High")

tc("FWD-005",
   "IPL to UCRI — P2P authentication failure at Handler (invalid/expired token)",
   "Adapter running.\nIPL operational.\nP2P token invalid/expired.",
   "1. IPL sends valid TransferEinsatzRequest via POST /einsatztransfer/erzeugeEinsatz\n2. Receiver validates (passes), forwards to Handler\n3. Handler attempts P2P authentication with invalid/expired token\n4. Verify authentication fails at UCRM\n5. Handler logs authentication failure details\n6. Error propagated back to IPL\n7. Verify NO message forwarded to UCRM via Dispatcher",
   "Authentication fails.\nIPL receives auth error response.\nNo message sent to UCRM.\nFailure details logged.",
   "High")

tc("FWD-006",
   "IPL to UCRI — Field mapping validation: TransferEinsatz to UCRM operation format",
   "Adapter running.\nIPL operational.\nUCRM accessible.",
   "1. Send TransferEinsatzRequest with sourceMandant='IPL_CENTER_001'\n2. Verify Handler maps to header.sender_client_id='IPL_CENTER_001'\n3. Verify targetMandant='UCRI_CENTER_001' → header.receiver_client_id='UCRI_CENTER_001'\n4. Verify einsatzId='EIN-12345' → operation.id='EIN-12345'\n5. Verify einsatznummerKurz='K-98765' → operation.name='K-98765'\n6. Verify kategorie='KAT-01' → operation.occassion.type_1='KAT-01'\n7. Verify einsatzStichwort='FEUER' → operation.occassion.keyword_1='FEUER'\n8. Verify alarmierungsstichwortText='Brand in Wohngebäude' → operation.occassion.keyword_1_text\n9. Verify no data loss in field translation",
   "All field mappings applied correctly (1:1 mapping per specification).\nNo data truncation or corruption.\nAll mapped fields present in outbound UCRM message.",
   "High")

tc("FWD-007",
   "IPL to UCRI — Missing required field: sourceMandant",
   "Adapter running.\nIPL operational.",
   "1. Send TransferEinsatzRequest without sourceMandant field\n2. Verify Receiver rejects the request with validation error\n3. Verify error response indicates 'sourceMandant is required'\n4. Verify no forwarding to Handler or UCRM",
   "Request rejected with clear validation error indicating missing sourceMandant.\nNo downstream processing.",
   "High")

tc("FWD-008",
   "IPL to UCRI — Missing required field: targetMandant",
   "Adapter running.\nIPL operational.",
   "1. Send TransferEinsatzRequest without targetMandant field\n2. Verify Receiver rejects the request with validation error\n3. Verify error response indicates 'targetMandant is required'\n4. Verify no forwarding to Handler or UCRM",
   "Request rejected with clear validation error indicating missing targetMandant.\nNo downstream processing.",
   "High")

tc("FWD-009",
   "IPL to UCRI — Missing required field: einsatzId",
   "Adapter running.\nIPL operational.",
   "1. Send TransferEinsatzRequest without einsatzId field\n2. Verify Receiver rejects the request with validation error\n3. Verify error response indicates 'einsatzId is required'\n4. Verify no forwarding to Handler or UCRM",
   "Request rejected with clear validation error indicating missing einsatzId.\nNo downstream processing.",
   "Medium")

tc("FWD-010",
   "IPL to UCRI — Empty payload sent to /einsatztransfer/erzeugeEinsatz",
   "Adapter running.\nIPL operational.",
   "1. Send POST /einsatztransfer/erzeugeEinsatz with empty body\n2. Verify Receiver rejects the request with validation error\n3. Verify error describes empty/invalid payload\n4. Verify no forwarding to Handler or UCRM",
   "Request rejected with clear error describing empty/invalid payload.\nNo downstream processing.",
   "Medium")

tc("FWD-011",
   "IPL to UCRI — UCRI dispatch center disabled (dispatch-center.ucri.enabled=false)",
   "Adapter running.\ndispatch-center.ucri.enabled=false",
   "1. Set dispatch-center.ucri.enabled=false\n2. Send TransferEinsatzRequest with targetMandant pointing to UCRI\n3. Verify Dispatcher does NOT route event to events-ucri topic\n4. Verify appropriate error or handling for disabled dispatch center",
   "Event not dispatched to UCRI.\nAppropriate handling for disabled target center.",
   "Medium")

tc("FWD-012",
   "IPL to UCRI — CorrelationId traceability across forward flow",
   "Adapter running.\nIPL operational.\nUCRM accessible.\nLogging enabled.",
   "1. IPL sends TransferEinsatzRequest with unique correlationId\n2. Verify Receiver logs receipt with the correlationId\n3. Verify Handler logs processing with same correlationId\n4. Verify Dispatcher logs forwarding with same correlationId\n5. Verify ACK response log (quittiereTranferEinsatz) contains same correlationId",
   "All log entries across Receiver, Handler, Dispatcher, and ACK contain matching correlationId.\nEnd-to-end traceability confirmed.",
   "High")

tc("FWD-013",
   "IPL to UCRI — Multiple events with same partition key delivered in order",
   "Adapter running.\nIPL operational.\nUCRM accessible.",
   "1. IPL sends Event-1 with partitionKey='INCIDENT-001' at T1\n2. IPL sends Event-2 with partitionKey='INCIDENT-001' at T2 (T2 > T1)\n3. IPL sends Event-3 with partitionKey='INCIDENT-001' at T3 (T3 > T2)\n4. Verify UCRM receives events in order: Event-1, Event-2, Event-3",
   "Events arrive at UCRM in FIFO order per partition key.",
   "Medium")

tc("FWD-014",
   "IPL to UCRI — P2P token expires during active session",
   "Adapter running.\nIPL operational.\nUCRM accessible.\nP2P token set to expire during test.",
   "1. IPL sends first TransferEinsatzRequest — P2P token is valid\n2. First event processed and delivered to UCRM successfully\n3. Wait for P2P token to expire\n4. IPL sends second TransferEinsatzRequest — token now expired\n5. Verify Handler detects expired token\n6. Verify error propagated to IPL; no message forwarded to UCRM\n7. Verify error logged with details",
   "First event succeeds.\nSecond event fails with P2P token expiry error.\nIPL notified; no message sent to UCRM.\nExpiry details logged.",
   "Medium")

tc("FWD-015",
   "IPL to UCRI — Verify EMPFANG ACK log entry content and correlationId match",
   "Adapter running.\nIPL operational.\nUCRM accessible.\nLogging enabled.",
   "1. Complete forward flow with UCRM Accept\n2. Check adapter logs for Einsatztransfer/quittiereTranferEinsatz\n3. Verify log contains 'EMPFANG'\n4. Verify correlationId in log matches original request correlationId",
   "Log entry shows EMPFANG with matching correlationId.",
   "High")

tc("FWD-016",
   "IPL to UCRI — Verify EINSATZ_ABGELEHNT ACK log entry content and correlationId match",
   "Adapter running.\nIPL operational.\nUCRM accessible.\nLogging enabled.",
   "1. Complete forward flow with UCRM Reject\n2. Check adapter logs for Einsatztransfer/quittiereTranferEinsatz\n3. Verify log contains 'EINSATZ_ABGELEHNT'\n4. Verify correlationId matches original request",
   "Log entry shows EINSATZ_ABGELEHNT with matching correlationId.",
   "High")

tc("FWD-017",
   "IPL to UCRI — Verify TECHNISCHER_FEHLER ACK log entry content and correlationId match",
   "Adapter running.\nIPL operational.\nUCRM accessible but timing out.\nLogging enabled.",
   "1. Complete forward flow where UCRM times out\n2. Check adapter logs for Einsatztransfer/quittiereTranferEinsatz\n3. Verify log contains 'TECHNISCHER_FEHLER'\n4. Verify correlationId matches original request",
   "Log entry shows TECHNISCHER_FEHLER with matching correlationId.",
   "High")

# ============================================================
# SECTION 3: UCRI TO IPL (REVERSE FLOW) TEST CASES
# ============================================================

tc("REV-001",
   "UCRI to IPL — Happy Path: Mission accepted by IPL (Accept ACK propagated to UCRM)",
   "Adapter running (in-memory transport).\nUCRM operational (or test client).\nIPL accessible.\nP2P connection established.",
   "1. UCRM sends mission payload to Adapter via POST messaging/send\n2. Verify Receiver validates the incoming mission data\n3. Verify Handler translates UCRM mission format to IPL format\n4. Verify Dispatcher forwards the translated mission to IPL\n5. IPL processes mission successfully\n6. IPL sends Accept ACK back to Adapter\n7. Handler translates IPL ACK to UCRM ACK format\n8. Dispatcher propagates Accept ACK to UCRM",
   "Mission delivered to IPL and accepted.\nAccept ACK propagated back to UCRM.\nEnd-to-end reverse flow completes.",
   "High")

tc("REV-002",
   "UCRI to IPL — Mission rejected by IPL (Reject ACK propagated to UCRM)",
   "Adapter running.\nUCRM operational.\nIPL accessible.",
   "1. UCRM sends mission payload to Adapter via messaging/send\n2. Receiver validates and Handler translates mission to IPL format\n3. Dispatcher forwards mission to IPL\n4. IPL rejects the mission\n5. IPL sends Reject ACK to Adapter\n6. Handler translates Reject ACK to UCRM format\n7. Dispatcher propagates Reject ACK to UCRM",
   "IPL rejects mission.\nReject ACK propagated to UCRM.",
   "High")

tc("REV-003",
   "UCRI to IPL — IPL does not respond within allowed time (timeout handling)",
   "Adapter running.\nUCRM operational.\nIPL configured to not respond or delayed.",
   "1. UCRM sends mission payload to Adapter via messaging/send\n2. Adapter forwards mission to IPL\n3. IPL does NOT respond within configured timeout\n4. Verify Adapter handles the timeout appropriately\n5. Verify UCRM is notified of the timeout/failure",
   "Timeout detected by Adapter.\nTimeout error logged.\nUCRM receives timeout notification or error response.",
   "High")

tc("REV-004",
   "UCRI to IPL — Invalid/malformed mission payload from UCRM rejected at Receiver",
   "Adapter running.\nUCRM operational.",
   "1. UCRM sends malformed/incomplete mission payload via messaging/send\n2. Verify Receiver rejects the payload with validation error\n3. Verify validation error logged at Receiver\n4. Verify NO mission forwarded to IPL\n5. Verify error response sent back to UCRM",
   "Malformed payload rejected at Receiver.\nNo forwarding to IPL.\nUCRM receives rejection/error notification.",
   "High")

tc("REV-005",
   "UCRI to IPL — Field mapping validation: UCRM mission to IPL format",
   "Adapter running.\nUCRM operational.\nIPL accessible.",
   "1. UCRM sends mission with header.sender_client_id='UCRI_CENTER_001'\n2. Verify Handler maps to IPL sourceMandant='UCRI_CENTER_001'\n3. Verify header.receiver_client_id='IPL_CENTER_001' → targetMandant='IPL_CENTER_001'\n4. Verify operation.id='OP-56789' → einsatzId='OP-56789'\n5. Verify operation.name='K-11111' → einsatznummerKurz='K-11111'\n6. Verify no data loss in reverse field translation",
   "All field mappings applied correctly in reverse direction.\nNo data loss or corruption.",
   "High")

tc("REV-006",
   "UCRI to IPL — CorrelationId traceability across reverse flow",
   "Adapter running.\nUCRM operational.\nIPL accessible.\nLogging enabled.",
   "1. UCRM sends mission with correlationId='REV-TRACE-001'\n2. Verify Receiver logs receipt with correlationId='REV-TRACE-001'\n3. Verify Handler logs translation with same correlationId\n4. Verify Dispatcher logs forwarding to IPL with same correlationId\n5. Verify IPL ACK response logged with same correlationId",
   "All log entries linked via 'REV-TRACE-001'.\nEnd-to-end reverse flow traceability confirmed.",
   "High")

tc("REV-007",
   "UCRI to IPL — Multiple missions from UCRM processed sequentially and independently",
   "Adapter running.\nUCRM operational.\nIPL accessible.",
   "1. UCRM sends Mission-A to Adapter\n2. UCRM sends Mission-B to Adapter\n3. Verify Mission-A is fully processed (forwarded to IPL, ACK returned to UCRM)\n4. Verify Mission-B is processed independently\n5. Verify each mission has distinct correlationId in logs",
   "Both missions processed independently.\nDistinct correlationIds in logs.\nNo data cross-contamination.",
   "Medium")

tc("REV-008",
   "UCRI to IPL — ACK from IPL propagated back to UCRM within acceptable time limits",
   "Adapter running.\nUCRM operational.\nIPL accessible.\nACK timeout configured.",
   "1. UCRM sends mission to Adapter\n2. Adapter forwards to IPL\n3. IPL sends Accept ACK back to Adapter\n4. Verify ACK propagated to UCRM within configured timeout (e.g., 10 seconds)\n5. Verify no timeout error triggered",
   "ACK reaches UCRM within acceptable time.\nNo TECHNISCHER_FEHLER logged.",
   "Medium")

# ============================================================
# SECTION 4: UCRI (CLIENT-1) TO UCRI (CLIENT-2) CROSS-CLIENT FLOW
# ============================================================

tc("U2U-001",
   "UCRI-1 to UCRI-2 — Happy Path: Mission dispatched and accepted by UCRI-2",
   "Adapter running.\nUCRI-1 (sender) and UCRI-2 (receiver) both accessible.\nP2P tokens valid for both clients.\ndispatch-center.ucri.enabled=true",
   "1. UCRI-1 (sourceMandant) sends mission/event to Adapter via messaging/send with targetMandant=UCRI-2\n2. Verify Receiver validates request (sourceMandant and targetMandant both UCRI clients)\n3. Verify Handler authenticates via P2P token for UCRI-2 target\n4. Verify Handler transforms UCRI-1 message to UCRI-2 compatible format\n5. Verify Dispatcher forwards message to UCRI-2 via messaging/send\n6. UCRI-2 processes and returns Accept response\n7. Verify Handler translates Accept to UCRI-1 ACK format\n8. Verify Dispatcher propagates Accept ACK back to UCRI-1",
   "Mission delivered to UCRI-2 and accepted.\nUCRI-1 receives Accept ACK confirming UCRI-2 acceptance.",
   "High")

tc("U2U-002",
   "UCRI-1 to UCRI-2 — Mission rejected by UCRI-2 (Reject ACK to UCRI-1)",
   "Adapter running.\nUCRI-1 and UCRI-2 accessible.\nP2P tokens valid.",
   "1. UCRI-1 sends mission/event to Adapter with targetMandant=UCRI-2\n2. Adapter validates, authenticates, transforms, and forwards to UCRI-2\n3. UCRI-2 rejects the mission and returns Reject response\n4. Handler translates Reject to UCRI-1 ACK format\n5. Dispatcher propagates Reject ACK back to UCRI-1",
   "UCRI-1 receives Reject ACK confirming UCRI-2 rejection.",
   "High")

tc("U2U-003",
   "UCRI-1 to UCRI-2 — Timeout: UCRI-2 does not respond within allowed time",
   "Adapter running.\nUCRI-1 accessible.\nUCRI-2 configured to not respond.\ntimeout-ms configured.",
   "1. UCRI-1 sends mission/event to Adapter with targetMandant=UCRI-2\n2. Adapter forwards to UCRI-2 via messaging/send\n3. UCRI-2 does NOT respond within configured timeout\n4. Verify Handler handles timeout condition\n5. Verify Dispatcher propagates timeout notification back to UCRI-1",
   "Timeout detected.\nUCRI-1 receives timeout error/notification.\nTimeout logged.",
   "High")

tc("U2U-004",
   "UCRI-2 to UCRI-1 — Reverse direction: Mission dispatched and accepted by UCRI-1",
   "Adapter running.\nUCRI-1 and UCRI-2 accessible.\nP2P tokens valid for both directions.",
   "1. UCRI-2 (sourceMandant) sends mission/event to Adapter with targetMandant=UCRI-1\n2. Adapter validates, authenticates, transforms, and forwards to UCRI-1\n3. UCRI-1 processes and returns Accept response\n4. Adapter propagates Accept ACK back to UCRI-2",
   "UCRI-2 receives Accept ACK confirming UCRI-1 acceptance.\nBi-directional cross-client capability confirmed.",
   "High")

tc("U2U-005",
   "UCRI-1 to UCRI-2 — Unrecognized sourceMandant (sender not valid UCRI client)",
   "Adapter running.\nUCRI-2 accessible.\nUCRI-1 using invalid/unrecognized sourceMandant.",
   "1. Send event with sourceMandant='UNKNOWN_CLIENT' and targetMandant='UCRI-2'\n2. Verify Receiver rejects with validation error indicating unrecognized sender\n3. Verify NO message forwarded to UCRI-2",
   "Request rejected with clear error about invalid sourceMandant.\nNo downstream processing.",
   "Medium")

tc("U2U-006",
   "UCRI-1 to UCRI-2 — Invalid targetMandant (receiver client not found)",
   "Adapter running.\nUCRI-1 accessible.\ntargetMandant points to non-existent UCRI client.",
   "1. Send event with sourceMandant='UCRI-1' and targetMandant='UCRI_NONEXISTENT'\n2. Verify Dispatcher cannot route to non-existent target\n3. Verify error response sent back to UCRI-1\n4. Verify error logged with routing failure details",
   "Routing failure or rejection at Dispatcher.\nUCRI-1 receives error notification.\nFailure details logged.",
   "Medium")

tc("U2U-007",
   "UCRI-1 to UCRI-2 — P2P authentication failure for cross-client connection",
   "Adapter running.\nUCRI-1 accessible.\nP2P token for UCRI-2 invalid/expired.",
   "1. UCRI-1 sends event with targetMandant=UCRI-2\n2. Receiver accepts the event\n3. Handler attempts P2P authentication for UCRI-2 with invalid token\n4. Verify authentication fails and error is logged\n5. Verify error notification propagated to UCRI-1\n6. Verify NO message delivered to UCRI-2",
   "P2P auth fails.\nUCRI-1 receives auth error.\nNo message sent to UCRI-2.",
   "High")

tc("U2U-008",
   "UCRI-1 to UCRI-2 — Field mapping validation for cross-client transfer",
   "Adapter running.\nUCRI-1 and UCRI-2 accessible.",
   "1. UCRI-1 sends mission with header.sender_client_id='UCRI_CENTER_001' (UCRI-1)\n2. Verify header.receiver_client_id='UCRI_CENTER_002' (UCRI-2) mapping\n3. Verify operation.id is preserved in cross-client transfer\n4. Verify operation.name is preserved\n5. Verify occasion fields (type_1, keyword_1, keyword_1_text) are preserved\n6. Verify no data loss in cross-client field translation",
   "All fields correctly mapped between UCRI-1 and UCRI-2.\nNo data loss or corruption.",
   "High")

tc("U2U-009",
   "UCRI-1 to UCRI-2 — CorrelationId traceability across cross-client flow",
   "Adapter running.\nUCRI-1 and UCRI-2 accessible.\nLogging enabled.",
   "1. UCRI-1 sends event with correlationId='U2U-TRACE-001'\n2. Verify Receiver logs receipt with correlationId='U2U-TRACE-001'\n3. Verify Handler logs processing with same correlationId\n4. Verify Dispatcher logs forwarding to UCRI-2 with same correlationId\n5. Verify UCRI-2 ACK response logged with same correlationId\n6. Verify end-to-end traceability: UCRI-1 → Adapter → UCRI-2 → Adapter → UCRI-1",
   "All log entries linked via 'U2U-TRACE-001'.\nFull round-trip traceability confirmed.",
   "High")

tc("U2U-010",
   "UCRI-1 to UCRI-2 — Simultaneous bi-directional transfers without interference",
   "Adapter running.\nUCRI-1 and UCRI-2 both operational.\nP2P tokens valid.",
   "1. UCRI-1 sends Mission-A to UCRI-2 (targetMandant=UCRI-2)\n2. Simultaneously, UCRI-2 sends Mission-B to UCRI-1 (targetMandant=UCRI-1)\n3. Verify Mission-A delivered to UCRI-2 correctly\n4. Verify Mission-B delivered to UCRI-1 correctly\n5. Verify UCRI-2 ACK for Mission-A propagated to UCRI-1\n6. Verify UCRI-1 ACK for Mission-B propagated to UCRI-2\n7. Verify no data cross-contamination between missions",
   "Both missions processed independently.\nACKs propagated correctly.\nNo interference or data cross-contamination.",
   "High")

tc("U2U-011",
   "UCRI-1 to UCRI-2 — Retry on transient Dispatcher failure for cross-client flow",
   "Adapter running.\nUCRI-1 accessible.\nUCRI-2 temporarily unavailable.\nretry-attempts=3; retry-delay-ms=1000",
   "1. UCRI-1 sends event with targetMandant=UCRI-2\n2. Dispatcher first attempt to UCRI-2 fails (transient error)\n3. Verify first retry after ~1000ms\n4. UCRI-2 becomes available; retry succeeds\n5. Verify UCRI-2 ACK propagated back to UCRI-1\n6. Verify retry count metrics updated",
   "Retry succeeds on subsequent attempt.\nUCRI-1 receives success ACK.\nMetrics reflect retry activity.",
   "Medium")

tc("U2U-012",
   "UCRI-1 to UCRI-2 — DLQ after all retries exhausted for cross-client flow",
   "Adapter running.\nUCRI-1 accessible.\nUCRI-2 permanently unavailable.\nretry-attempts=3.\nDLQ configured.",
   "1. UCRI-1 sends event with targetMandant=UCRI-2\n2. Dispatcher fails to deliver to UCRI-2 after all 3 retry attempts\n3. Verify event sent to DLQ topic (events-ucri-dlq)\n4. Verify DLQ entry contains original payload and failure reason\n5. Verify UCRI-1 notified of delivery failure",
   "Event in DLQ after 3 retries.\nDLQ entry has full context.\nUCRI-1 receives error/timeout notification.",
   "Medium")

# ============================================================
# SECTION 5: UCRM TEST CLIENT VALIDATION
# ============================================================

tc("TC-001",
   "UCRM Test Client — Installation, prerequisites, and basic connectivity",
   "UCRM Test Client installation package available.\nDeployment guide accessible.",
   "1. Verify all prerequisites listed in deployment guide are met\n2. Install UCRM Test Client per deployment guide\n3. Verify UCRM Test Client starts successfully\n4. Execute basic connectivity test to Adapter",
   "Installation completes without errors.\nTest Client running.\nConnectivity to Adapter confirmed.",
   "High")

tc("TC-002",
   "UCRM Test Client — Send mission and receive ACK through Adapter",
   "UCRM Test Client installed and running.\nAdapter running.\nIPL accessible (or mock).",
   "1. Use UCRM Test Client to send a mission to the Adapter\n2. Verify mission forwarded to IPL\n3. IPL sends Accept ACK\n4. Verify UCRM Test Client receives the Accept ACK",
   "Mission delivered to IPL.\nAccept ACK received by Test Client.",
   "High")

# Save
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
filename = f"c:\\Apparao Maroju\\AI-SSI\\Project-IGNIS(POC)\\UCRI2.0-Output\\TestCases_{timestamp}.xlsx"
wb.save(filename)
print(f"Test cases saved to: {filename}")
