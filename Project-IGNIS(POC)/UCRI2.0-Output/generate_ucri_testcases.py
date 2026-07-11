import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "UCRI Test Cases"

# Headers
headers = ["TID", "Test Case Name", "Category", "Description", "Pre-conditions", "Step No", "Step Description", "Expected Result", "Priority"]
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
wrap_alignment = Alignment(wrap_text=True, vertical='top')

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = wrap_alignment

# Column widths
col_widths = [12, 40, 30, 55, 55, 10, 60, 55, 10]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[chr(64 + i) if i <= 9 else None].width = w
ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 40
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 55
ws.column_dimensions['E'].width = 55
ws.column_dimensions['F'].width = 10
ws.column_dimensions['G'].width = 60
ws.column_dimensions['H'].width = 55
ws.column_dimensions['I'].width = 10

row = 2

def add_test_case(tid, name, category, description, preconditions, steps, priority):
    """Add a test case with steps. steps = list of (step_desc, expected_result)"""
    global row
    for i, (step_desc, expected) in enumerate(steps):
        if i == 0:
            ws.cell(row=row, column=1, value=tid).alignment = wrap_alignment
            ws.cell(row=row, column=2, value=name).alignment = wrap_alignment
            ws.cell(row=row, column=3, value=category).alignment = wrap_alignment
            ws.cell(row=row, column=4, value=description).alignment = wrap_alignment
            ws.cell(row=row, column=5, value=preconditions).alignment = wrap_alignment
        else:
            for c in range(1, 6):
                ws.cell(row=row, column=c, value="").alignment = wrap_alignment
        ws.cell(row=row, column=6, value=i + 1).alignment = wrap_alignment
        ws.cell(row=row, column=7, value=step_desc).alignment = wrap_alignment
        ws.cell(row=row, column=8, value=expected).alignment = wrap_alignment
        if i == 0:
            ws.cell(row=row, column=9, value=priority).alignment = wrap_alignment
        else:
            ws.cell(row=row, column=9, value="").alignment = wrap_alignment
        for c in range(1, 10):
            ws.cell(row=row, column=c).border = thin_border
        row += 1
    # blank row between test cases
    row += 1

# ============================================================
# SECTION 1: IN-MEMORY TRANSPORT TEST CASES
# ============================================================

add_test_case(
    "IMT-001", "In-Memory Transport Initialization", "In-Memory Transport",
    "Verify Adapter starts correctly with in-memory transport mode configured (Vert.x EventBus)",
    "Adapter service deployed; adapter.transport=in-memory configured; adapter.receiver.enabled=true; adapter.handler.enabled=true; adapter.dispatcher.enabled=true",
    [
        ("Set configuration property adapter.transport=in-memory", "Configuration accepted without error"),
        ("Start the Adapter service (Receiver, Handler, Dispatcher)", "Adapter starts successfully without errors"),
        ("Check startup logs for transport mode indication", "Logs indicate 'in-memory transport active' or equivalent Vert.x EventBus initialization message"),
        ("Verify no external Kafka broker connection is attempted", "No Kafka connection attempts in logs; no external infrastructure dependency"),
        ("Verify Receiver, Handler, and Dispatcher modules are all initialized", "All three modules report 'started' or 'initialized' in logs"),
    ],
    "High"
)

add_test_case(
    "IMT-002", "In-Memory End-to-End Flow — Receiver to Handler to Dispatcher", "In-Memory Transport",
    "Verify complete message flow from Receiver through Handler to Dispatcher using in-memory transport (Vert.x EventBus)",
    "Adapter running with in-memory transport; all modules enabled; dispatch-center.ucri.enabled=true",
    [
        ("Send a valid event message to the Receiver via EventReceiverPort", "Receiver accepts the event without errors"),
        ("Verify Receiver validates the incoming event", "Validation passes; event is forwarded to Handler via in-memory EventBus"),
        ("Verify Handler receives the event from Receiver", "Handler logs receipt of event with correct correlationId"),
        ("Verify Handler processes and transforms the event", "Handler transforms the event to the target format"),
        ("Verify Handler forwards the processed event to Dispatcher via in-memory EventBus", "Dispatcher receives the processed event"),
        ("Verify Dispatcher routes the event to correct dispatch center topic (events-ucri)", "Event is published to events-ucri topic"),
        ("Verify envelope metadata (correlationId, headers) is preserved end-to-end", "correlationId and headers match original event across all three services"),
    ],
    "High"
)

add_test_case(
    "IMT-003", "In-Memory Transport — Event Ordering Preservation", "In-Memory Transport",
    "Verify that events with the same partition key are delivered in published order via in-memory transport",
    "Adapter running with in-memory transport; all modules enabled",
    [
        ("Publish Event-A with partitionKey='PK-001' at time T1", "Event-A accepted by Receiver"),
        ("Publish Event-B with partitionKey='PK-001' at time T2 (T2 > T1)", "Event-B accepted by Receiver"),
        ("Publish Event-C with partitionKey='PK-001' at time T3 (T3 > T2)", "Event-C accepted by Receiver"),
        ("Verify Dispatcher processes events in order: Event-A, Event-B, Event-C", "Events delivered in published order (A → B → C) at Dispatcher output"),
        ("Verify ordering is preserved based on partition key", "All events with same partitionKey maintain FIFO order"),
    ],
    "High"
)

add_test_case(
    "IMT-004", "In-Memory Transport — Envelope Metadata Preservation", "In-Memory Transport",
    "Verify that message envelope metadata (correlationId, headers, partitionKey) is preserved across Receiver → Handler → Dispatcher",
    "Adapter running with in-memory transport; all modules enabled",
    [
        ("Create event with correlationId='CORR-12345', headers={'source':'IPL','type':'CreateRequest'}, partitionKey='PK-001'", "Event created with all metadata fields populated"),
        ("Send event to Receiver", "Receiver logs event receipt with correlationId='CORR-12345'"),
        ("Verify Handler receives event with same correlationId", "Handler log shows correlationId='CORR-12345'"),
        ("Verify Dispatcher receives event with same correlationId and headers", "Dispatcher log shows correlationId='CORR-12345' and all headers intact"),
        ("Verify outbound event published by Dispatcher contains all original metadata", "Published event envelope contains correlationId='CORR-12345', headers={'source':'IPL','type':'CreateRequest'}, partitionKey='PK-001'"),
    ],
    "High"
)

add_test_case(
    "IMT-005", "In-Memory Transport — Invalid Event Rejected by Receiver", "In-Memory Transport",
    "Verify that an invalid/malformed event is rejected at the Receiver and not forwarded to Handler or Dispatcher",
    "Adapter running with in-memory transport; all modules enabled",
    [
        ("Send a malformed event (missing required fields) to Receiver", "Receiver rejects the event with a validation error"),
        ("Check Receiver logs for validation failure reason", "Logs contain clear validation error description"),
        ("Verify Handler does NOT receive the rejected event", "No Handler processing log for the rejected event"),
        ("Verify Dispatcher does NOT receive or publish the rejected event", "No Dispatcher log entry for the rejected event"),
    ],
    "High"
)

add_test_case(
    "IMT-006", "In-Memory Transport — Receiver Module Disabled", "In-Memory Transport / Module Config",
    "Verify that disabling the Receiver module prevents event ingestion while Handler and Dispatcher remain unaffected",
    "Adapter deployed; adapter.receiver.enabled=false; adapter.handler.enabled=true; adapter.dispatcher.enabled=true",
    [
        ("Start Adapter with adapter.receiver.enabled=false", "Adapter starts successfully"),
        ("Verify Receiver module is NOT instantiated", "Logs confirm Receiver module is disabled; no event subscriptions created"),
        ("Attempt to send an event to the Receiver endpoint", "Event is not consumed; no processing occurs"),
        ("Verify no resources consumed by Receiver", "No Receiver threads or subscriptions active"),
    ],
    "Medium"
)

add_test_case(
    "IMT-007", "In-Memory Transport — Handler Module Disabled", "In-Memory Transport / Module Config",
    "Verify that disabling the Handler module stops event processing while Receiver still accepts events",
    "Adapter deployed; adapter.receiver.enabled=true; adapter.handler.enabled=false; adapter.dispatcher.enabled=true",
    [
        ("Start Adapter with adapter.handler.enabled=false", "Adapter starts successfully"),
        ("Send a valid event to the Receiver", "Receiver accepts and validates the event"),
        ("Verify Handler does NOT process the event", "No Handler processing log; event stops after Receiver"),
        ("Verify Dispatcher does NOT receive or dispatch the event", "No Dispatcher activity for this event"),
    ],
    "Medium"
)

add_test_case(
    "IMT-008", "In-Memory Transport — Dispatcher Module Disabled", "In-Memory Transport / Module Config",
    "Verify that disabling the Dispatcher module prevents event routing while Receiver and Handler remain functional",
    "Adapter deployed; adapter.receiver.enabled=true; adapter.handler.enabled=true; adapter.dispatcher.enabled=false",
    [
        ("Start Adapter with adapter.dispatcher.enabled=false", "Adapter starts successfully"),
        ("Send a valid event to Receiver", "Receiver accepts and validates the event"),
        ("Verify Handler processes the event", "Handler transformation completes successfully"),
        ("Verify Dispatcher does NOT route the event to any dispatch center", "No outbound event published; Dispatcher module inactive"),
    ],
    "Medium"
)

add_test_case(
    "IMT-009", "In-Memory Transport — Invalid Transport Configuration", "In-Memory Transport / Config",
    "Verify that an invalid transport configuration value causes startup failure with a descriptive error",
    "Adapter deployed; adapter.transport=invalid_value",
    [
        ("Set adapter.transport=invalid_value in configuration", "Configuration is set"),
        ("Attempt to start the Adapter service", "Startup fails"),
        ("Check error output/logs for descriptive error message", "Error message lists valid transport options (in-memory, kafka)"),
        ("Correct configuration to adapter.transport=in-memory and restart", "Adapter starts successfully"),
    ],
    "Medium"
)

add_test_case(
    "IMT-010", "In-Memory Transport — Missing Transport Configuration", "In-Memory Transport / Config",
    "Verify that missing transport configuration property causes startup failure",
    "Adapter deployed; adapter.transport property removed from configuration",
    [
        ("Remove adapter.transport property from configuration", "Property is absent"),
        ("Attempt to start the Adapter service", "Startup fails with a descriptive error"),
        ("Verify error message indicates missing transport configuration", "Error message clearly states transport property is required"),
    ],
    "Medium"
)

add_test_case(
    "IMT-011", "In-Memory Transport — Retry Decorator on Publish Failure", "In-Memory Transport / Cross-Cutting",
    "Verify the retry decorator retries failed publish operations with configurable backoff",
    "Adapter running with in-memory transport; dispatch-center.ucri.retry-attempts=3; dispatch-center.ucri.retry-delay-ms=1000",
    [
        ("Simulate a transient Dispatcher publish failure for the first attempt", "Publish fails; retry mechanism triggered"),
        ("Verify first retry occurs after ~1000ms delay", "Retry attempt 1 logged with timestamp ~1s after initial failure"),
        ("Verify second retry if first retry also fails", "Retry attempt 2 logged"),
        ("Verify third and final retry if second also fails", "Retry attempt 3 logged; maximum retries reached"),
        ("Verify metrics reflect all retry attempts", "Publish failure count and retry count metrics updated"),
    ],
    "Medium"
)

add_test_case(
    "IMT-012", "In-Memory Transport — Dead Letter Queue (DLQ)", "In-Memory Transport / Cross-Cutting",
    "Verify events that fail permanently after all retries are sent to the Dead Letter Queue",
    "Adapter running with in-memory transport; DLQ configured (dispatch-center.ucri.dlq-topic=events-ucri-dlq); retry-attempts=3",
    [
        ("Send event that will cause permanent publish failure after all retries", "Event fails 3 retries"),
        ("Verify event is sent to DLQ topic (events-ucri-dlq)", "Event present in DLQ topic"),
        ("Verify DLQ entry contains the original event payload", "Original payload matches sent event"),
        ("Verify DLQ entry contains error context (reason for failure)", "Error context/reason attached to DLQ entry"),
    ],
    "Medium"
)

add_test_case(
    "IMT-013", "In-Memory Transport — Logging and Traceability Across All 3 Services", "In-Memory Transport / Observability",
    "Verify comprehensive logging with correlationId traceability across Receiver, Handler, and Dispatcher",
    "Adapter running with in-memory transport; logging enabled",
    [
        ("Send a valid event with correlationId='TRACE-001' to Receiver", "Receiver logs event receipt with correlationId='TRACE-001'"),
        ("Verify Receiver logs contain event validation details", "Receiver log entry shows validation pass with correlationId='TRACE-001'"),
        ("Verify Handler logs contain transformation details with same correlationId", "Handler log entry shows processing with correlationId='TRACE-001'"),
        ("Verify Dispatcher logs contain dispatch details with same correlationId", "Dispatcher log entry shows outbound publish with correlationId='TRACE-001'"),
        ("Verify end-to-end traceability: all log entries linkable via correlationId", "All three services' logs are correlated via 'TRACE-001'"),
    ],
    "High"
)

# ============================================================
# SECTION 2: IPL TO UCRI (FORWARD FLOW) TEST CASES
# ============================================================

add_test_case(
    "FWD-001", "IPL to UCRI — Happy Path: Event Accepted (EMPFANG)", "IPL to UCRI / Forward Flow",
    "Verify end-to-end forward flow: IPL sends event via Adapter to UCRM, UCRM accepts, Adapter propagates EMPFANG ACK back to IPL",
    "Adapter running (in-memory transport); IPL operational (or mock); UCRM accessible; P2P token valid; dispatch-center.ucri.enabled=true",
    [
        ("IPL sends valid TransferEinsatzRequest to Adapter via POST /einsatztransfer/erzeugeEinsatz", "Adapter Receiver accepts the request; HTTP 200/202 returned"),
        ("Verify Receiver validates the incoming request (format, required fields)", "Validation passes; no errors logged"),
        ("Verify Handler authenticates to UCRM via P2P token", "P2P authentication successful"),
        ("Verify Handler transforms IPL payload to UCRM format (sourceMandant → header.sender_client_id, targetMandant → header.receiver_client_id, einsatzId → operation.id, etc.)", "All field mappings applied correctly"),
        ("Verify Dispatcher forwards message to UCRM via POST messaging/send", "Message sent to UCRM successfully"),
        ("UCRM processes and returns Accept response", "Adapter receives Accept response from UCRM"),
        ("Verify Handler translates Accept response to IPL ACK format", "ACK payload contains EMPFANG status"),
        ("Verify Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Accept status", "IPL receives EMPFANG ACK"),
        ("Verify log entry: Einsatztransfer/quittiereTranferEinsatz contains 'EMPFANG'", "Log shows EMPFANG for this transaction"),
    ],
    "High"
)

add_test_case(
    "FWD-002", "IPL to UCRI — Event Rejected (EINSATZ_ABGELEHNT)", "IPL to UCRI / Forward Flow",
    "Verify forward flow where UCRM rejects the event: Adapter propagates EINSATZ_ABGELEHNT ACK to IPL",
    "Adapter running; IPL operational; UCRM accessible; P2P token valid",
    [
        ("IPL sends valid TransferEinsatzRequest to Adapter via POST /einsatztransfer/erzeugeEinsatz", "Adapter Receiver accepts the request"),
        ("Receiver validates incoming request", "Validation passes"),
        ("Handler authenticates to UCRM and transforms payload", "Authentication and transformation successful"),
        ("Dispatcher forwards message to UCRM via messaging/send", "Message sent to UCRM"),
        ("UCRM processes and returns Reject response", "Adapter receives Reject response from UCRM"),
        ("Handler translates Reject response to IPL ACK format", "ACK payload contains EINSATZ_ABGELEHNT status"),
        ("Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Reject status", "IPL receives EINSATZ_ABGELEHNT ACK"),
        ("Verify log entry: Einsatztransfer/quittiereTranferEinsatz contains 'EINSATZ_ABGELEHNT'", "Log shows EINSATZ_ABGELEHNT for this transaction"),
    ],
    "High"
)

add_test_case(
    "FWD-003", "IPL to UCRI — Timeout (TECHNISCHER_FEHLER)", "IPL to UCRI / Forward Flow",
    "Verify forward flow where UCRM does not respond within allowed time: Adapter propagates TECHNISCHER_FEHLER ACK to IPL",
    "Adapter running; IPL operational; UCRM accessible but configured to not respond (or delayed); timeout-ms configured",
    [
        ("IPL sends valid TransferEinsatzRequest to Adapter via POST /einsatztransfer/erzeugeEinsatz", "Adapter Receiver accepts the request"),
        ("Receiver validates and Handler authenticates/transforms payload", "Validation and auth pass; message transformed"),
        ("Dispatcher forwards message to UCRM via messaging/send", "Message sent to UCRM"),
        ("UCRM does NOT respond within the configured timeout (e.g., 10 seconds)", "Timeout condition detected by Handler"),
        ("Handler detects timeout and prepares TECHNISCHER_FEHLER ACK", "ACK payload contains TECHNISCHER_FEHLER status"),
        ("Dispatcher calls POST /einsatztransfer/quittiereTranferEinsatz with Timeout error status", "IPL receives TECHNISCHER_FEHLER ACK"),
        ("Verify log entry: Einsatztransfer/quittiereTranferEinsatz contains 'TECHNISCHER_FEHLER'", "Log shows TECHNISCHER_FEHLER for this transaction"),
    ],
    "High"
)

add_test_case(
    "FWD-004", "IPL to UCRI — Validation Failure at Receiver", "IPL to UCRI / Forward Flow",
    "Verify that an invalid/malformed TransferEinsatzRequest from IPL is rejected at the Receiver and not forwarded further",
    "Adapter running; IPL operational",
    [
        ("IPL sends an invalid/malformed TransferEinsatzRequest (e.g., missing einsatzId, invalid format) to POST /einsatztransfer/erzeugeEinsatz", "Adapter returns validation error response to IPL"),
        ("Verify Receiver logs validation failure reason", "Log contains specific validation error details"),
        ("IPL handles the validation response from Adapter", "IPL processes the error response correctly"),
        ("Verify NO message is forwarded to Handler", "No Handler processing log for this request"),
        ("Verify NO message is sent to UCRM via Dispatcher", "No messaging/send call to UCRM"),
        ("Verify NO Einsatztransfer/quittiereTranferEinsatz log entry is generated", "No ACK log entry for the rejected request"),
    ],
    "High"
)

add_test_case(
    "FWD-005", "IPL to UCRI — P2P Authentication Failure", "IPL to UCRI / Forward Flow",
    "Verify that P2P authentication failure at Handler prevents message forwarding to UCRM and returns error to IPL",
    "Adapter running; IPL operational; P2P token invalid/expired",
    [
        ("IPL sends valid TransferEinsatzRequest to POST /einsatztransfer/erzeugeEinsatz", "Receiver accepts and validates (passes)"),
        ("Receiver forwards event to Handler", "Handler receives the event"),
        ("Handler attempts P2P authentication with invalid/expired token", "Authentication fails at UCRM"),
        ("Handler logs authentication failure details", "Log contains P2P auth failure reason"),
        ("Error is propagated back to IPL", "IPL receives authentication error response"),
        ("Verify NO message forwarded to UCRM via Dispatcher", "No messaging/send call made"),
    ],
    "High"
)

add_test_case(
    "FWD-006", "IPL to UCRI — Field Mapping Validation (TransferEinsatz to UCRM)", "IPL to UCRI / Data Mapping",
    "Verify all field mappings from IPL TransferEinsatz format to UCRM operation format are correctly applied",
    "Adapter running; IPL operational; UCRM accessible",
    [
        ("Send TransferEinsatzRequest with sourceMandant='IPL_CENTER_001'", "Handler maps to header.sender_client_id='IPL_CENTER_001'"),
        ("Verify targetMandant='UCRI_CENTER_001' mapping", "Handler maps to header.receiver_client_id='UCRI_CENTER_001'"),
        ("Verify transferEinsatz.einsatzId='EIN-12345' mapping", "Handler maps to operation.id='EIN-12345'"),
        ("Verify transferEinsatz.einsatznummerKurz='K-98765' mapping", "Handler maps to operation.name='K-98765'"),
        ("Verify transferEinsatz.kategorie='KAT-01' mapping", "Handler maps to operation.occassion.type_1='KAT-01'"),
        ("Verify transferEinsatz.einsatzStichwort='FEUER' mapping", "Handler maps to operation.occassion.keyword_1='FEUER'"),
        ("Verify transferEinsatz.alarmierungsstichwortText='Brand in Wohngebäude' mapping", "Handler maps to operation.occassion.keyword_1_text='Brand in Wohngebäude'"),
        ("Verify no data loss in field translation", "All mapped fields present in UCRM message; no truncation or corruption"),
    ],
    "High"
)

add_test_case(
    "FWD-007", "IPL to UCRI — Missing sourceMandant in Request", "IPL to UCRI / Validation",
    "Verify that a TransferEinsatzRequest missing the required sourceMandant field is rejected",
    "Adapter running; IPL operational",
    [
        ("Send TransferEinsatzRequest without sourceMandant field", "Receiver rejects the request with validation error"),
        ("Verify error response indicates 'sourceMandant is required'", "Clear validation error message returned to IPL"),
        ("Verify no forwarding to Handler or UCRM", "No Handler or Dispatcher processing for this request"),
    ],
    "High"
)

add_test_case(
    "FWD-008", "IPL to UCRI — Missing targetMandant in Request", "IPL to UCRI / Validation",
    "Verify that a TransferEinsatzRequest missing the required targetMandant field is rejected",
    "Adapter running; IPL operational",
    [
        ("Send TransferEinsatzRequest without targetMandant field", "Receiver rejects the request with validation error"),
        ("Verify error response indicates 'targetMandant is required'", "Clear validation error message returned to IPL"),
        ("Verify no forwarding to Handler or UCRM", "No Handler or Dispatcher processing for this request"),
    ],
    "High"
)

add_test_case(
    "FWD-009", "IPL to UCRI — Missing einsatzId in Request", "IPL to UCRI / Validation",
    "Verify that a TransferEinsatzRequest missing the required einsatzId field is rejected",
    "Adapter running; IPL operational",
    [
        ("Send TransferEinsatzRequest without einsatzId field", "Receiver rejects the request with validation error"),
        ("Verify error response indicates 'einsatzId is required'", "Clear validation error message returned to IPL"),
        ("Verify no forwarding to Handler or UCRM", "No Handler or Dispatcher processing for this request"),
    ],
    "Medium"
)

add_test_case(
    "FWD-010", "IPL to UCRI — Empty Payload", "IPL to UCRI / Validation",
    "Verify that an empty payload sent to /einsatztransfer/erzeugeEinsatz is rejected",
    "Adapter running; IPL operational",
    [
        ("Send POST /einsatztransfer/erzeugeEinsatz with empty body", "Receiver rejects the request with validation error"),
        ("Verify error response indicates invalid payload", "Error message clearly describes empty/invalid payload"),
        ("Verify no forwarding to Handler or UCRM", "No Handler or Dispatcher processing"),
    ],
    "Medium"
)

add_test_case(
    "FWD-011", "IPL to UCRI — UCRI Dispatch Center Disabled", "IPL to UCRI / Config",
    "Verify that when UCRI dispatch center is disabled, events targeted for UCRI are not dispatched",
    "Adapter running; dispatch-center.ucri.enabled=false",
    [
        ("Set dispatch-center.ucri.enabled=false", "Configuration applied"),
        ("Send TransferEinsatzRequest with targetMandant pointing to UCRI", "Receiver accepts the event"),
        ("Verify Dispatcher does NOT route event to events-ucri topic", "No event published to UCRI dispatch center"),
        ("Verify appropriate error or handling for disabled dispatch center", "Event handled per system behavior for disabled target"),
    ],
    "Medium"
)

add_test_case(
    "FWD-012", "IPL to UCRI — CorrelationId Traceability in Forward Flow", "IPL to UCRI / Observability",
    "Verify correlationId is maintained across all components in the forward flow for end-to-end traceability",
    "Adapter running; IPL operational; UCRM accessible; logging enabled",
    [
        ("IPL sends TransferEinsatzRequest with a unique correlationId", "Receiver logs receipt with the correlationId"),
        ("Verify Handler logs processing with the same correlationId", "Handler log entry contains matching correlationId"),
        ("Verify Dispatcher logs forwarding to UCRM with the same correlationId", "Dispatcher log entry contains matching correlationId"),
        ("Verify ACK response (quittiereTranferEinsatz) log contains the same correlationId", "ACK log entry linkable to original request via correlationId"),
    ],
    "High"
)

add_test_case(
    "FWD-013", "IPL to UCRI — Multiple Events Same Partition Key", "IPL to UCRI / Ordering",
    "Verify that multiple events from IPL with the same partition key are delivered to UCRM in published order",
    "Adapter running; IPL operational; UCRM accessible",
    [
        ("IPL sends Event-1 with partitionKey='INCIDENT-001' at T1", "Event-1 received by Adapter"),
        ("IPL sends Event-2 with partitionKey='INCIDENT-001' at T2 (T2 > T1)", "Event-2 received by Adapter"),
        ("IPL sends Event-3 with partitionKey='INCIDENT-001' at T3 (T3 > T2)", "Event-3 received by Adapter"),
        ("Verify UCRM receives events in order: Event-1, Event-2, Event-3", "Events arrive at UCRM in FIFO order per partition key"),
    ],
    "Medium"
)

add_test_case(
    "FWD-014", "IPL to UCRI — P2P Token Expired During Active Session", "IPL to UCRI / Authentication",
    "Verify Adapter handles scenario where P2P token expires during an active messaging session",
    "Adapter running; IPL operational; UCRM accessible; P2P token set to expire during test",
    [
        ("IPL sends first TransferEinsatzRequest — P2P token is valid", "First event processed and delivered to UCRM successfully"),
        ("Wait for P2P token to expire", "Token expires"),
        ("IPL sends second TransferEinsatzRequest — P2P token is now expired", "Handler detects expired token during authentication"),
        ("Verify appropriate error handling for expired token", "Error propagated to IPL; no message forwarded to UCRM"),
        ("Verify error logged with details", "Log contains P2P token expiry error details"),
    ],
    "Medium"
)

# ============================================================
# SECTION 3: UCRI TO IPL (REVERSE FLOW) TEST CASES
# ============================================================

add_test_case(
    "REV-001", "UCRI to IPL — Happy Path: Mission Accepted by IPL", "UCRI to IPL / Reverse Flow",
    "Verify end-to-end reverse flow: UCRM sends mission to IPL via Adapter, IPL accepts and sends Accept ACK back to UCRM",
    "Adapter running (in-memory transport); UCRM operational (or test client); IPL accessible; P2P connection established",
    [
        ("UCRM sends mission payload to Adapter via POST messaging/send", "Adapter Receiver accepts the mission payload"),
        ("Verify Receiver validates the incoming mission data", "Validation passes"),
        ("Verify Handler translates UCRM mission format to IPL format", "Mission data transformed to IPL-compatible format"),
        ("Verify Dispatcher forwards the translated mission to IPL", "IPL receives the mission"),
        ("IPL processes the mission successfully", "Mission processed without errors"),
        ("IPL sends Accept ACK back to Adapter", "Adapter Receiver receives Accept ACK from IPL"),
        ("Handler translates IPL ACK to UCRM ACK format", "ACK translated"),
        ("Dispatcher propagates Accept ACK to UCRM", "UCRM receives Accept ACK confirming mission acceptance"),
    ],
    "High"
)

add_test_case(
    "REV-002", "UCRI to IPL — Mission Rejected by IPL", "UCRI to IPL / Reverse Flow",
    "Verify reverse flow where IPL rejects the mission: Adapter propagates Reject ACK to UCRM",
    "Adapter running; UCRM operational; IPL accessible",
    [
        ("UCRM sends mission payload to Adapter via messaging/send", "Adapter receives the mission"),
        ("Receiver validates and Handler translates mission to IPL format", "Translation complete"),
        ("Dispatcher forwards mission to IPL", "IPL receives the mission"),
        ("IPL rejects the mission", "IPL sends Reject ACK to Adapter"),
        ("Adapter Receiver receives Reject ACK", "Reject ACK received"),
        ("Handler translates Reject ACK to UCRM format", "ACK translated"),
        ("Dispatcher propagates Reject ACK to UCRM", "UCRM receives Reject ACK confirming mission rejection"),
    ],
    "High"
)

add_test_case(
    "REV-003", "UCRI to IPL — IPL Does Not Respond (Timeout)", "UCRI to IPL / Reverse Flow",
    "Verify reverse flow where IPL does not respond within allowed time: Adapter handles timeout scenario",
    "Adapter running; UCRM operational; IPL configured to not respond (or delayed)",
    [
        ("UCRM sends mission payload to Adapter via messaging/send", "Adapter receives the mission"),
        ("Adapter forwards mission to IPL", "Mission sent to IPL"),
        ("IPL does NOT respond within the configured timeout", "Timeout condition detected by Adapter"),
        ("Verify Adapter handles the timeout appropriately", "Timeout error logged"),
        ("Verify UCRM is notified of the timeout/failure", "UCRM receives timeout notification or error response"),
    ],
    "High"
)

add_test_case(
    "REV-004", "UCRI to IPL — Invalid Mission Payload from UCRM", "UCRI to IPL / Reverse Flow",
    "Verify that an invalid/malformed mission payload from UCRM is rejected at the Adapter Receiver",
    "Adapter running; UCRM operational",
    [
        ("UCRM sends malformed/incomplete mission payload via messaging/send", "Adapter Receiver rejects the payload with validation error"),
        ("Verify validation error logged at Receiver", "Log contains validation failure details"),
        ("Verify NO mission is forwarded to IPL", "No Dispatcher forwarding to IPL"),
        ("Verify error response sent back to UCRM", "UCRM receives rejection/error notification"),
    ],
    "High"
)

add_test_case(
    "REV-005", "UCRI to IPL — Field Mapping Validation (UCRM to IPL)", "UCRI to IPL / Data Mapping",
    "Verify that UCRM mission fields are correctly mapped to IPL format during reverse flow",
    "Adapter running; UCRM operational; IPL accessible",
    [
        ("UCRM sends mission with header.sender_client_id='UCRI_CENTER_001'", "Handler maps to IPL sourceMandant='UCRI_CENTER_001'"),
        ("Verify header.receiver_client_id='IPL_CENTER_001' mapping", "Handler maps to IPL targetMandant='IPL_CENTER_001'"),
        ("Verify operation.id='OP-56789' mapping", "Handler maps to IPL einsatzId='OP-56789'"),
        ("Verify operation.name='K-11111' mapping", "Handler maps to IPL einsatznummerKurz='K-11111'"),
        ("Verify no data loss in reverse field translation", "All mapped fields present in IPL message; no truncation or corruption"),
    ],
    "High"
)

add_test_case(
    "REV-006", "UCRI to IPL — Reverse Flow CorrelationId Traceability", "UCRI to IPL / Observability",
    "Verify correlationId is maintained across all components in the reverse flow",
    "Adapter running; UCRM operational; IPL accessible; logging enabled",
    [
        ("UCRM sends mission with correlationId='REV-TRACE-001'", "Receiver logs receipt with correlationId='REV-TRACE-001'"),
        ("Verify Handler logs translation with same correlationId", "Handler log entry contains 'REV-TRACE-001'"),
        ("Verify Dispatcher logs forwarding to IPL with same correlationId", "Dispatcher log entry contains 'REV-TRACE-001'"),
        ("Verify IPL ACK response is logged with same correlationId", "ACK log traceable via 'REV-TRACE-001'"),
    ],
    "High"
)

add_test_case(
    "REV-007", "UCRI to IPL — Multiple Missions Sequential Processing", "UCRI to IPL / Reverse Flow",
    "Verify that multiple missions from UCRM are processed sequentially and independently",
    "Adapter running; UCRM operational; IPL accessible",
    [
        ("UCRM sends Mission-A to Adapter", "Mission-A accepted and processing starts"),
        ("UCRM sends Mission-B to Adapter", "Mission-B accepted and queued"),
        ("Verify Mission-A is fully processed (forwarded to IPL, ACK returned to UCRM)", "Mission-A end-to-end complete"),
        ("Verify Mission-B is then processed independently", "Mission-B end-to-end complete with independent processing"),
        ("Verify each mission has distinct correlationId in logs", "Logs show separate correlationIds for A and B"),
    ],
    "Medium"
)

add_test_case(
    "REV-008", "UCRI to IPL — ACK Propagation Timing", "UCRI to IPL / Reverse Flow",
    "Verify ACK from IPL is propagated back to UCRM within acceptable time limits",
    "Adapter running; UCRM operational; IPL accessible; ACK timeout configured",
    [
        ("UCRM sends mission to Adapter", "Mission received by Adapter"),
        ("Adapter forwards to IPL", "IPL receives mission"),
        ("IPL sends Accept ACK back to Adapter", "Adapter receives ACK"),
        ("Verify ACK is propagated to UCRM within configured timeout (e.g., 10 seconds)", "UCRM receives ACK within acceptable time"),
        ("Verify no timeout error triggered for this flow", "No TECHNISCHER_FEHLER logged"),
    ],
    "Medium"
)

# ============================================================
# SECTION 4: UCRI (CLIENT-1) TO UCRI (CLIENT-2) FLOW
# ============================================================

add_test_case(
    "U2U-001", "UCRI-1 to UCRI-2 — Happy Path: Mission Dispatched and Accepted", "UCRI to UCRI / Cross-Client",
    "Verify end-to-end flow from one UCRI client (UCRI-1) to another UCRI client (UCRI-2) via the Adapter, where UCRI-2 accepts the mission",
    "Adapter running; UCRI-1 (sender) and UCRI-2 (receiver) both accessible; P2P tokens valid for both clients; dispatch-center.ucri.enabled=true",
    [
        ("UCRI-1 (sourceMandant) sends mission/event to Adapter via messaging/send with targetMandant=UCRI-2", "Adapter Receiver accepts the payload from UCRI-1"),
        ("Verify Receiver validates the incoming request from UCRI-1", "Validation passes; sourceMandant and targetMandant both identify UCRI clients"),
        ("Verify Handler authenticates via P2P token for UCRI-2 target", "P2P authentication successful for UCRI-2 connection"),
        ("Verify Handler transforms UCRI-1 message format to UCRI-2 compatible format", "Message correctly transformed (if format differences exist between clients)"),
        ("Verify Dispatcher forwards the message to UCRI-2 via messaging/send", "Message delivered to UCRI-2"),
        ("UCRI-2 processes and returns Accept response", "Adapter receives Accept from UCRI-2"),
        ("Verify Handler translates Accept response to UCRI-1 ACK format", "ACK payload prepared for UCRI-1"),
        ("Verify Dispatcher propagates Accept ACK back to UCRI-1", "UCRI-1 receives confirmation that UCRI-2 accepted the mission"),
    ],
    "High"
)

add_test_case(
    "U2U-002", "UCRI-1 to UCRI-2 — Mission Rejected by UCRI-2", "UCRI to UCRI / Cross-Client",
    "Verify cross-client flow where UCRI-2 rejects the mission from UCRI-1: Reject ACK propagated back to UCRI-1",
    "Adapter running; UCRI-1 and UCRI-2 accessible; P2P tokens valid",
    [
        ("UCRI-1 sends mission/event to Adapter with targetMandant=UCRI-2", "Adapter receives the event from UCRI-1"),
        ("Adapter validates, authenticates, transforms, and forwards to UCRI-2", "Message delivered to UCRI-2"),
        ("UCRI-2 rejects the mission and returns Reject response", "Adapter receives Reject from UCRI-2"),
        ("Handler translates Reject to UCRI-1 ACK format", "Reject ACK prepared"),
        ("Dispatcher propagates Reject ACK back to UCRI-1", "UCRI-1 receives notification that UCRI-2 rejected the mission"),
    ],
    "High"
)

add_test_case(
    "U2U-003", "UCRI-1 to UCRI-2 — Timeout (UCRI-2 Does Not Respond)", "UCRI to UCRI / Cross-Client",
    "Verify cross-client flow where UCRI-2 does not respond within the allowed time: Timeout handled and UCRI-1 notified",
    "Adapter running; UCRI-1 accessible; UCRI-2 configured to not respond (or delayed); timeout-ms configured",
    [
        ("UCRI-1 sends mission/event to Adapter with targetMandant=UCRI-2", "Adapter receives the event"),
        ("Adapter forwards to UCRI-2 via messaging/send", "Message sent to UCRI-2"),
        ("UCRI-2 does NOT respond within configured timeout", "Timeout detected by Adapter"),
        ("Verify Handler handles timeout condition", "Timeout error logged"),
        ("Verify Dispatcher propagates timeout notification back to UCRI-1", "UCRI-1 receives timeout error/notification"),
    ],
    "High"
)

add_test_case(
    "U2U-004", "UCRI-2 to UCRI-1 — Reverse Direction: Mission Dispatched and Accepted", "UCRI to UCRI / Cross-Client",
    "Verify bi-directional capability: UCRI-2 sends mission to UCRI-1, UCRI-1 accepts",
    "Adapter running; UCRI-1 and UCRI-2 accessible; P2P tokens valid for both directions",
    [
        ("UCRI-2 (sourceMandant) sends mission/event to Adapter with targetMandant=UCRI-1", "Adapter Receiver accepts payload from UCRI-2"),
        ("Adapter validates, authenticates, transforms, and forwards to UCRI-1", "Message delivered to UCRI-1"),
        ("UCRI-1 processes and returns Accept response", "Adapter receives Accept from UCRI-1"),
        ("Adapter propagates Accept ACK back to UCRI-2", "UCRI-2 receives confirmation that UCRI-1 accepted the mission"),
    ],
    "High"
)

add_test_case(
    "U2U-005", "UCRI-1 to UCRI-2 — Invalid sourceMandant (Sender Not Recognized)", "UCRI to UCRI / Validation",
    "Verify that a request from an unrecognized UCRI sender client is rejected",
    "Adapter running; UCRI-2 accessible; UCRI-1 using invalid/unrecognized sourceMandant",
    [
        ("Send event with sourceMandant='UNKNOWN_CLIENT' and targetMandant='UCRI-2'", "Adapter Receiver rejects the request"),
        ("Verify validation error logged indicating unrecognized sender", "Error log contains 'invalid sourceMandant' or similar"),
        ("Verify NO message forwarded to UCRI-2", "No messaging/send call to UCRI-2"),
    ],
    "Medium"
)

add_test_case(
    "U2U-006", "UCRI-1 to UCRI-2 — Invalid targetMandant (Receiver Not Found)", "UCRI to UCRI / Validation",
    "Verify that a request with an invalid/unknown targetMandant is handled appropriately",
    "Adapter running; UCRI-1 accessible; targetMandant points to non-existent UCRI client",
    [
        ("Send event with sourceMandant='UCRI-1' and targetMandant='UCRI_NONEXISTENT'", "Adapter processes the request"),
        ("Verify Dispatcher cannot route to non-existent target", "Routing failure or rejection at Dispatcher"),
        ("Verify error response sent back to UCRI-1", "UCRI-1 receives error notification about invalid target"),
        ("Verify error logged with details", "Log contains routing failure details"),
    ],
    "Medium"
)

add_test_case(
    "U2U-007", "UCRI-1 to UCRI-2 — P2P Authentication Failure for Cross-Client", "UCRI to UCRI / Authentication",
    "Verify that P2P authentication failure when connecting to UCRI-2 prevents message delivery and notifies UCRI-1",
    "Adapter running; UCRI-1 accessible; P2P token for UCRI-2 invalid/expired",
    [
        ("UCRI-1 sends event with targetMandant=UCRI-2", "Adapter Receiver accepts the event"),
        ("Handler attempts P2P authentication for UCRI-2 with invalid token", "Authentication fails"),
        ("Verify error logged for P2P auth failure", "Log contains authentication failure details"),
        ("Verify error notification propagated to UCRI-1", "UCRI-1 receives P2P auth error response"),
        ("Verify NO message delivered to UCRI-2", "No messaging/send call to UCRI-2"),
    ],
    "High"
)

add_test_case(
    "U2U-008", "UCRI-1 to UCRI-2 — Field Mapping for Cross-Client Transfer", "UCRI to UCRI / Data Mapping",
    "Verify field mapping is correctly applied when transferring between two UCRI clients",
    "Adapter running; UCRI-1 and UCRI-2 accessible",
    [
        ("UCRI-1 sends mission with header.sender_client_id='UCRI_CENTER_001' (UCRI-1)", "Handler receives with correct sender identity"),
        ("Verify header.receiver_client_id='UCRI_CENTER_002' (UCRI-2) mapping", "Target client correctly identified as UCRI-2"),
        ("Verify operation.id is preserved in cross-client transfer", "operation.id matches between sender and receiver message"),
        ("Verify operation.name is preserved", "operation.name matches"),
        ("Verify occasion fields (type_1, keyword_1, keyword_1_text) are preserved", "All occasion fields correctly mapped"),
        ("Verify no data loss in cross-client field translation", "All fields present in delivered message to UCRI-2"),
    ],
    "High"
)

add_test_case(
    "U2U-009", "UCRI-1 to UCRI-2 — CorrelationId Traceability", "UCRI to UCRI / Observability",
    "Verify correlationId traceability across the entire cross-client flow from UCRI-1 to UCRI-2 and back",
    "Adapter running; UCRI-1 and UCRI-2 accessible; logging enabled",
    [
        ("UCRI-1 sends event with correlationId='U2U-TRACE-001'", "Receiver logs receipt with correlationId='U2U-TRACE-001'"),
        ("Verify Handler logs processing with same correlationId", "Handler log entry contains 'U2U-TRACE-001'"),
        ("Verify Dispatcher logs forwarding to UCRI-2 with same correlationId", "Dispatcher log entry contains 'U2U-TRACE-001'"),
        ("Verify UCRI-2 ACK response is logged with same correlationId", "ACK traceable via 'U2U-TRACE-001'"),
        ("Verify end-to-end traceability from UCRI-1 → Adapter → UCRI-2 → Adapter → UCRI-1", "All log entries linked via 'U2U-TRACE-001'"),
    ],
    "High"
)

add_test_case(
    "U2U-010", "UCRI-1 to UCRI-2 — Simultaneous Bi-Directional Transfer", "UCRI to UCRI / Cross-Client",
    "Verify the Adapter handles simultaneous transfers in both directions (UCRI-1→UCRI-2 and UCRI-2→UCRI-1) without interference",
    "Adapter running; UCRI-1 and UCRI-2 both operational; P2P tokens valid",
    [
        ("UCRI-1 sends Mission-A to UCRI-2 (event with targetMandant=UCRI-2)", "Adapter accepts Mission-A from UCRI-1"),
        ("Simultaneously, UCRI-2 sends Mission-B to UCRI-1 (event with targetMandant=UCRI-1)", "Adapter accepts Mission-B from UCRI-2"),
        ("Verify Mission-A is delivered to UCRI-2 correctly", "UCRI-2 receives Mission-A with correct data"),
        ("Verify Mission-B is delivered to UCRI-1 correctly", "UCRI-1 receives Mission-B with correct data"),
        ("Verify UCRI-2 ACK for Mission-A is propagated back to UCRI-1", "UCRI-1 receives ACK for Mission-A"),
        ("Verify UCRI-1 ACK for Mission-B is propagated back to UCRI-2", "UCRI-2 receives ACK for Mission-B"),
        ("Verify no data cross-contamination between Mission-A and Mission-B", "Each mission processed independently; correlationIds and data are distinct"),
    ],
    "High"
)

add_test_case(
    "U2U-011", "UCRI-1 to UCRI-2 — Retry on Dispatcher Failure", "UCRI to UCRI / Cross-Cutting",
    "Verify retry mechanism works for cross-client UCRI-to-UCRI flows when Dispatcher encounters transient failure",
    "Adapter running; UCRI-1 accessible; UCRI-2 temporarily unavailable; retry-attempts=3; retry-delay-ms=1000",
    [
        ("UCRI-1 sends event with targetMandant=UCRI-2", "Adapter receives the event"),
        ("Dispatcher attempts to forward to UCRI-2; first attempt fails (transient error)", "Retry triggered"),
        ("Verify first retry occurs after ~1000ms", "Retry attempt 1 logged"),
        ("UCRI-2 becomes available; retry succeeds", "Message delivered to UCRI-2 on retry"),
        ("Verify UCRI-2 ACK propagated back to UCRI-1", "UCRI-1 receives success ACK"),
        ("Verify retry count metrics updated", "Metrics show 1 initial failure + 1 successful retry"),
    ],
    "Medium"
)

add_test_case(
    "U2U-012", "UCRI-1 to UCRI-2 — DLQ After All Retries Exhausted", "UCRI to UCRI / Cross-Cutting",
    "Verify that when all retries are exhausted for UCRI-to-UCRI flow, the event is sent to DLQ",
    "Adapter running; UCRI-1 accessible; UCRI-2 permanently unavailable; retry-attempts=3; DLQ configured",
    [
        ("UCRI-1 sends event with targetMandant=UCRI-2", "Adapter receives the event"),
        ("Dispatcher fails to deliver to UCRI-2 after all 3 retry attempts", "All retries exhausted"),
        ("Verify event sent to DLQ topic (events-ucri-dlq)", "Event present in DLQ"),
        ("Verify DLQ entry contains original payload and failure reason", "DLQ entry has full context"),
        ("Verify UCRI-1 notified of delivery failure", "UCRI-1 receives error/timeout notification"),
    ],
    "Medium"
)

# ============================================================
# SECTION 5: UCRM TEST CLIENT VALIDATION
# ============================================================

add_test_case(
    "TC-001", "UCRM Test Client — Installation and Prerequisites", "Test Client / Setup",
    "Verify UCRM Test Client can be installed and all prerequisites are met",
    "UCRM Test Client installation package available; deployment guide accessible",
    [
        ("Verify all prerequisites listed in deployment guide are met", "All prerequisites confirmed"),
        ("Install UCRM Test Client per deployment guide", "Installation completes without errors"),
        ("Verify UCRM Test Client starts successfully", "Test Client is running"),
        ("Execute basic connectivity test to Adapter", "Connectivity confirmed"),
    ],
    "High"
)

add_test_case(
    "TC-002", "UCRM Test Client — Send Mission and Receive ACK", "Test Client / Functional",
    "Verify UCRM Test Client can send a mission and receive an ACK response through the Adapter",
    "UCRM Test Client installed and running; Adapter running; IPL accessible (or mock)",
    [
        ("Use UCRM Test Client to send a mission to the Adapter", "Mission accepted by Adapter"),
        ("Verify mission is forwarded to IPL", "IPL receives the mission"),
        ("IPL sends Accept ACK", "ACK received by Adapter"),
        ("Verify UCRM Test Client receives the Accept ACK", "Test Client displays Accept ACK"),
    ],
    "High"
)

# Save file
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
filename = f"c:\\Apparao Maroju\\AI-SSI\\Project-IGNIS(POC)\\UCRI2.0-Output\\TestCases_{timestamp}.xlsx"
wb.save(filename)
print(f"Test cases saved to: {filename}")
