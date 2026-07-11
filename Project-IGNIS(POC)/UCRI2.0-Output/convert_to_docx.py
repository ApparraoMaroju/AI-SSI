from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re
from datetime import datetime

def set_cell_shading(cell, fill_color):
    """Set cell background color"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), fill_color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_formatted_text(paragraph, text):
    """Add text with basic formatting (bold, inline code)"""
    # Handle bold text
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # Handle inline code
            code_parts = re.split(r'(`[^`]+`)', part)
            for code_part in code_parts:
                if code_part.startswith('`') and code_part.endswith('`'):
                    run = paragraph.add_run(code_part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                else:
                    paragraph.add_run(code_part)

def create_word_document():
    doc = Document()
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    # Title
    title = doc.add_heading('UCRI 2.0 Phase II – Integration Testing', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Overview Section
    doc.add_heading('Overview', level=1)
    overview_text = """As part of the end-to-end integration testing phase for UCRI 2.0 Phase II, comprehensive test scenarios were developed to validate the interface connectivity, data flow, and communication between IPL (Incident Processing Layer) and UCRI (Unified Central Routing Infrastructure) systems. The following test scenarios ensure that information is exchanged, processed, and synchronized as expected within the integrated environment."""
    doc.add_paragraph(overview_text)
    
    # Integration Scope
    doc.add_heading('Integration Scope', level=2)
    scope_items = [
        ('System A (Initiator/Receiver):', 'IPL (Incident Processing Layer)'),
        ('System B (Partner/Mock Service):', 'UCRI (Unified Central Routing Infrastructure)'),
        ('Communication Protocols:', 'REST API, Event-driven messaging'),
        ('Testing Focus:', 'Bidirectional event flow validation, error handling, edge cases, and boundary conditions'),
    ]
    for label, value in scope_items:
        p = doc.add_paragraph()
        run = p.add_run(label + ' ')
        run.bold = True
        p.add_run(value)
    
    # Test Categories Summary
    doc.add_heading('Test Categories Summary', level=1)
    
    # Create summary table
    summary_table = doc.add_table(rows=5, cols=3)
    summary_table.style = 'Table Grid'
    
    # Header row
    header_cells = summary_table.rows[0].cells
    headers = ['Category', 'Description', 'Test Count']
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_shading(header_cells[i], '4472C4')
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    # Data rows
    summary_data = [
        ('Functional', 'Core business flow validations for IPL-UCRI integration', '18'),
        ('Negative', 'Error handling, invalid inputs, and failure scenarios', '10'),
        ('Edge', 'Concurrent flows, race conditions, and stress scenarios', '4'),
        ('Boundary', 'Payload limits and boundary condition testing', '0'),
    ]
    for row_idx, (cat, desc, count) in enumerate(summary_data, start=1):
        row = summary_table.rows[row_idx]
        row.cells[0].text = cat
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = desc
        row.cells[2].text = count
    
    doc.add_paragraph()  # Add spacing
    
    # Test Matrix Section
    doc.add_heading('Test Matrix', level=1)
    
    # Test cases data
    test_cases = [
        {
            'id': 'UCRI2-EDGE-01',
            'name': 'Verify system handles simultaneous forward (IPL to UCRI) and reverse (UCRI to IPL) flows',
            'category': 'Edge',
            'prereq': '1. Adapter service running with adapter.transport=in-memory\n2. All modules enabled\n3. Both IPL and UCRM systems accessible',
            'steps': '1. Initiate an IPL to UCRI forward flow.\n2. Simultaneously initiate a UCRI to IPL reverse flow.\n3. Verify both flows complete successfully.\n4. Verify no message corruption or cross-contamination between flows.\n5. Verify correlationIds maintain separation across flows.',
            'expected': '• Forward flow started.\n• Reverse flow started.\n• Both flows complete without errors.\n• Messages isolated correctly. Each flow processes its own data.\n• Each flow has its own unique correlationId trail.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-EDGE-02',
            'name': 'Verify system handles rapid sequential events without message loss',
            'category': 'Edge',
            'prereq': '1. Adapter service running with adapter.transport=in-memory\n2. All modules enabled',
            'steps': '1. Send 10 events in rapid succession to the Adapter Receiver.\n2. Verify all 10 events are received by the Receiver.\n3. Verify all 10 events are processed through Handler.\n4. Verify all 10 events are dispatched by Dispatcher.\n5. Verify ordering is preserved for events with the same partition key.\n6. Verify no events are lost (count at Dispatcher matches count at Receiver).',
            'expected': '• Events sent quickly.\n• All 10 events received. Receiver logs show 10 events.\n• Handler processes all 10 events.\n• Dispatcher routes all 10 events.\n• Event order maintained.\n• Event count: 10 in = 10 out.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-EDGE-03',
            'name': 'Verify system handles simultaneous IPL to UCRI-1, IPL to UCRI-2, and UCRI-1 to UCRI-2 flows',
            'category': 'Edge',
            'prereq': '1. Adapter service running with adapter.transport=in-memory\n2. IPL, UCRI-1, UCRI-2 dispatch centers enabled\n3. All systems accessible',
            'steps': '1. Simultaneously initiate: Flow-A (IPL to UCRI-1), Flow-B (IPL to UCRI-2), Flow-C (UCRI-1 to UCRI-2).\n2. Verify Flow-A routes to events-ucri-1 topic.\n3. Verify Flow-B routes to events-ucri-2 topic.\n4. Verify Flow-C routes from UCRI-1 to events-ucri-2 topic.\n5. Verify all three flows complete independently without interference.\n6. Verify each flow has its own correlationId trail.',
            'expected': '• Three flows started concurrently.\n• Flow-A correctly routed to UCRI-1.\n• Flow-B correctly routed to UCRI-2.\n• Flow-C correctly routed to UCRI-2 from UCRI-1.\n• All flows complete. No cross-contamination or routing errors.\n• Three distinct correlationIds maintained.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IMT-01',
            'name': 'Verify In-Memory transport initialization and Receiver, Handler, Dispatcher flow activation',
            'category': 'Functional',
            'prereq': '1. Adapter service code deployed\n2. Configuration: adapter.transport=in-memory\n3. adapter.receiver.enabled=true\n4. adapter.handler.enabled=true\n5. adapter.dispatcher.enabled=true',
            'steps': '1. Set configuration property adapter.transport=in-memory and start the Adapter service.\n2. Verify startup logs indicate in-memory transport active.\n3. Verify that Receiver service is active and accepting events.\n4. Verify that Handler service is active and processing events.\n5. Verify that Dispatcher service is active and routing events.',
            'expected': '• Service starts without errors.\n• Logs show in-memory transport (Vert.x EventBus) activated.\n• Receiver module initialized. Event subscriptions created.\n• Handler module initialized. Ready to consume from in-memory bus.\n• Dispatcher module initialized. Ready to dispatch to configured dispatch centers.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IMT-02',
            'name': 'Verify end-to-end message flow via In-Memory transport (Receiver, Handler, Dispatcher)',
            'category': 'Functional',
            'prereq': '1. Adapter configured with adapter.transport=in-memory\n2. All modules enabled\n3. UCRI dispatch center configured (dispatch-center.ucri.topic=events-ucri)',
            'steps': '1. Send a valid event payload to the Receiver.\n2. Monitor application logs to trace the event passing from Receiver to Handler via in-memory bus.\n3. Verify Handler processes the event (transformation, business logic).\n4. Verify the Dispatcher routes the event to the correct dispatch center topic (events-ucri).\n5. Verify the message envelope metadata (correlationId, headers) is preserved end-to-end.\n6. Verify the payload is intact (no data loss).',
            'expected': '• Receiver accepts and validates the event.\n• Handler receives the event from Receiver. Logs confirm event handoff.\n• Handler processes and publishes transformed event to Dispatcher via in-memory bus.\n• Event appears on the UCRI dispatch center topic.\n• correlationId and headers are identical at each stage.\n• Payload content unchanged throughout the flow.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IPL-UCRI-01',
            'name': 'Verify that an IPL incident is successfully delivered to UCRM through the Adapter and UCRM accepts the incident (EMPFANG)',
            'category': 'Functional',
            'prereq': '1. Adapter service (Receiver, Handler, Dispatcher) is deployed and running with adapter.transport=in-memory\n2. IPL Mock is running and Postman is configured correctly\n3. UCRM is accessible via P2P connection\n4. Valid P2P authentication token is provisioned\n5. UCRI dispatch center enabled\n6. Logging infrastructure is operational',
            'steps': '1. Send a valid IPL TransferEinsatzRequest via REST API.\n   Method: POST\n   Endpoint: http://<hostname>:9080/einsatztransfer/erzeugeEinsatz\n2. Verify that the Receiver service accepts and validates the incoming request.\n3. Verify that the Handler service transforms the IPL payload into a UCRM-compliant request.\n4. Verify that the Handler authenticates to UCRM via P2P token.\n5. Verify that the Dispatcher forwards the UCRM-compliant request to UCRM via messaging/send endpoint.\n6. Verify that UCRM processes the message and returns an Accept response.\n7. Verify that the Handler translates the UCRM Accept ACK into IPL format.\n8. Verify the adapter logs for the ACK outcome.\n9. Verify the IPL Mock for the Acknowledgement.\n10. Verify end-to-end traceability via correlationId.',
            'expected': '• Incident (REST API) should be sent successfully. Postman receives HTTP 200 response.\n• Receiver validates the request successfully. No validation errors in logs.\n• Handler successfully converts the IPL TransferEinsatzRequest into the UCRM request format.\n• P2P authentication succeeds.\n• Message forwarded successfully to UCRM.\n• UCRM Accept response received by the Adapter.\n• IPL ACK is propagated to IPL Mock.\n• Log entry contains EMPFANG, confirming incident acceptance.\n• IPL Mock receives ACK Type=EMPFANG.\n• All log entries contain the same correlationId.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IPL-UCRI-02',
            'name': 'Verify that an IPL incident is delivered to UCRM through the Adapter and UCRM rejects the incident (EINSATZ_ABGELEHNT)',
            'category': 'Functional',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. IPL Mock is running and Postman is configured correctly\n3. UCRM is accessible via P2P connection\n4. Valid P2P authentication token is provisioned\n5. UCRI dispatch center enabled\n6. Logging infrastructure is operational',
            'steps': '1. Send a valid IPL TransferEinsatzRequest via REST API.\n2. Verify that the Receiver validates the request and forwards to Handler.\n3. Verify that the Handler transforms the IPL payload into UCRM-compliant format.\n4. Verify that the Dispatcher forwards the request to UCRM.\n5. Verify that UCRM processes the message and returns a Reject response.\n6. Verify that the Handler translates the UCRM Reject ACK into IPL format.\n7. Verify the adapter logs for the ACK outcome.\n8. Verify the IPL Mock for the Acknowledgement.',
            'expected': '• Incident (REST API) should be sent successfully.\n• Receiver validation passes.\n• Transformation and authentication succeed.\n• Message forwarded to UCRM successfully.\n• UCRM Reject response received by the Adapter.\n• Reject ACK is propagated to IPL Mock.\n• Log entry contains EINSATZ_ABGELEHNT.\n• IPL Mock receives ACK Type=EINSATZ_ABGELEHNT.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IPL-UCRI-03',
            'name': 'Verify timeout handling when UCRM does not respond within allowed time (TECHNISCHER_FEHLER)',
            'category': 'Functional',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. IPL Mock is running and Postman is configured correctly\n3. UCRM is either unavailable or configured to delay beyond timeout threshold\n4. Valid P2P authentication token is provisioned\n5. UCRI dispatch center enabled with timeout configured\n6. Logging infrastructure is operational',
            'steps': '1. Send a valid IPL TransferEinsatzRequest via REST API.\n2. Verify that the Receiver validates the request and forwards to Handler.\n3. Verify that the Handler transforms and authenticates.\n4. Verify that UCRM does not respond within the configured timeout period.\n5. Verify that the Handler detects the timeout condition.\n6. Verify that the Dispatcher calls /einsatztransfer/quittiereTranferEinsatz with the timeout error.\n7. Verify the adapter logs for the timeout ACK outcome.\n8. Verify the IPL Mock for the Acknowledgement.',
            'expected': '• Incident (REST API) should be sent successfully.\n• Receiver validation passes.\n• Dispatcher attempts to send to UCRM.\n• No response received from UCRM within allowed time.\n• Timeout detected. Handler processes the timeout as a technical error.\n• Timeout error propagated to IPL Mock.\n• Log entry contains TECHNISCHER_FEHLER.\n• IPL Mock receives ACK Type=TECHNISCHER_FEHLER.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IPL-UCRI-04',
            'name': 'Verify that an IPL incident is not delivered to UCRM when the Adapter Receiver rejects invalid data',
            'category': 'Negative',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. IPL Mock is running and Postman is configured correctly\n3. UCRM is accessible\n4. Logging infrastructure is operational',
            'steps': '1. Send an IPL incident with invalid/malformed data via REST API.\n2. Verify the Receiver service validates the incoming request.\n3. Verify that the Receiver returns an error response to the caller.\n4. Verify that no message was forwarded to the Handler service.\n5. Verify that no message was forwarded to UCRM via Dispatcher.\n6. Verify that no Einsatztransfer/quittiereTranferEinsatz log entry is generated.',
            'expected': '• Incident (REST API) request is sent.\n• Receiver validation fails. A clear rejection reason is logged.\n• Postman should receive an error response (4xx/5xx).\n• Handler logs show no event received.\n• Dispatcher logs show no dispatch attempted.\n• No ACK log entry generated since request was rejected at Receiver level.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-IPL-UCRI-05',
            'name': 'Verify that an IPL incident is not delivered to UCRM when mandatory fields are missing',
            'category': 'Negative',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. IPL Mock is running and Postman is configured correctly\n3. UCRM is accessible\n4. Logging infrastructure is operational',
            'steps': '1. Send an IPL incident with missing mandatory fields (e.g., no sourceMandant or targetMandant) via REST API.\n2. Verify the Receiver service validates and rejects the request due to missing mandatory fields.\n3. Verify the error response returned to Postman.\n4. Verify that no message was forwarded to Handler or Dispatcher.\n5. Verify the IPL Mock for the Acknowledgement.',
            'expected': '• Incident (REST API) request is sent.\n• Receiver validation fails with a clear error indicating which mandatory fields are missing.\n• Postman should receive an error response (4xx/5xx).\n• Handler and Dispatcher logs show no event received.\n• IPL Mock should return a transport error acknowledgment.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-U2U-01',
            'name': 'Verify that a mission from UCRI-1 client is successfully delivered to UCRI-2 client through the Adapter and UCRI-2 accepts (EMPFANG)',
            'category': 'Functional',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. UCRI-1 dispatch center enabled\n3. UCRI-2 dispatch center enabled\n4. Both UCRI-1 and UCRI-2 are accessible via P2P connection\n5. Valid P2P authentication tokens provisioned for both UCRI clients\n6. Logging infrastructure is operational',
            'steps': '1. UCRI-1 sends a valid mission payload to the Adapter with header.sender_client_id=UCRI-1 and header.receiver_client_id=UCRI-2.\n2. Verify that the Receiver validates the incoming UCRI-1 payload.\n3. Verify that the Handler identifies the target dispatch center as UCRI-2.\n4. Verify that the Handler authenticates to UCRI-2 via P2P token.\n5. Verify that the Handler transforms the payload if format differences exist.\n6. Verify that the Dispatcher routes the message to the UCRI-2 dispatch center topic.\n7. Verify that UCRI-2 receives and processes the mission and returns an Accept response.\n8. Verify that the Handler translates the UCRI-2 Accept ACK.\n9. Verify the adapter logs for the ACK outcome.\n10. Verify end-to-end traceability via correlationId.',
            'expected': '• Adapter Receiver accepts the mission payload from UCRI-1 successfully.\n• Receiver validates the mission. No validation errors.\n• Handler resolves UCRI-2 as the target dispatch center.\n• P2P authentication to UCRI-2 succeeds.\n• Payload transformed/passthrough as needed.\n• Message published to events-ucri-2 topic.\n• UCRI-2 Accept response received by the Adapter.\n• Accept ACK routed back to UCRI-1.\n• Log entry contains EMPFANG.\n• All log entries contain the same correlationId.',
            'priority': 'Medium',
            'status': 'Draft'
        },
        {
            'id': 'UCRI2-UCRI-IPL-01',
            'name': 'Verify that a UCRM mission is successfully delivered to IPL through the Adapter (reverse flow) and IPL accepts the mission',
            'category': 'Functional',
            'prereq': '1. Adapter service deployed and running with adapter.transport=in-memory\n2. UCRM is accessible and can send missions\n3. IPL system is operational (or mock IPL running)\n4. Reverse flow support enabled\n5. Logging infrastructure is operational',
            'steps': '1. UCRM sends a valid mission payload to the Adapter via messaging/send endpoint.\n2. Verify that the Receiver service accepts and validates the incoming UCRM mission.\n3. Verify that the Handler service translates the UCRM mission into IPL format.\n4. Verify that the Dispatcher forwards the translated mission to IPL.\n5. Verify that IPL processes the mission successfully and sends Accept ACK to Adapter.\n6. Verify that the Receiver receives IPL Accept ACK and Handler translates it.\n7. Verify that the Dispatcher propagates the Accept ACK to UCRM.\n8. Verify end-to-end traceability via correlationId across all 3 services.',
            'expected': '• Adapter Receiver accepts the mission payload successfully.\n• Receiver validates the UCRM mission. No validation errors.\n• Handler successfully translates UCRM mission fields to IPL-compatible format.\n• IPL receives the mission from the Adapter via Dispatcher.\n• IPL processes the mission and sends Accept acknowledgment back.\n• Handler translates IPL Accept ACK to UCRM format.\n• UCRM receives Accept acknowledgment from Adapter.\n• correlationId preserved across Receiver, Handler, and Dispatcher logs.',
            'priority': 'Medium',
            'status': 'Draft'
        },
    ]
    
    # Create a table for each test case
    for idx, tc in enumerate(test_cases):
        # Test Case Header
        p = doc.add_paragraph()
        run = p.add_run(f"{tc['id']}: {tc['name']}")
        run.bold = True
        run.font.size = Pt(11)
        
        # Category badge
        p2 = doc.add_paragraph()
        run = p2.add_run(f"Category: {tc['category']} | Priority: {tc['priority']} | Status: {tc['status']}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        # Create test case table
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        table.autofit = True
        
        # Pre-requisites
        table.rows[0].cells[0].text = 'Pre-Requisite'
        table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(table.rows[0].cells[0], 'E7E6E6')
        table.rows[0].cells[1].text = tc['prereq']
        
        # Verification Steps
        table.rows[1].cells[0].text = 'Verification Detail'
        table.rows[1].cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(table.rows[1].cells[0], 'E7E6E6')
        table.rows[1].cells[1].text = tc['steps']
        
        # Expected Outcome
        table.rows[2].cells[0].text = 'Expected Outcome'
        table.rows[2].cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(table.rows[2].cells[0], 'E7E6E6')
        table.rows[2].cells[1].text = tc['expected']
        
        # Reference & Status
        table.rows[3].cells[0].text = 'Reference & Status'
        table.rows[3].cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(table.rows[3].cells[0], 'E7E6E6')
        table.rows[3].cells[1].text = f"Test ID: {tc['id']}\nCategory: {tc['category']}\nPriority: {tc['priority']}\nStatus: {tc['status']}"
        
        # Set column widths
        for row in table.rows:
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(5.5)
        
        doc.add_paragraph()  # Add spacing
        
        if idx < len(test_cases) - 1:
            # Add a visual separator
            p = doc.add_paragraph()
            p.add_run('─' * 80)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
    
    # Add note about additional test cases
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run('Note: ')
    run.bold = True
    p.add_run('This document contains a representative sample of 12 out of 32 test cases. The complete test matrix includes additional test cases for IMT (In-Memory Transport), U2U (UCRI-to-UCRI), and UCRI-IPL (reverse flow) scenarios.')
    
    # Document Information Section
    doc.add_heading('Document Information', level=1)
    
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    
    info_data = [
        ('Document Version', '0.4'),
        ('Phase', 'UCRI 2.0 Phase II'),
        ('Generated Date', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('Test Cases Count', '32'),
        ('Document Type', 'Integration Test Plan'),
    ]
    
    for row_idx, (field, value) in enumerate(info_data):
        row = info_table.rows[row_idx]
        row.cells[0].text = field
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], 'E7E6E6')
        row.cells[1].text = value
    
    # Appendix
    doc.add_heading('Appendix: Test Execution Guidelines', level=1)
    
    doc.add_heading('Pre-requisites for Test Execution', level=2)
    prereq_items = [
        'Adapter service running with adapter.transport=in-memory',
        'All relevant modules enabled (IPL, UCRI dispatch centers)',
        'Postman/API testing tools configured',
        'Logging infrastructure operational',
        'Metrics/monitoring endpoint accessible',
    ]
    for item in prereq_items:
        doc.add_paragraph(item, style='List Number')
    
    doc.add_heading('Test Environment Configuration', level=2)
    config_items = [
        ('Transport Mode:', 'In-memory (for unit/integration testing)'),
        ('Logging Level:', 'DEBUG (for detailed verification)'),
        ('Timeout Settings:', 'Standard production values'),
    ]
    for label, value in config_items:
        p = doc.add_paragraph()
        run = p.add_run(label + ' ')
        run.bold = True
        p.add_run(value)
    
    doc.add_heading('Test Evidence Requirements', level=2)
    evidence_items = [
        'API request/response payloads',
        'System logs with correlationId traces',
        'Metrics dashboard screenshots',
        'Error logs for negative scenarios',
    ]
    for item in evidence_items:
        doc.add_paragraph(item, style='List Bullet')
    
    # Footer
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run('This document is auto-generated from UCRI-Phase II-TestCases V0.4.xlsx')
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Save the document
    output_path = r'c:\Apparao Maroju\AI-SSI\Project-IGNIS(POC)\UCRI2.0-Output\UCRI2-Phase-II-IntegrationTestPlan.docx'
    doc.save(output_path)
    print(f"Word document saved to: {output_path}")

if __name__ == '__main__':
    create_word_document()
