import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Delete Address - Test Cases"

# Styles
header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
data_font = Font(name='Calibri', size=10)
wrap_align = Alignment(wrap_text=True, vertical='top')
header_align = Alignment(wrap_text=True, vertical='center', horizontal='center')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

headers = [
    "TID", "Test Case Name", "Category", "Description",
    "Pre-conditions", "Step No", "Step Description", "Expected Result", "Priority"
]

col_widths = [12, 40, 15, 50, 50, 8, 55, 55, 10]

# Write headers
for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
ws.auto_filter.ref = f"A1:I1"

# Test case data
test_cases = [
    {"id": "TC-DEL-001", "name": "Delete non-default address — Main Flow", "category": "Functional - Positive",
     "desc": "Verify a non-default address with no contract assignments can be permanently deleted via burger menu",
     "precond": "1. User is authenticated and has write access to partner data.\n2. Partner has multiple addresses.\n3. At least one non-default address exists with no contract assignments.\n4. Reference date is current or future.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed with all partner addresses as compact address cards."),
         (2, "Locate the non-default address card and click the burger menu.", "Burger menu is displayed with options: Edit, Delete, Relocation."),
         (3, "Click the \"Delete\" action from the burger menu.", "Confirmation pop-up dialogue is displayed with message: \"Delete the address? <compact address> will be permanently deleted.\" with \"Cancel\" and \"Yes, delete\" buttons."),
         (4, "Click \"Yes, delete\" to confirm the deletion.", "Address is permanently deleted. Address is immediately removed from the Address Data overview. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-002", "name": "Delete non-default address — Cancel deletion", "category": "Functional - Negative",
     "desc": "Verify that cancelling the deletion of a non-default address preserves the address unchanged",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. Non-default address exists.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed with all partner addresses."),
         (2, "Click the burger menu on a non-default address card.", "Burger menu is displayed with Edit, Delete, Relocation options."),
         (3, "Click the \"Delete\" action from the burger menu.", "Confirmation pop-up dialogue is displayed: \"Delete the address? <compact address> will be permanently deleted.\""),
         (4, "Click \"Cancel\" in the pop-up dialogue.", "Pop-up is closed. Address remains unchanged in the Address Data section. No deletion occurs. No new session/version is created.")
     ]},

    {"id": "TC-DEL-003", "name": "Delete default address with replacement — Main Flow", "category": "Functional - Positive",
     "desc": "Verify that when deleting a default address with other addresses available, user is prompted to select a new default and deletion proceeds after selection",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses (at least 2).\n3. The selected address is the default address.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed with all addresses. Default address is identifiable."),
         (2, "Click the burger menu on the default address card.", "Burger menu is displayed with Edit, Delete, Relocation options."),
         (3, "Click the \"Delete\" action from the burger menu.", "Pop-up dialogue is displayed with message: \"Replace Default Address — <compact address> is currently the default address. Please choose a different address as the default.\" A dropdown titled \"New Default Address\" is shown listing all other partner addresses (excluding the address being deleted). \"Delete & Replace Default\" and \"Cancel\" buttons are displayed."),
         (4, "Select a new default address from the \"New Default Address\" dropdown.", "Address B is selected as the new default address."),
         (5, "Click \"Delete & Replace Default\" to confirm.", "The old default address (Address A) is permanently deleted and removed from the overview. Address B is set as the new default address. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-004", "name": "Delete default address — Cancel operation", "category": "Functional - Negative",
     "desc": "Verify that cancelling the deletion of a default address preserves all addresses and default assignment unchanged",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. The selected address is the default.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed."),
         (2, "Click the burger menu on the default address card and click \"Delete\".", "Pop-up dialogue \"Replace Default Address\" is displayed with dropdown and Cancel/Delete & Replace Default buttons."),
         (3, "Click \"Cancel\" in the pop-up dialogue.", "Pop-up is closed. All addresses remain unchanged. Default address assignment is preserved. No deletion occurs. No session/version is created.")
     ]},

    {"id": "TC-DEL-005", "name": "Delete default address — Verify dropdown excludes deleted address", "category": "Functional - Positive",
     "desc": "Verify the \"New Default Address\" dropdown lists all other addresses except the one being deleted",
     "precond": "1. User is authenticated with write access.\n2. Partner has 3 or more addresses.\n3. Selected address is the default.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module, open Address Data section.", "Address Data section displayed with 3 addresses."),
         (2, "Click burger menu on Address A (default) and click \"Delete\".", "Pop-up \"Replace Default Address\" is displayed."),
         (3, "Inspect the \"New Default Address\" dropdown list.", "Dropdown lists only Address B and Address C. Address A (the address being deleted) is NOT listed in the dropdown.")
     ]},

    {"id": "TC-DEL-006", "name": "Delete the only address — Confirm deletion", "category": "Functional - Positive",
     "desc": "Verify that when a partner has only one address, it can be deleted after confirmation with appropriate warning message",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.\n3. The address is the only and default address.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section displays one address card."),
         (2, "Click the burger menu on the only address card.", "Burger menu is displayed with Edit, Delete, Relocation options."),
         (3, "Click the \"Delete\" action from the burger menu.", "Pop-up dialogue is displayed with message: \"Delete the only address? <compact address> is the only address. Do you want to delete it?\" with \"Cancel\" and \"Yes, delete\" buttons."),
         (4, "Click \"Yes, delete\" to confirm the deletion.", "Address is permanently deleted and immediately removed from the Address Data overview. Address Data section is now empty. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-007", "name": "Delete the only address — Cancel deletion", "category": "Functional - Negative",
     "desc": "Verify that cancelling deletion of the only address preserves it unchanged",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section displays one address card."),
         (2, "Click burger menu on the only address card and click \"Delete\".", "Pop-up dialogue \"Delete the only address?\" is displayed."),
         (3, "Click \"Cancel\" in the pop-up dialogue.", "Pop-up is closed. The address remains unchanged. No deletion occurs.")
     ]},

    {"id": "TC-DEL-008", "name": "Delete action available for current reference date", "category": "Functional - Positive",
     "desc": "Verify the Delete action is available in the burger menu when the reference date is current",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to current date.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Partner Module, set reference date to current date, and open the Address Data section.", "Address Data section is displayed."),
         (2, "Click the burger menu on an address card.", "\"Delete\" action is available and clickable.")
     ]},

    {"id": "TC-DEL-009", "name": "Delete action available for future reference date", "category": "Functional - Positive",
     "desc": "Verify the Delete action is available in the burger menu when the reference date is in the future",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to a future date.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module, set reference date to a future date, and open the Address Data section.", "Address Data section is displayed."),
         (2, "Click the burger menu on an address card.", "\"Delete\" action is available and clickable.")
     ]},

    {"id": "TC-DEL-010", "name": "Delete action NOT available for past reference date", "category": "Functional - Negative",
     "desc": "Verify the Delete action is NOT available in the burger menu when the reference date is in the past",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to a past date.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Partner Module, set reference date to a past date, and open the Address Data section.", "Address Data section is displayed."),
         (2, "Click the burger menu on an address card.", "\"Delete\" action is NOT available or is disabled/hidden.")
     ]},

    {"id": "TC-DEL-011", "name": "Soft delete for address with valid-from date in the past", "category": "Functional - Positive",
     "desc": "Verify an address with a valid-from date in the past is soft deleted (not permanently removed from database)",
     "precond": "1. User is authenticated with write access.\n2. Partner has an address with valid-from date in the past.\n3. Reference date is current.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed."),
         (2, "Click burger menu on the address with past valid-from date and click \"Delete\".", "Appropriate confirmation pop-up is displayed."),
         (3, "Confirm the deletion.", "Address is soft deleted (marked as deleted but retained in database). Address is removed from the Address Data overview. A new session/version is created in the Change History reflecting the soft delete.")
     ]},

    {"id": "TC-DEL-012", "name": "Address already deleted by another user", "category": "Concurrency - Negative",
     "desc": "Verify the system handles the scenario where another user has already deleted the address before confirmation",
     "precond": "1. User A and User B both have write access to the same partner.\n2. Both users have the Address Data section open.\n3. Partner has at least one address.",
     "priority": "High",
     "steps": [
         (1, "User A navigates to the Partner Module and opens the Address Data section.", "Address Data section is displayed for User A."),
         (2, "User B deletes the same address and confirms deletion (in a separate session).", "User B successfully deletes the address."),
         (3, "User A clicks the burger menu on the same address and clicks \"Delete\", then confirms.", "The system refreshes the view. The address is no longer visible. No deletion operation takes place. System handles gracefully without error.")
     ]},

    {"id": "TC-DEL-013", "name": "Browser closed before confirming deletion", "category": "Functional - Negative",
     "desc": "Verify no deletion occurs if the browser is closed before the user confirms the action",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module, open the Address Data section.", "Address Data section is displayed."),
         (2, "Click burger menu on an address and click \"Delete\".", "Confirmation pop-up dialogue is displayed."),
         (3, "Close the browser without clicking \"Yes, delete\" or \"Cancel\".", "No deletion takes place. Address remains unchanged in the system."),
         (4, "Reopen the browser, log in, and navigate back to the partner's Address Data section.", "The address is still present and unchanged.")
     ]},

    {"id": "TC-DEL-014", "name": "Verify pop-up message for non-default address deletion", "category": "UI - Positive",
     "desc": "Validate the exact confirmation message and button labels for deleting a non-default address",
     "precond": "1. User is authenticated with write access.\n2. Partner has a non-default address \"Mailbox 1234, 12345 Sample City\".",
     "priority": "Low",
     "steps": [
         (1, "Navigate to the Address Data section and click burger menu on the non-default address.", "Burger menu is displayed."),
         (2, "Click \"Delete\" from the burger menu.", "Pop-up displays title: \"Delete the address?\" Body: \"Mailbox 1234, 12345 Sample City will be permanently deleted\". Two buttons: \"Cancel\" and \"Yes, delete\" are visible.")
     ]},

    {"id": "TC-DEL-015", "name": "Verify pop-up message for default address deletion (multiple addresses)", "category": "UI - Positive",
     "desc": "Validate the exact confirmation message, dropdown, and button labels for deleting a default address when other addresses exist",
     "precond": "1. User is authenticated with write access.\n2. Partner has default address and other addresses.",
     "priority": "Low",
     "steps": [
         (1, "Navigate to the Address Data section and click burger menu on the default address.", "Burger menu is displayed."),
         (2, "Click \"Delete\" from the burger menu.", "Pop-up displays title: \"Replace Default Address\". Body: \"<compact address> is currently the default address. Please choose a different address as the default.\" Dropdown titled \"New Default Address\" is displayed. Two buttons: \"Cancel\" and \"Delete & Replace Default\" are visible.")
     ]},

    {"id": "TC-DEL-016", "name": "Verify pop-up message for only address deletion", "category": "UI - Positive",
     "desc": "Validate the exact confirmation message and button labels for deleting the only address",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.",
     "priority": "Low",
     "steps": [
         (1, "Navigate to the Address Data section and click burger menu on the only address.", "Burger menu is displayed."),
         (2, "Click \"Delete\" from the burger menu.", "Pop-up displays title: \"Delete the only address?\" Body: \"<compact address> is the only address. Do you want to delete it?\" Two buttons: \"Cancel\" and \"Yes, delete\" are visible.")
     ]},

    {"id": "TC-DEL-017", "name": "Verify Delete action is present in burger menu for each address", "category": "UI - Positive",
     "desc": "Confirm the Delete action is available in the burger menu of every address card",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. Reference date is current.",
     "priority": "Medium",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section shows 3 address cards."),
         (2, "Click the burger menu on each address card one at a time.", "Each burger menu displays the \"Delete\" action along with \"Edit\" and \"Relocation\" options.")
     ]},

    {"id": "TC-DEL-018", "name": "Verify burger menu options — Edit, Delete, Relocation", "category": "UI - Positive",
     "desc": "Validate that the burger menu contains exactly three options: Edit, Delete, and Relocation",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "priority": "Low",
     "steps": [
         (1, "Navigate to the Address Data section and click the burger menu on any address card.", "Burger menu displays exactly three actions: \"Edit\", \"Delete\", \"Relocation\". Options are displayed in correct order.")
     ]},

    {"id": "TC-DEL-019", "name": "Expand/Collapse Address Data section using chevron", "category": "UI - Positive",
     "desc": "Verify the Address Data section can be opened and closed using the chevron control",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "priority": "Low",
     "steps": [
         (1, "Navigate to the Partner Module. Locate the Address Data section (collapsed).", "Address Data section header is visible with a chevron control."),
         (2, "Click the chevron to expand the Address Data section.", "Address Data section expands and displays all partner addresses as compact address cards."),
         (3, "Click the chevron again to collapse the section.", "Address Data section collapses and address cards are hidden.")
     ]},

    {"id": "TC-DEL-020", "name": "Verify deletion creates new session/version in Change History", "category": "Audit - Positive",
     "desc": "Validate that a successful address deletion is recorded in the Change History with appropriate details",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one deletable address.",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Address Data section and delete an address (confirm deletion).", "Address is permanently deleted and removed from overview."),
         (2, "Navigate to the Change History section of the partner.", "A new session/version entry is visible in the Change History. The entry contains: timestamp, user ID, and the deletion action. The change is traceable and readable.")
     ]},

    {"id": "TC-DEL-021", "name": "Verify cancelled deletion does NOT create Change History entry", "category": "Audit - Negative",
     "desc": "Validate that a cancelled deletion does not create any entry in the Change History",
     "precond": "1. User is authenticated with write access.\n2. Partner has addresses.",
     "priority": "Medium",
     "steps": [
         (1, "Initiate delete action on an address and click \"Cancel\" in the pop-up.", "Address remains unchanged."),
         (2, "Navigate to the Change History section.", "No new session/version entry is created for the cancelled action.")
     ]},

    {"id": "TC-DEL-022", "name": "Verify deletion is irreversible", "category": "Functional - Positive",
     "desc": "Confirm that once an address is deleted and confirmed, it cannot be recovered or undone",
     "precond": "1. User is authenticated with write access.\n2. Partner has a non-default address.",
     "priority": "High",
     "steps": [
         (1, "Note down the address details of the address to be deleted.", "Address details are recorded."),
         (2, "Delete the address via burger menu → Delete → \"Yes, delete\".", "Address is permanently deleted and removed from overview."),
         (3, "Verify there is no \"Undo\" option or restore functionality available.", "No undo or restore functionality is available. The deleted address cannot be recovered."),
         (4, "Navigate away and return to the Address Data section.", "The deleted address is not displayed.")
     ]},

    {"id": "TC-DEL-023", "name": "Verify address overview updates immediately after deletion", "category": "UI - Positive",
     "desc": "Confirm the Address Data section reflects the updated state immediately after a successful deletion",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.",
     "priority": "High",
     "steps": [
         (1, "Note the current number of addresses displayed.", "3 address cards displayed."),
         (2, "Delete one non-default address and confirm.", "Address is deleted."),
         (3, "Observe the Address Data section immediately after confirmation.", "Address Data section immediately shows only 2 address cards. The deleted address is no longer visible. No page refresh is required.")
     ]},

    {"id": "TC-DEL-024", "name": "Verify address overview unchanged after cancellation", "category": "UI - Negative",
     "desc": "Confirm the Address Data section reflects no changes after a cancelled deletion",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.",
     "priority": "Medium",
     "steps": [
         (1, "Note the current state of addresses displayed.", "3 address cards displayed."),
         (2, "Initiate delete on an address and click \"Cancel\".", "Pop-up is closed."),
         (3, "Observe the Address Data section.", "All 3 address cards remain displayed and unchanged.")
     ]},

    {"id": "TC-DEL-025", "name": "Read-only user cannot delete addresses", "category": "Security - Negative",
     "desc": "Verify a user without write permission cannot access or execute the Delete action",
     "precond": "1. User is authenticated with read-only access to partner data.\n2. Partner has at least one address.",
     "priority": "Critical",
     "steps": [
         (1, "Navigate to the Partner Module and open the Address Data section.", "Address Data section is displayed with partner addresses."),
         (2, "Click the burger menu on an address card.", "\"Delete\" action is either not visible or disabled/greyed out for the read-only user.")
     ]},

    {"id": "TC-DEL-026", "name": "Unauthenticated user redirected to login", "category": "Security - Negative",
     "desc": "Verify an unauthenticated user is redirected to the login page when attempting to access the Partner Module",
     "precond": "1. User is not authenticated (no active session).",
     "priority": "Critical",
     "steps": [
         (1, "Attempt to navigate directly to the Partner Module's Address Data section URL.", "User is automatically redirected to the login page. Access to the Partner Module is denied.")
     ]},

    {"id": "TC-DEL-027", "name": "Session expired during deletion process", "category": "Security - Negative",
     "desc": "Verify the system handles session expiry during an active deletion process",
     "precond": "1. User is authenticated with write access.\n2. Session timeout is configured (e.g., 15 minutes).",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Address Data section and click \"Delete\" on an address.", "Confirmation pop-up is displayed."),
         (2, "Wait for the session to expire (e.g., 15+ minutes of inactivity).", "Session expires. User is automatically logged out."),
         (3, "Attempt to confirm the deletion by clicking \"Yes, delete\".", "No deletion occurs. User is redirected to the login page. Address remains unchanged.")
     ]},

    {"id": "TC-DEL-028", "name": "System error during deletion — user-friendly error message", "category": "Error Handling - Negative",
     "desc": "Verify the system displays a user-friendly error message when a technical error occurs during deletion (e.g., API failure)",
     "precond": "1. User is authenticated with write access.\n2. A technical error is simulated (e.g., API failure, network issue).",
     "priority": "High",
     "steps": [
         (1, "Navigate to the Address Data section, click burger menu → Delete on an address.", "Confirmation pop-up is displayed."),
         (2, "Confirm the deletion (while a backend error is simulated).", "A user-friendly error message is displayed (no technical details exposed). The address remains unchanged. The error is logged with timestamp, user ID, and affected functionality.")
     ]},

    {"id": "TC-DEL-029", "name": "Network failure during deletion", "category": "Error Handling - Negative",
     "desc": "Verify deletion does not proceed if network connectivity is lost during the confirmation step",
     "precond": "1. User is authenticated with write access.\n2. Partner has addresses.",
     "priority": "High",
     "steps": [
         (1, "Click burger menu → Delete on an address. Confirmation pop-up is displayed.", "Confirmation pop-up is displayed."),
         (2, "Disconnect from the network.", "Network is disconnected."),
         (3, "Click \"Yes, delete\" to confirm.", "Deletion does not proceed. A user-friendly error message is displayed. Address remains unchanged after reconnection.")
     ]},

    {"id": "TC-DEL-030", "name": "Verify deletion action is logged", "category": "Audit - Positive",
     "desc": "Confirm that a successful address deletion is logged with timestamp, user ID, and affected functionality",
     "precond": "1. User is authenticated with write access.\n2. Partner has a deletable address.",
     "priority": "High",
     "steps": [
         (1, "Delete an address via burger menu → Delete → confirm.", "Address is deleted."),
         (2, "Verify the system logs.", "Log entry exists with: timestamp, user ID, action performed (delete address), and affected partner/address details. Logs are GDPR-compliant (no sensitive personal data exposed in plain text).")
     ]},
]

row = 2
for tc in test_cases:
    first_step = True
    for step_no, step_desc, expected in tc["steps"]:
        if first_step:
            # Write the main test case info in the first row
            ws.cell(row=row, column=1, value=tc["id"]).font = data_font
            ws.cell(row=row, column=2, value=tc["name"]).font = data_font
            ws.cell(row=row, column=3, value=tc["category"]).font = data_font
            ws.cell(row=row, column=4, value=tc["desc"]).font = data_font
            ws.cell(row=row, column=5, value=tc["precond"]).font = data_font
            ws.cell(row=row, column=9, value=tc["priority"]).font = data_font
            first_step = False
        else:
            # Leave main info blank for subsequent steps
            pass

        # Write step info
        ws.cell(row=row, column=6, value=step_no).font = data_font
        ws.cell(row=row, column=7, value=step_desc).font = data_font
        ws.cell(row=row, column=8, value=expected).font = data_font

        # Apply formatting to all cells in the row
        for c in range(1, 10):
            ws.cell(row=row, column=c).alignment = wrap_align
            ws.cell(row=row, column=c).border = thin_border

        row += 1

    # Insert a blank row after each test case
    row += 1

# Save
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
filename = f"TestCases_{timestamp}.xlsx"
filepath = f"c:\\Apparao Maroju\\AI-SSI\\Project-Nexina(POC)\\Nex-Output\\{filename}"
wb.save(filepath)
print(f"Excel file saved: {filepath}")
print(f"Total test cases: {len(test_cases)}")
