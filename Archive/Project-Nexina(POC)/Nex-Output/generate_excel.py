import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Delete Address - Test Cases"

# Styles
header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
group_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
group_font = Font(name='Calibri', bold=True, size=11)
data_font = Font(name='Calibri', size=10)
wrap_align = Alignment(wrap_text=True, vertical='top')
header_align = Alignment(wrap_text=True, vertical='center', horizontal='center')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

headers = [
    "Scenario Number", "Test Summary", "Test Type", "Description",
    "Test Repository", "Precondition", "Test Steps", "Test Data", "Expected Results"
]

col_widths = [16, 40, 12, 50, 35, 50, 55, 45, 55]

# Write headers
for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
ws.auto_filter.ref = f"A1:I1"

# Test case data
test_cases = [
    # Group separator
    {"group": "Scenario Group 1: Delete Non-Default Address (No Contract Assignments) — Role: User with Write Access"},

    {"id": "TC-DEL-001", "summary": "Delete non-default address — Main Flow", "type": "Positive",
     "desc": "Verify a non-default address with no contract assignments can be permanently deleted via burger menu",
     "repo": "Delete Address - Main Flow",
     "precond": "1. User is authenticated and has write access to partner data.\n2. Partner has multiple addresses.\n3. At least one non-default address exists with no contract assignments.\n4. Reference date is current or future.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with multiple addresses (e.g., Address A — default, Address B — non-default)",
          "Address Data section is displayed with all partner addresses as compact address cards."),
         ("Step 2: Locate the non-default address card and click the burger menu.",
          "Address B (non-default, no contract assignments)",
          "Burger menu is displayed with options: Edit, Delete, Relocation."),
         ("Step 3: Click the \"Delete\" action from the burger menu.",
          "",
          "Confirmation pop-up dialogue is displayed with message: \"Delete the address? <compact address> will be permanently deleted.\" with \"Cancel\" and \"Yes, delete\" buttons."),
         ("Step 4: Click \"Yes, delete\" to confirm the deletion.",
          "",
          "Address is permanently deleted. Address is immediately removed from the Address Data overview. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-002", "summary": "Delete non-default address — Cancel deletion", "type": "Negative",
     "desc": "Verify that cancelling the deletion of a non-default address preserves the address unchanged",
     "repo": "Delete Address - Alternative Flow",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. Non-default address exists.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with multiple addresses",
          "Address Data section is displayed with all partner addresses."),
         ("Step 2: Click the burger menu on a non-default address card.",
          "Non-default address",
          "Burger menu is displayed with Edit, Delete, Relocation options."),
         ("Step 3: Click the \"Delete\" action from the burger menu.",
          "",
          "Confirmation pop-up dialogue is displayed: \"Delete the address? <compact address> will be permanently deleted.\""),
         ("Step 4: Click \"Cancel\" in the pop-up dialogue.",
          "",
          "Pop-up is closed. Address remains unchanged in the Address Data section. No deletion occurs. No new session/version is created.")
     ]},

    {"group": "Scenario Group 2: Delete Default Address (Multiple Addresses Exist)"},

    {"id": "TC-DEL-003", "summary": "Delete default address with replacement — Main Flow", "type": "Positive",
     "desc": "Verify that when deleting a default address with other addresses available, user is prompted to select a new default and deletion proceeds after selection",
     "repo": "Delete Address - Default Address (Multiple)",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses (at least 2).\n3. The selected address is the default address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with Address A (default) and Address B, Address C (non-default)",
          "Address Data section is displayed with all addresses. Default address is identifiable."),
         ("Step 2: Click the burger menu on the default address card.",
          "Address A (default)",
          "Burger menu is displayed with Edit, Delete, Relocation options."),
         ("Step 3: Click the \"Delete\" action from the burger menu.",
          "",
          "Pop-up dialogue is displayed with message: \"Replace Default Address — <compact address> is currently the default address. Please choose a different address as the default.\" A dropdown titled \"New Default Address\" is shown listing all other partner addresses (excluding the address being deleted). \"Delete & Replace Default\" and \"Cancel\" buttons are displayed."),
         ("Step 4: Select a new default address from the \"New Default Address\" dropdown.",
          "Select Address B from dropdown",
          "Address B is selected as the new default address."),
         ("Step 5: Click \"Delete & Replace Default\" to confirm.",
          "",
          "The old default address (Address A) is permanently deleted and removed from the overview. Address B is set as the new default address. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-004", "summary": "Delete default address — Cancel operation", "type": "Negative",
     "desc": "Verify that cancelling the deletion of a default address preserves all addresses and default assignment unchanged",
     "repo": "Delete Address - Default Address Cancel",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. The selected address is the default.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with default address and other addresses",
          "Address Data section is displayed."),
         ("Step 2: Click the burger menu on the default address card and click \"Delete\".",
          "Default address card",
          "Pop-up dialogue \"Replace Default Address\" is displayed with dropdown and Cancel/Delete & Replace Default buttons."),
         ("Step 3: Click \"Cancel\" in the pop-up dialogue.",
          "",
          "Pop-up is closed. All addresses remain unchanged. Default address assignment is preserved. No deletion occurs. No session/version is created.")
     ]},

    {"id": "TC-DEL-005", "summary": "Delete default address — Verify dropdown excludes deleted address", "type": "Positive",
     "desc": "Verify the \"New Default Address\" dropdown lists all other addresses except the one being deleted",
     "repo": "Delete Address - Default Address Dropdown",
     "precond": "1. User is authenticated with write access.\n2. Partner has 3 or more addresses.\n3. Selected address is the default.",
     "steps": [
         ("Step 1: Navigate to the Partner Module, open Address Data section.",
          "Partner with Address A (default), Address B, Address C",
          "Address Data section displayed with 3 addresses."),
         ("Step 2: Click burger menu on Address A (default) and click \"Delete\".",
          "",
          "Pop-up \"Replace Default Address\" is displayed."),
         ("Step 3: Inspect the \"New Default Address\" dropdown list.",
          "",
          "Dropdown lists only Address B and Address C. Address A (the address being deleted) is NOT listed in the dropdown.")
     ]},

    {"group": "Scenario Group 3: Delete the Only Address (Single Address Partner)"},

    {"id": "TC-DEL-006", "summary": "Delete the only address — Confirm deletion", "type": "Positive",
     "desc": "Verify that when a partner has only one address, it can be deleted after confirmation with appropriate warning message",
     "repo": "Delete Address - Only Address",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.\n3. The address is the only and default address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with a single address (Address A)",
          "Address Data section displays one address card."),
         ("Step 2: Click the burger menu on the only address card.",
          "Address A (only address)",
          "Burger menu is displayed with Edit, Delete, Relocation options."),
         ("Step 3: Click the \"Delete\" action from the burger menu.",
          "",
          "Pop-up dialogue is displayed with message: \"Delete the only address? <compact address> is the only address. Do you want to delete it?\" with \"Cancel\" and \"Yes, delete\" buttons."),
         ("Step 4: Click \"Yes, delete\" to confirm the deletion.",
          "",
          "Address is permanently deleted and immediately removed from the Address Data overview. Address Data section is now empty. A new session/version is created in the Change History.")
     ]},

    {"id": "TC-DEL-007", "summary": "Delete the only address — Cancel deletion", "type": "Negative",
     "desc": "Verify that cancelling deletion of the only address preserves it unchanged",
     "repo": "Delete Address - Only Address Cancel",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with a single address",
          "Address Data section displays one address card."),
         ("Step 2: Click burger menu on the only address card and click \"Delete\".",
          "",
          "Pop-up dialogue \"Delete the only address?\" is displayed."),
         ("Step 3: Click \"Cancel\" in the pop-up dialogue.",
          "",
          "Pop-up is closed. The address remains unchanged. No deletion occurs.")
     ]},

    {"group": "Scenario Group 4: Reference Date Validations"},

    {"id": "TC-DEL-008", "summary": "Delete action available for current reference date", "type": "Positive",
     "desc": "Verify the Delete action is available in the burger menu when the reference date is current",
     "repo": "Delete Address - Reference Date",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to current date.",
     "steps": [
         ("Step 1: Navigate to the Partner Module, set reference date to current date, and open the Address Data section.",
          "Reference date = today's date",
          "Address Data section is displayed."),
         ("Step 2: Click the burger menu on an address card.",
          "",
          "Burger menu is displayed. \"Delete\" action is available and clickable.")
     ]},

    {"id": "TC-DEL-009", "summary": "Delete action available for future reference date", "type": "Positive",
     "desc": "Verify the Delete action is available in the burger menu when the reference date is in the future",
     "repo": "Delete Address - Reference Date Future",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to a future date.",
     "steps": [
         ("Step 1: Navigate to the Partner Module, set reference date to a future date, and open the Address Data section.",
          "Reference date = future date (e.g., 2027-01-01)",
          "Address Data section is displayed."),
         ("Step 2: Click the burger menu on an address card.",
          "",
          "Burger menu is displayed. \"Delete\" action is available and clickable.")
     ]},

    {"id": "TC-DEL-010", "summary": "Delete action NOT available for past reference date", "type": "Negative",
     "desc": "Verify the Delete action is NOT available in the burger menu when the reference date is in the past",
     "repo": "Delete Address - Reference Date Past",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.\n3. Reference date is set to a past date.",
     "steps": [
         ("Step 1: Navigate to the Partner Module, set reference date to a past date, and open the Address Data section.",
          "Reference date = past date (e.g., 2025-01-01)",
          "Address Data section is displayed."),
         ("Step 2: Click the burger menu on an address card.",
          "",
          "Burger menu is displayed. \"Delete\" action is NOT available or is disabled/hidden.")
     ]},

    {"group": "Scenario Group 5: Soft Delete — Address with Valid-From Date in the Past"},

    {"id": "TC-DEL-011", "summary": "Soft delete for address with valid-from date in the past", "type": "Positive",
     "desc": "Verify an address with a valid-from date in the past is soft deleted (not permanently removed from database)",
     "repo": "Delete Address - Soft Delete",
     "precond": "1. User is authenticated with write access.\n2. Partner has an address with valid-from date in the past.\n3. Reference date is current.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with an address having valid-from date = 2025-01-01 (past)",
          "Address Data section is displayed."),
         ("Step 2: Click burger menu on the address with past valid-from date and click \"Delete\".",
          "",
          "Appropriate confirmation pop-up is displayed."),
         ("Step 3: Confirm the deletion.",
          "",
          "Address is soft deleted (marked as deleted but retained in database). Address is removed from the Address Data overview. A new session/version is created in the Change History reflecting the soft delete.")
     ]},

    {"group": "Scenario Group 6: Concurrent User Scenarios"},

    {"id": "TC-DEL-012", "summary": "Address already deleted by another user", "type": "Negative",
     "desc": "Verify the system handles the scenario where another user has already deleted the address before confirmation",
     "repo": "Delete Address - Concurrent Deletion",
     "precond": "1. User A and User B both have write access to the same partner.\n2. Both users have the Address Data section open.\n3. Partner has at least one address.",
     "steps": [
         ("Step 1: User A navigates to the Partner Module and opens the Address Data section.",
          "Shared partner data between User A and User B",
          "Address Data section is displayed for User A."),
         ("Step 2: User B deletes the same address and confirms deletion (in a separate session).",
          "",
          "User B successfully deletes the address."),
         ("Step 3: User A clicks the burger menu on the same address and clicks \"Delete\", then confirms.",
          "",
          "The system refreshes the view. The address is no longer visible. No deletion operation takes place. System handles gracefully without error.")
     ]},

    {"id": "TC-DEL-013", "summary": "Browser closed before confirming deletion", "type": "Negative",
     "desc": "Verify no deletion occurs if the browser is closed before the user confirms the action",
     "repo": "Delete Address - Browser Close",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module, open the Address Data section.",
          "Any partner with addresses",
          "Address Data section is displayed."),
         ("Step 2: Click burger menu on an address and click \"Delete\".",
          "",
          "Confirmation pop-up dialogue is displayed."),
         ("Step 3: Close the browser without clicking \"Yes, delete\" or \"Cancel\".",
          "",
          "No deletion takes place. Address remains unchanged in the system."),
         ("Step 4: Reopen the browser, log in, and navigate back to the partner's Address Data section.",
          "",
          "The address is still present and unchanged.")
     ]},

    {"group": "Scenario Group 7: UI Validations — Confirmation Pop-Up Dialogues"},

    {"id": "TC-DEL-014", "summary": "Verify pop-up message for non-default address deletion", "type": "Positive",
     "desc": "Validate the exact confirmation message and button labels for deleting a non-default address",
     "repo": "Delete Address - UI Validation",
     "precond": "1. User is authenticated with write access.\n2. Partner has a non-default address \"Mailbox 1234, 12345 Sample City\".",
     "steps": [
         ("Step 1: Navigate to the Address Data section and click burger menu on the non-default address.",
          "Address: \"Mailbox 1234, 12345 Sample City\" (non-default)",
          "Burger menu is displayed."),
         ("Step 2: Click \"Delete\" from the burger menu.",
          "",
          "Pop-up displays title: \"Delete the address?\" Body: \"Mailbox 1234, 12345 Sample City will be permanently deleted\". Two buttons: \"Cancel\" and \"Yes, delete\" are visible.")
     ]},

    {"id": "TC-DEL-015", "summary": "Verify pop-up message for default address deletion (multiple addresses)", "type": "Positive",
     "desc": "Validate the exact confirmation message, dropdown, and button labels for deleting a default address when other addresses exist",
     "repo": "Delete Address - UI Validation Default",
     "precond": "1. User is authenticated with write access.\n2. Partner has default address and other addresses.",
     "steps": [
         ("Step 1: Navigate to the Address Data section and click burger menu on the default address.",
          "Default address: \"Mailbox 1234, 12345 Sample City\"; Other addresses exist",
          "Burger menu is displayed."),
         ("Step 2: Click \"Delete\" from the burger menu.",
          "",
          "Pop-up displays title: \"Replace Default Address\". Body: \"<compact address> is currently the default address. Please choose a different address as the default.\" Dropdown titled \"New Default Address\" is displayed. Two buttons: \"Cancel\" and \"Delete & Replace Default\" are visible.")
     ]},

    {"id": "TC-DEL-016", "summary": "Verify pop-up message for only address deletion", "type": "Positive",
     "desc": "Validate the exact confirmation message and button labels for deleting the only address",
     "repo": "Delete Address - UI Validation Only Address",
     "precond": "1. User is authenticated with write access.\n2. Partner has exactly one address.",
     "steps": [
         ("Step 1: Navigate to the Address Data section and click burger menu on the only address.",
          "Single address: \"Mailbox 1234, 12345 Sample City\"",
          "Burger menu is displayed."),
         ("Step 2: Click \"Delete\" from the burger menu.",
          "",
          "Pop-up displays title: \"Delete the only address?\" Body: \"<compact address> is the only address. Do you want to delete it?\" Two buttons: \"Cancel\" and \"Yes, delete\" are visible.")
     ]},

    {"group": "Scenario Group 8: Burger Menu Validations"},

    {"id": "TC-DEL-017", "summary": "Verify Delete action is present in burger menu for each address", "type": "Positive",
     "desc": "Confirm the Delete action is available in the burger menu of every address card",
     "repo": "Delete Address - Burger Menu",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.\n3. Reference date is current.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Partner with 3 addresses",
          "Address Data section shows 3 address cards."),
         ("Step 2: Click the burger menu on each address card one at a time.",
          "",
          "Each burger menu displays the \"Delete\" action along with \"Edit\" and \"Relocation\" options.")
     ]},

    {"id": "TC-DEL-018", "summary": "Verify burger menu options — Edit, Delete, Relocation", "type": "Positive",
     "desc": "Validate that the burger menu contains exactly three options: Edit, Delete, and Relocation",
     "repo": "Delete Address - Burger Menu Options",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "steps": [
         ("Step 1: Navigate to the Address Data section and click the burger menu on any address card.",
          "Any address card",
          "Burger menu displays exactly three actions: \"Edit\", \"Delete\", \"Relocation\". Options are displayed in correct order.")
     ]},

    {"group": "Scenario Group 9: Address Data Section — Chevron Toggle"},

    {"id": "TC-DEL-019", "summary": "Expand/Collapse Address Data section using chevron", "type": "Positive",
     "desc": "Verify the Address Data section can be opened and closed using the chevron control",
     "repo": "Delete Address - UI Navigation",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module. Locate the Address Data section (collapsed).",
          "Any partner",
          "Address Data section header is visible with a chevron control."),
         ("Step 2: Click the chevron to expand the Address Data section.",
          "",
          "Address Data section expands and displays all partner addresses as compact address cards."),
         ("Step 3: Click the chevron again to collapse the section.",
          "",
          "Address Data section collapses and address cards are hidden.")
     ]},

    {"group": "Scenario Group 10: Change History / Audit Trail Validations"},

    {"id": "TC-DEL-020", "summary": "Verify deletion creates new session/version in Change History", "type": "Positive",
     "desc": "Validate that a successful address deletion is recorded in the Change History with appropriate details",
     "repo": "Delete Address - Change History",
     "precond": "1. User is authenticated with write access.\n2. Partner has at least one deletable address.",
     "steps": [
         ("Step 1: Navigate to the Address Data section and delete an address (confirm deletion).",
          "Non-default address to delete",
          "Address is permanently deleted and removed from overview."),
         ("Step 2: Navigate to the Change History section of the partner.",
          "",
          "A new session/version entry is visible in the Change History. The entry contains: timestamp, user ID, and the deletion action. The change is traceable and readable.")
     ]},

    {"id": "TC-DEL-021", "summary": "Verify cancelled deletion does NOT create Change History entry", "type": "Negative",
     "desc": "Validate that a cancelled deletion does not create any entry in the Change History",
     "repo": "Delete Address - Change History Cancel",
     "precond": "1. User is authenticated with write access.\n2. Partner has addresses.",
     "steps": [
         ("Step 1: Initiate delete action on an address and click \"Cancel\" in the pop-up.",
          "",
          "Address remains unchanged."),
         ("Step 2: Navigate to the Change History section.",
          "",
          "No new session/version entry is created for the cancelled action.")
     ]},

    {"group": "Scenario Group 11: Irreversibility of Deletion"},

    {"id": "TC-DEL-022", "summary": "Verify deletion is irreversible", "type": "Positive",
     "desc": "Confirm that once an address is deleted and confirmed, it cannot be recovered or undone",
     "repo": "Delete Address - Irreversibility",
     "precond": "1. User is authenticated with write access.\n2. Partner has a non-default address.",
     "steps": [
         ("Step 1: Note down the address details of the address to be deleted.",
          "Non-default address with known details",
          "Address details are recorded."),
         ("Step 2: Delete the address via burger menu → Delete → \"Yes, delete\".",
          "",
          "Address is permanently deleted and removed from overview."),
         ("Step 3: Verify there is no \"Undo\" option or restore functionality available.",
          "",
          "No undo or restore functionality is available. The deleted address cannot be recovered."),
         ("Step 4: Navigate away and return to the Address Data section.",
          "",
          "The deleted address is not displayed.")
     ]},

    {"group": "Scenario Group 12: Immediate UI Refresh After Deletion"},

    {"id": "TC-DEL-023", "summary": "Verify address overview updates immediately after deletion", "type": "Positive",
     "desc": "Confirm the Address Data section reflects the updated state immediately after a successful deletion",
     "repo": "Delete Address - Immediate Refresh",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.",
     "steps": [
         ("Step 1: Note the current number of addresses displayed.",
          "Partner with 3 addresses",
          "3 address cards displayed."),
         ("Step 2: Delete one non-default address and confirm.",
          "",
          "Address is deleted."),
         ("Step 3: Observe the Address Data section immediately after confirmation.",
          "",
          "Address Data section immediately shows only 2 address cards. The deleted address is no longer visible. No page refresh is required.")
     ]},

    {"id": "TC-DEL-024", "summary": "Verify address overview unchanged after cancellation", "type": "Negative",
     "desc": "Confirm the Address Data section reflects no changes after a cancelled deletion",
     "repo": "Delete Address - Immediate Refresh Cancel",
     "precond": "1. User is authenticated with write access.\n2. Partner has multiple addresses.",
     "steps": [
         ("Step 1: Note the current state of addresses displayed.",
          "Partner with 3 addresses",
          "3 address cards displayed."),
         ("Step 2: Initiate delete on an address and click \"Cancel\".",
          "",
          "Pop-up is closed."),
         ("Step 3: Observe the Address Data section.",
          "",
          "All 3 address cards remain displayed and unchanged.")
     ]},

    {"group": "Scenario Group 13: Authorization / Role-Based Access — User WITHOUT Write Access (Read-Only)"},

    {"id": "TC-DEL-025", "summary": "Read-only user cannot delete addresses", "type": "Negative",
     "desc": "Verify a user without write permission cannot access or execute the Delete action",
     "repo": "Delete Address - Authorization",
     "precond": "1. User is authenticated with read-only access to partner data.\n2. Partner has at least one address.",
     "steps": [
         ("Step 1: Navigate to the Partner Module and open the Address Data section.",
          "Read-only user account",
          "Address Data section is displayed with partner addresses."),
         ("Step 2: Click the burger menu on an address card.",
          "",
          "Burger menu is displayed. \"Delete\" action is either not visible or disabled/greyed out for the read-only user.")
     ]},

    {"group": "Scenario Group 14: Authentication Validations — Unauthenticated User"},

    {"id": "TC-DEL-026", "summary": "Unauthenticated user redirected to login", "type": "Negative",
     "desc": "Verify an unauthenticated user is redirected to the login page when attempting to access the Partner Module",
     "repo": "Delete Address - Authentication",
     "precond": "1. User is not authenticated (no active session).",
     "steps": [
         ("Step 1: Attempt to navigate directly to the Partner Module's Address Data section URL.",
          "No authentication token/session",
          "User is automatically redirected to the login page. Access to the Partner Module is denied.")
     ]},

    {"id": "TC-DEL-027", "summary": "Session expired during deletion process", "type": "Negative",
     "desc": "Verify the system handles session expiry during an active deletion process",
     "repo": "Delete Address - Session Expiry",
     "precond": "1. User is authenticated with write access.\n2. Session timeout is configured (e.g., 15 minutes).",
     "steps": [
         ("Step 1: Navigate to the Address Data section and click \"Delete\" on an address.",
          "",
          "Confirmation pop-up is displayed."),
         ("Step 2: Wait for the session to expire (e.g., 15+ minutes of inactivity).",
          "",
          "Session expires. User is automatically logged out."),
         ("Step 3: Attempt to confirm the deletion by clicking \"Yes, delete\".",
          "",
          "No deletion occurs. User is redirected to the login page. Address remains unchanged.")
     ]},

    {"group": "Scenario Group 15: Error Handling"},

    {"id": "TC-DEL-028", "summary": "System error during deletion — user-friendly error message", "type": "Negative",
     "desc": "Verify the system displays a user-friendly error message when a technical error occurs during deletion (e.g., API failure)",
     "repo": "Delete Address - Error Handling",
     "precond": "1. User is authenticated with write access.\n2. A technical error is simulated (e.g., API failure, network issue).",
     "steps": [
         ("Step 1: Navigate to the Address Data section, click burger menu → Delete on an address.",
          "",
          "Confirmation pop-up is displayed."),
         ("Step 2: Confirm the deletion (while a backend error is simulated).",
          "",
          "A user-friendly error message is displayed (no technical details exposed). The address remains unchanged. The error is logged with timestamp, user ID, and affected functionality.")
     ]},

    {"id": "TC-DEL-029", "summary": "Network failure during deletion", "type": "Negative",
     "desc": "Verify deletion does not proceed if network connectivity is lost during the confirmation step",
     "repo": "Delete Address - Network Failure",
     "precond": "1. User is authenticated with write access.\n2. Partner has addresses.",
     "steps": [
         ("Step 1: Click burger menu → Delete on an address. Confirmation pop-up is displayed.",
          "",
          "Confirmation pop-up is displayed."),
         ("Step 2: Disconnect from the network.",
          "",
          "Network is disconnected."),
         ("Step 3: Click \"Yes, delete\" to confirm.",
          "",
          "Deletion does not proceed. A user-friendly error message is displayed. Address remains unchanged after reconnection.")
     ]},

    {"group": "Scenario Group 16: Logging"},

    {"id": "TC-DEL-030", "summary": "Verify deletion action is logged", "type": "Positive",
     "desc": "Confirm that a successful address deletion is logged with timestamp, user ID, and affected functionality",
     "repo": "Delete Address - Logging",
     "precond": "1. User is authenticated with write access.\n2. Partner has a deletable address.",
     "steps": [
         ("Step 1: Delete an address via burger menu → Delete → confirm.",
          "Non-default address",
          "Address is deleted."),
         ("Step 2: Verify the system logs.",
          "",
          "Log entry exists with: timestamp, user ID, action performed (delete address), and affected partner/address details. Logs are GDPR-compliant (no sensitive personal data exposed in plain text).")
     ]},
]

row = 2
for tc in test_cases:
    if "group" in tc:
        # Write group header row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=1, value=tc["group"])
        cell.font = group_font
        cell.fill = group_fill
        cell.alignment = Alignment(vertical='center')
        cell.border = thin_border
        for c in range(2, 10):
            ws.cell(row=row, column=c).fill = group_fill
            ws.cell(row=row, column=c).border = thin_border
        ws.row_dimensions[row].height = 25
        row += 1
        continue

    first_step = True
    for step_text, test_data, expected in tc["steps"]:
        if first_step:
            ws.cell(row=row, column=1, value=tc["id"]).font = data_font
            ws.cell(row=row, column=2, value=tc["summary"]).font = data_font
            ws.cell(row=row, column=3, value=tc["type"]).font = data_font
            ws.cell(row=row, column=4, value=tc["desc"]).font = data_font
            ws.cell(row=row, column=5, value=tc["repo"]).font = data_font
            ws.cell(row=row, column=6, value=tc["precond"]).font = data_font
            first_step = False
        else:
            ws.cell(row=row, column=1, value=tc["id"]).font = data_font

        ws.cell(row=row, column=7, value=step_text).font = data_font
        ws.cell(row=row, column=8, value=test_data).font = data_font
        ws.cell(row=row, column=9, value=expected).font = data_font

        for c in range(1, 10):
            ws.cell(row=row, column=c).alignment = wrap_align
            ws.cell(row=row, column=c).border = thin_border

        row += 1

# Save
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
filename = f"TestCases_{timestamp}.xlsx"
filepath = f"c:\\Apparao Maroju\\AI-SSI\\Project-Nexina(POC)\\Nex-Output\\{filename}"
wb.save(filepath)
print(f"Excel file saved: {filepath}")
print(f"Total test scenarios: 30")
print(f"Total rows (including step rows): {row - 2}")
